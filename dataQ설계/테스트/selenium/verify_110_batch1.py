# -*- coding: utf-8 -*-
"""
110번 기존 기능 보완 1묶음(표준사전·표준 승인·대시보드 — 2.4·3.12·8.5 절) 조치 검증 — API + 화면

  S1. 통합 검색 (2.4-2-a·2-b·3·4) — 사전 구분, 도메인 승인 조건, 컬럼 사용·승인 조건과 모델명·테이블 한글명
  S2. 대시보드 (2.4-8-b·10-b·11) — 표준 현황 응답, 지운 API, 구조 변경 감지는 완료(DONE) 진단만
  S3. 표준 사전 등록·수정·삭제 (3.12-3-b·5-e-2·7-a·b·c·d·g·9-b·10-a·b·12-d·15-e, 4-b 확인)
  S4. 엑셀 다운로드 (3.12-6-a·b·c) — 고른 사전만, 검색어, 0건이면 헤더만
  S5. 엑셀 일괄등록 (3.12-7-e·12-b·12-c) — 등록 경로·상세 건, 요청시스템, 영문약어 규칙·유사어, 코드 이력 한 건
  S6. 표준 승인 (8.5-4·6·7·8·9, 8.5-3b 확인) — 도메인 반려 가드, 사유 50자, 반려 이력, 연쇄 반려 요청자, 요청일
  S7. 표준사전 관리 (3.12-11-b) — 기본 사전·없는 사전 사용 중지 거부
  S8. 화면 (2.4-2-a·7, 3.12-7-f·11-a·12-a·12-f·13-c·13-e·14-c·15-f, 8.5-3a·5·6·13·14) — 셀레니움 Edge(기본 창 없음)

필요:
  - q-center(28091)·q-executor(28098) 기동
  - 환경 변수 PGPASSWORD (개발 DB admin), DATAQ_TEST_PW (시험 계정 space·jyjang 비밀번호), 선택 PSQL (psql.exe 경로)
  - 인자로 돌릴 묶음을 고를 수 있다 (예: python verify_110_batch1.py S1 S3). 없으면 전부
시험 데이터(시험 사전과 그 안의 단어·용어·코드·도메인, 모델, 구조 진단 이력, 승인·변경 이력, 알림)는 끝에서 지운다.
"""
import base64, io, json, os, subprocess, sys, tempfile, time, traceback
import requests
import openpyxl

BASE = "http://localhost:28091"
TAG = time.strftime("%H%M%S")
PSQL = os.environ.get("PSQL", r"C:\Program Files\PostgreSQL\13\bin\psql.exe")
PW = os.environ["DATAQ_TEST_PW"]
results = []
cleanup = []          # 끝에서 거꾸로 실행할 정리 함수
FX = {}               # 공통 시험 데이터 (사전·도메인·단어)
TMP = tempfile.mkdtemp(prefix="v110_")


def step(name, fn):
    print(f"\n=== {name}")
    try:
        fn()
        results.append((name, "PASS"))
        print("  >> PASS")
    except Exception:
        traceback.print_exc()
        results.append((name, "FAIL"))


def psql(sql):
    r = subprocess.run([PSQL, "-h", "localhost", "-p", "25433", "-U", "admin", "-d", "postgres", "-w", "-At", "-c", sql],
                       capture_output=True, text=True, timeout=60, encoding="utf-8", errors="replace",
                       env=dict(os.environ, PGCLIENTENCODING="UTF8"))
    if r.returncode != 0:
        raise RuntimeError("psql 실패: " + (r.stderr or "")[:300])
    return (r.stdout or "").strip()


def login_api(user, pwd=None):
    s = requests.Session()
    r = s.post(BASE + "/login", data={"id": user, "password": base64.b64encode((pwd or PW).encode()).decode()},
               allow_redirects=False, timeout=10)
    assert r.status_code == 200, f"로그인 실패 {user} {r.status_code}"
    return s


def post(s, path, body, timeout=30):
    r = s.post(BASE + path, json=body, timeout=timeout)
    try:
        return r.json()
    except Exception:
        return {"http": r.status_code, "text": r.text[:300]}


def get(s, path, params=None, timeout=30):
    r = s.get(BASE + path, params=params, timeout=timeout)
    try:
        return r.json()
    except Exception:
        return {"http": r.status_code, "text": r.text[:300]}


def contents(r):
    c = r.get("contents")
    return json.loads(c) if isinstance(c, str) and c else (c or {})


def code(r):
    return r.get("resultCode")


def msg(r):
    return r.get("resultMessage") or ""


def q(v):
    return "'" + str(v).replace("'", "''") + "'"


def hist(target_id, extra=""):
    """대상의 변경 이력 — 유형:대상유형:등록경로 (시간순)"""
    return psql(f"select string_agg(change_type||':'||target_type||':'||coalesce(change_source,''), ',' order by change_dt, change_id) "
                f"from quality.tb_change_history where target_id='{target_id}' {extra}")


def word(s, nm, abrv, dict_id, synonyms=None, clsf="N"):
    r = post(s, "/api/std/createWord", {"wordNm": nm, "wordEngAbrvNm": abrv, "wordEngNm": abrv.title(), "wordDesc": "110 시험",
                                        "wordClsfYn": clsf, "dictId": dict_id, "allophSynmLst": synonyms})
    wid = psql(f"select word_id from quality.tb_word where word_nm={q(nm)} and dict_id='{dict_id}'")
    assert wid, f"단어 등록 실패: {r}"
    return wid


# ---------------------------------------------------------------- 공통 준비
def fixture(sa):
    """시험 사전 D 와 그 안의 도메인 그룹·분류·도메인(승인), 단어 W1(승인)"""
    r = post(sa, "/api/dict/save", {"dictNm": "ZZ110사전" + TAG, "dictDesc": "110 시험", "termLastWordClsfYn": "N",
                                    "upperYn": "Y", "wordEngNmDupYn": "Y", "termEngMaxLen": "", "oldStdYn": "N"})
    assert code(r) == 200, r
    d = contents(r)["dictId"]
    FX["dict"] = d
    FX["grp"], FX["clsf"], FX["dom"] = "ZZ110그룹" + TAG, "ZZ110분류" + TAG, "ZZ110도메인" + TAG

    def _drop():
        for sql in [
            f"delete from quality.tb_aprv_stats where req_item_nm like 'ZZ110%'",
            f"delete from quality.tb_code_integ where dict_id='{d}'",
            f"delete from quality.tb_code_data where dict_id='{d}'",
            f"delete from quality.tb_terms_words where terms_id in (select terms_id from quality.tb_terms where dict_id='{d}')",
            f"delete from quality.tb_terms where dict_id='{d}'",
            f"delete from quality.tb_word where dict_id='{d}'",
            f"delete from quality.tb_domain where dict_id='{d}'",
            f"delete from quality.tb_domain_clsf where dict_id='{d}'",
            f"delete from quality.tb_domain_grp where dict_id='{d}'",
            f"delete from quality.tb_std_dict where dict_id='{d}'",
        ]:
            psql(sql)
    cleanup.append(_drop)
    psql(f"insert into quality.tb_domain_grp (domain_grp_id, domain_grp_nm, comm_stnd_yn, dict_id) values ('Z1G{TAG}', {q(FX['grp'])}, 'N', '{d}')")
    psql(f"insert into quality.tb_domain_clsf (domain_clsf_id, domain_clsf_nm, domain_grp_nm, comm_stnd_yn, dict_id) "
         f"values ('Z1C{TAG}', {q(FX['clsf'])}, {q(FX['grp'])}, 'N', '{d}')")
    psql(f"insert into quality.tb_domain (domain_id, domain_nm, domain_grp_nm, domain_clsf_nm, domain_desc, data_type, data_len, aprv_yn, dict_id) "
         f"values ('Z1D{TAG}', {q(FX['dom'])}, {q(FX['grp'])}, {q(FX['clsf'])}, '110 시험', 'VARCHAR', 30, 'Y', '{d}')")
    FX["domId"] = f"Z1D{TAG}"
    FX["w1nm"], FX["w1ab"] = "ZZ110일" + TAG, "ZZA" + TAG
    FX["w1"] = word(sa, FX["w1nm"], FX["w1ab"], d)
    print(f"  시험 사전={d} 도메인={FX['dom']} 단어 W1={FX['w1nm']}")


def make_model(s, nm, dict_id):
    r = post(s, "/api/dm/createDataModel", {"dataModelNm": nm, "dataModelDsId": None, "modelType": "PHYSICAL", "ver": "1",
                                            "dictId": dict_id, "stndReqYn": "N"})
    dm = psql(f"select dm_id from quality.tb_data_model where dm_nm={q(nm)}")
    assert code(r) == 200 and dm, r

    def _drop():
        for t in ("tb_data_model_change_history", "tb_data_model_attr", "tb_data_model_constraint", "tb_data_model_index",
                  "tb_data_model_obj", "tb_data_model_clct", "tb_data_model_schema", "tb_udp_val", "tb_term_resolve_history"):
            psql(f"delete from quality.{t} where dm_id='{dm}'")
        psql(f"delete from quality.tb_struct_diag_history where data_model_id='{dm}'")
        psql(f"delete from quality.tb_data_model where dm_id='{dm}'")
    cleanup.append(_drop)
    return dm


def xlsx(path, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def xlsx_rows(content):
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


WORD_HDR = ["No", "제정차수", "단어명", "단어영문약어명", "단어영문명", "단어설명", "형식단어여부", "도메인분류명", "이음동의어목록",
            "금칙어목록", "요청시스템", "표준여부"]
TERM_HDR = ["No", "제정차수", "용어명", "용어설명", "용어영문약어명", "도메인명", "허용값", "저장형식", "표현형식", "코드그룹명",
            "소관기관명", "이음동의어목록", "요청시스템", "표준여부", "개인정보여부", "암호화여부"]


def upload(s, path_api, file_path, dict_id):
    with open(file_path, "rb") as f:
        r = s.post(BASE + path_api, files={"file": (os.path.basename(file_path), f.read(),
                                                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                   data={"dictId": dict_id}, timeout=30)
    return r


def wait_bulk(target_type, since, timeout=40):
    """그 뒤에 생긴 일괄등록 이력 {changeId, cnt, source, summary}"""
    end = time.time() + timeout
    while time.time() < end:
        v = psql(f"select change_id||'|'||change_cnt||'|'||coalesce(change_source,'')||'|'||summary from quality.tb_change_history "
                 f"where change_type='BULK_INSERT' and target_type='{target_type}' and change_user_id='space' "
                 f"and change_dt >= to_timestamp('{since}', 'YYYYMMDDHH24MISS') order by change_dt desc limit 1")
        if v:
            cid, cnt, src, summ = v.split("|", 3)
            return {"changeId": cid, "cnt": int(cnt), "source": src, "summary": summ}
        time.sleep(1)
    raise AssertionError(f"{target_type} 일괄등록 이력이 {timeout}초 안에 안 생겼다")


# ---------------------------------------------------------------- S1
def s1(sa, su):
    d = FX["dict"]

    def _dict():
        kw = FX["w1nm"]
        in_d = get(sa, "/api/std/search", {"keyword": kw, "dictId": d})
        in_def = get(sa, "/api/std/search", {"keyword": kw, "dictId": "DEFAULT"})
        no_dict = get(sa, "/api/std/search", {"keyword": kw})
        print(f"  단어 {kw}: 시험 사전={len(in_d['words'])} 기본 사전={len(in_def['words'])} 사전 없이={len(no_dict['words'])}")
        assert len(in_d["words"]) == 1 and in_d["words"][0]["wordNm"] == kw
        assert len(in_def["words"]) == 0 and len(no_dict["words"]) == 0
    step("S1-1. 통합 검색 — 고른 사전 안에서만 찾는다 (사전을 안 보내면 기본 사전)", _dict)

    def _domain():
        pend = "ZZ110대기도메인" + TAG
        psql(f"insert into quality.tb_domain (domain_id, domain_nm, domain_grp_nm, domain_clsf_nm, domain_desc, data_type, data_len, aprv_yn, dict_id) "
             f"values ('Z1P{TAG}', {q(pend)}, {q(FX['grp'])}, {q(FX['clsf'])}, '110 시험', 'VARCHAR', 10, 'N', '{d}')")
        a = get(sa, "/api/std/search", {"keyword": FX["dom"], "dictId": d})["domains"]
        p = get(sa, "/api/std/search", {"keyword": pend, "dictId": d})["domains"]
        print(f"  승인 도메인 {len(a)}건 / 승인 대기 도메인 {len(p)}건")
        assert len(a) == 1 and len(p) == 0
    step("S1-2. 통합 검색 — 도메인도 승인된 것만 (단어·용어와 같다)", _domain)

    def _columns():
        nm = "ZZ110모델_" + TAG
        dm = make_model(sa, nm, d)
        FX["dm"], FX["dmNm"] = dm, nm
        obj = "ZZ110_T" + TAG
        assert code(post(sa, "/api/dm/addObj", {"dataModelId": dm, "objOwner": "", "objNm": obj, "objNmKr": "시험테이블"})) == 200
        r = post(sa, "/api/dm/saveAttrs", {"dataModelId": dm, "objNm": obj, "objOwner": "", "attrs": [
            {"mode": "ADD", "attrNm": "C1", "attrNmKr": "열하나", "dataType": "VARCHAR", "dataLen": 10},
            {"mode": "ADD", "attrNm": "C2", "attrNmKr": "열둘", "dataType": "VARCHAR", "dataLen": 10}]})
        assert code(r) == 200, r
        cols = get(sa, "/api/std/search", {"keyword": obj})["columns"]
        print(f"  컬럼 {len(cols)}건: {cols}")
        assert len(cols) == 2 and all(c["dataModelNm"] == nm and c["objNmKr"] == "시험테이블" for c in cols)
        psql(f"update quality.tb_data_model_attr set use_yn='N' where dm_id='{dm}' and attr_nm='C2'")
        a = get(sa, "/api/std/search", {"keyword": obj})["columns"]
        psql(f"update quality.tb_data_model_attr set aprv_status='DRAFT' where dm_id='{dm}' and attr_nm='C1'")
        b = get(sa, "/api/std/search", {"keyword": obj})["columns"]
        psql(f"update quality.tb_data_model_attr set use_yn='Y', aprv_status='APPROVED' where dm_id='{dm}'")
        print(f"  삭제 대기(USE_YN=N) 뒤 {len(a)}건 / 승인 대기(DRAFT) 뒤 {len(b)}건")
        assert [c["attrNm"] for c in a] == ["C1"] and b == []
    step("S1-3. 통합 검색 — 컬럼은 사용 중·승인된 것만, 모델명·테이블 한글명을 함께", _columns)


# ---------------------------------------------------------------- S2
def s2(sa, su):
    def _dash():
        r = get(sa, "/api/search/getDashboardInfo")
        # 없는 주소는 화면(index.html)으로 넘어간다 — JSON 이 아니면 API 가 없는 것
        gone = sa.get(BASE + "/api/search/getDataModelStat", timeout=10).headers.get("Content-Type", "")
        gone2 = sa.get(BASE + "/api/search/getTopDataModelList", timeout=10).headers.get("Content-Type", "")
        print(f"  표준 현황 키={sorted(r.keys())} / getDataModelStat 응답={gone} getTopDataModelList 응답={gone2}")
        assert sorted(r.keys()) == ["domainCnt", "termsCnt", "wordCnt"] and "json" not in gone and "json" not in gone2
    step("S2-1. 대시보드 표준 현황은 건수 3개만, 호출 화면이 없던 API 2개는 없다", _dash)

    def _struct():
        dm = FX.get("dm") or make_model(sa, "ZZ110모델_" + TAG, FX["dict"])
        FX["dm"] = dm
        psql(f"insert into quality.tb_struct_diag_history (diag_id, data_model_id, status, diag_dt, total_tables, total_columns, added_columns) "
             f"values ('Z1S1{TAG}', '{dm}', 'DONE', now() - interval '1 hour', 1, 2, 1), "
             f"('Z1S2{TAG}', '{dm}', 'RUNNING', now(), 0, 0, 0), ('Z1S3{TAG}', '{dm}', 'FAILED', now() - interval '1 minute', 0, 0, 0)")
        r = get(sa, "/api/std/structDiag/latestByModel", {"dataModelId": dm})
        h = r.get("history") or {}
        print(f"  최신 구조 진단 = {h.get('diagId')} {h.get('status')} (실행 중·실패가 더 최근)")
        assert h.get("diagId") == f"Z1S1{TAG}" and h.get("status") == "DONE"
    step("S2-2. 구조 변경 감지(latestByModel)는 최신 완료(DONE) 진단 — 실행 중·실패 진단을 건너뜀", _struct)


# ---------------------------------------------------------------- S3
def s3(sa, su):
    d, dom = FX["dict"], FX["dom"]
    st = {}

    def _word():
        w2nm = "ZZ110이" + TAG
        r = post(sa, "/api/std/createWord", {"wordNm": w2nm, "wordEngAbrvNm": FX["w1ab"], "wordEngNm": "Zzb" + TAG, "wordDesc": "110",
                                            "wordClsfYn": "N", "dictId": d})
        print(f"  같은 사전에 같은 영문약어로 등록 = {code(r)} {msg(r)}")
        assert code(r) == 500 and msg(r) == f"이미 등록된 단어 영문약어입니다: '{FX['w1ab']}'"
        r = post(sa, "/api/std/updateWord", {"id": FX["w1"], "wordNm": FX["w1nm"], "wordEngAbrvNm": FX["w1ab"], "wordEngNm": FX["w1ab"].title(),
                                            "wordDesc": "110 수정", "wordClsfYn": "N"})
        cur = psql(f"select curr_value from quality.tb_change_history where target_id='{FX['w1']}' and change_type='UPDATE'")
        print(f"  단어 수정={code(r)} / 이력 변경값의 승인상태 = {'aprvYn=Y' in cur}")
        assert code(r) == 200 and "aprvYn=Y" in cur and "aprvYn=N" not in cur
    step("S3-1. 단어 — 영문약어 중복은 지금 인덱스(tb_word_ux_2)로 안내, 수정 이력의 승인상태는 그대로(Y)", _word)

    def _term_domain():
        r = post(sa, "/api/std/createTerms", {"termsNm": "ZZ110무도메인" + TAG, "termsEngAbrvNm": FX["w1ab"], "domainNm": "",
                                             "dictId": d, "wordList": [{"wordId": FX["w1"], "wordNm": FX["w1nm"], "wordOrd": 0}]})
        print(f"  도메인 없이 용어 등록 = {code(r)} {msg(r)}")
        assert code(r) == 500 and msg(r) == "도메인명은 필수입니다."
    step("S3-2. 용어 등록 — 도메인이 비면 '도메인명은 필수입니다.' (DB 오류 원문 아님)", _term_domain)

    def _code_hist():
        grp = "Z1K" + TAG
        st["grp"], st["cnm"] = grp, "ZZ110코드" + TAG
        r = post(sa, "/api/std/createCode", {"termsNm": st["cnm"], "termsEngAbrvNm": FX["w1ab"], "termsDesc": "110", "domainNm": dom,
                                            "codeGrp": grp, "codeTp": "SINGLE", "integSubGrps": [], "dictId": d,
                                            "wordList": [{"wordId": FX["w1"], "wordNm": FX["w1nm"], "wordOrd": 0}]})
        assert code(r) == 200, r
        cid = psql(f"select terms_id from quality.tb_terms where terms_nm={q(st['cnm'])} and dict_id='{d}'")
        st["cid"] = cid
        h1 = hist(cid)
        for v in ("01", "02"):
            r = post(sa, "/api/std/createCodeData", {"codeGrp": grp, "codeNm": "ZZ110값" + TAG, "codeEngNm": "VAL", "codeVal": v,
                                                    "codeValDesc": v, "parentCodeVal": None, "dictId": d})
            assert code(r) == 200, r
        body = {"id": cid, "termsNm": st["cnm"], "termsEngAbrvNm": FX["w1ab"], "termsDesc": "110 수정", "domainNm": dom, "codeGrp": grp,
                "codeTp": "SINGLE", "integSubGrps": [], "commStndYn": "N", "wordList": [], "allophSynmLst": [], "dictId": d}
        r = post(sa, "/api/std/updateCode", body)
        words = psql(f"select count(*) from quality.tb_terms_words where terms_id='{cid}'")
        h2 = hist(cid)
        r2 = post(su, "/api/std/updateCode", dict(body, termsDesc="권한 없는 수정"))
        h3 = hist(cid)
        print(f"  코드 등록 이력={h1}\n  단어 구성 없이 수정={code(r)} {msg(r)} 구성 단어={words} 이력={h2}\n  권한 없는 수정={code(r2)} {msg(r2)} 이력={h3}")
        assert h1 == "INSERT:CODE:ADMIN_DIRECT"
        assert code(r) == 200 and words == "1" and h2 == "INSERT:CODE:ADMIN_DIRECT,UPDATE:CODE:ADMIN_DIRECT"
        assert code(r2) == 500 and h3 == h2
        st["body"] = body
    step("S3-3. 코드 — 등록 이력은 CODE 한 건, 단어 구성 없이 수정 200(구성 그대로), 실패한 수정은 이력 없음", _code_hist)

    def _last_word():
        body = st["body"]
        save = {"dictId": d, "dictNm": "ZZ110사전" + TAG, "termEngMaxLen": ""}
        assert code(post(sa, "/api/dict/save", dict(save, termLastWordClsfYn="Y"))) == 200
        try:
            r1 = post(sa, "/api/std/updateCode", dict(body, wordList=[{"termsId": st["cid"], "wordId": FX["w1"], "wordNm": FX["w1nm"], "wordOrd": 0}]))
            r2 = post(sa, "/api/std/updateCode", dict(body, termsDesc="구성 그대로"))
        finally:
            assert code(post(sa, "/api/dict/save", dict(save, termLastWordClsfYn="N"))) == 200
        print(f"  형식단어 강제 사전에서 비형식단어로 구성 수정={code(r1)} {msg(r1)} / 구성 그대로 수정={code(r2)}")
        assert code(r1) == 500 and "마지막 단어는 형식단어" in msg(r1) and code(r2) == 200
    step("S3-4. 용어·코드 수정 — 사전이 형식단어를 강제하면 등록처럼 마지막 단어 검사 (구성을 바꿀 때)", _last_word)

    def _domain_del():
        r = post(sa, "/api/std/deleteDomains", [{"id": FX["domId"]}])
        left = psql(f"select count(*) from quality.tb_domain where domain_id='{FX['domId']}'")
        print(f"  용어가 쓰는 도메인 삭제 = {code(r)} {msg(r)} / 남음={left}")
        assert code(r) == 500 and msg(r).startswith("다음 용어에서 사용 중이므로 삭제할 수 없습니다: ") and st["cnm"] in msg(r) and left == "1"
        d2 = "ZZ110도메인둘" + TAG
        psql(f"insert into quality.tb_domain (domain_id, domain_nm, domain_grp_nm, domain_clsf_nm, domain_desc, data_type, data_len, aprv_yn, dict_id) "
             f"values ('Z1E{TAG}', {q(d2)}, {q(FX['grp'])}, {q(FX['clsf'])}, '110 시험', 'VARCHAR', 10, 'Y', '{d}')")
        r = post(sa, "/api/std/deleteDomains", [{"id": f"Z1E{TAG}"}])
        h = psql(f"select target_nm||'|'||(prev_value like '%domainNm={d2}%') from quality.tb_change_history where target_id='Z1E{TAG}' and change_type='DELETE'")
        print(f"  쓰지 않는 도메인 삭제 = {code(r)} / 이력 대상명|이전 값 = {h}")
        assert code(r) == 200 and h == f"{d2}|true"
    step("S3-5. 도메인 삭제 — 쓰는 용어가 있으면 이름을 알려 거부, 지우면 항목별 이력(대상명·이전 값)", _domain_del)

    def _code_del():
        d_ = d
        vid = psql(f"select code_id from quality.tb_code_data where dict_id='{d_}' and code_grp='{st['grp']}' and code_val='02'")
        r = post(sa, "/api/std/deleteCodeDatas", [{"id": vid}])
        h = psql(f"select target_nm||'|'||(prev_value like '%codeVal=02%') from quality.tb_change_history where target_id='{vid}' and change_type='DELETE'")
        print(f"  코드값 삭제 = {code(r)} / 이력 = {h}")
        assert code(r) == 200 and h == f"ZZ110값{TAG}|true"
        r = post(sa, "/api/std/deleteCodeList", [{"id": st["cid"]}])
        left = psql(f"select count(*) from quality.tb_code_data where dict_id='{d_}' and code_grp='{st['grp']}'")
        h = hist(st["cid"])
        prev = psql(f"select prev_value like '%termsNm={st['cnm']}%' from quality.tb_change_history where target_id='{st['cid']}' and change_type='DELETE'")
        withv = psql(f"select count(*) from quality.tb_change_history where target_type='CODE_DATA' and target_nm='{st['grp']}' "
                     f"and summary like '%코드값 1건 함께 삭제%'")
        print(f"  코드 삭제 = {code(r)} {msg(r)} / 남은 코드값={left} / 이력={h} 이전 값={prev} / 함께 지운 코드값 이력={withv}")
        assert code(r) == 200 and left == "0" and h.endswith("DELETE:CODE:ADMIN_DIRECT") and prev == "t" and withv == "1"
    step("S3-6. 코드값·코드 삭제 — 항목별 이력, 코드를 지우면 그 코드값도 함께 지움", _code_del)

    def _word_del():
        w2 = word(sa, "ZZ110이" + TAG, "ZZB" + TAG, d)
        r = post(sa, "/api/std/deleteWords", [{"id": w2}])
        h = psql(f"select target_nm||'|'||(prev_value like '%wordNm=ZZ110이{TAG}%') from quality.tb_change_history where target_id='{w2}' and change_type='DELETE'")
        print(f"  단어 삭제(화면처럼 id 만) = {code(r)} / 이력 = {h}")
        assert code(r) == 200 and h == f"ZZ110이{TAG}|true"
    step("S3-7. 단어 삭제 — id 만 보내도 이력 대상명·이전 값", _word_del)

    def _cand():
        pend = "ZZ110대기도메인" + TAG
        exists = psql(f"select count(*) from quality.tb_domain where domain_nm={q(pend)}")
        if exists == "0":
            psql(f"insert into quality.tb_domain (domain_id, domain_nm, domain_grp_nm, domain_clsf_nm, domain_desc, data_type, data_len, aprv_yn, dict_id) "
                 f"values ('Z1P{TAG}', {q(pend)}, {q(FX['grp'])}, {q(FX['clsf'])}, '110 시험', 'VARCHAR', 10, 'N', '{d}')")
        a = [x["domainNm"] for x in post(sa, "/api/std/getDomainList", {"dictId": d, "schAprvYn": "Y"})]
        b = [x["domainNm"] for x in post(sa, "/api/std/getDomainList", {"dictId": d})]
        c = [x["domainNm"] for x in get(sa, "/api/std/getDomainInfoByClsfNm", {"clsfNm": FX["clsf"], "dictId": d, "schAprvYn": "Y"})]
        e = [x["domainNm"] for x in get(sa, "/api/std/getDomainInfoByClsfNm", {"clsfNm": FX["clsf"], "dictId": d})]
        print(f"  목록(승인만)={a} 목록(조건 없음)={b}\n  분류별(승인만)={c} 분류별(조건 없음)={e}")
        assert pend not in a and pend in b and pend not in c and pend in e and dom in a and dom in c
    step("S3-8. 용어·코드 화면의 도메인 후보 — 승인 조건(schAprvYn=Y)으로 승인 대기 도메인 제외", _cand)

    def _perm():
        r = post(su, "/api/std/createDomainGroup", {"domainGrpNm": "ZZ110권한" + TAG, "commStndYn": "N", "dictId": d})
        print(f"  일반 사용자 도메인 그룹 등록 = {code(r)} {msg(r)}")
        assert code(r) == 403
    step("S3-9. (이미 고쳐짐 확인, 109 #13) 일반 사용자의 도메인 그룹 등록 403", _perm)


# ---------------------------------------------------------------- S4
def s4(sa, su):
    d = FX["dict"]

    def _dl(path, params):
        r = sa.get(BASE + path, params=params, timeout=30)
        assert r.status_code == 200, (path, r.status_code)
        assert len(r.content) > 0, (path, "0바이트")
        return xlsx_rows(r.content)

    def _words():
        rows = _dl("/api/std/downloadWords", {"dictId": d})
        names = sorted(x[2] for x in rows[1:])
        want = sorted(psql(f"select word_nm from quality.tb_word where dict_id='{d}' and aprv_yn='Y'").splitlines())
        dflt = _dl("/api/std/downloadWords", {"searchKey": FX["w1nm"]})
        print(f"  시험 사전 단어 다운로드 {names} (DB {want}) / 사전 없이(기본 사전) 같은 이름 {len(dflt) - 1}건")
        assert names == want and FX["w1nm"] in names and len(dflt) == 1
    step("S4-1. 단어 다운로드 — 고른 사전만 (사전 없이 부르면 기본 사전)", _words)

    def _search():
        g = _dl("/api/std/downloadDomainGroups", {"dictId": d, "searchKey": "ZZ110그룹"})
        c = _dl("/api/std/downloadDomainClsfs", {"dictId": d, "searchKey": "없는검색어" + TAG})
        print(f"  도메인 그룹(검색어) 행={[x[1] for x in g[1:]]} / 없는 검색어 분류 = 헤더 {c[0][:3]}… 행 {len(c) - 1}")
        assert [x[1] for x in g[1:]] == [FX["grp"]] and len(c) == 1 and c[0][0] == "No"
    step("S4-2. 도메인 그룹·분류 다운로드 — 검색어로 거르고, 0건이면 헤더만 있는 파일 (0바이트 아님)", _search)


# ---------------------------------------------------------------- S5
def s5(sa, su):
    d = FX["dict"]

    def _words():
        syn = "ZZ110유사" + TAG
        word(sa, "ZZ110셋" + TAG, "ZZC" + TAG, d, synonyms=[syn])
        sys_nm = psql("select sys_nm from quality.tb_sys_info order by sys_cd limit 1")
        sys_cd = psql(f"select sys_cd from quality.tb_sys_info where sys_nm={q(sys_nm)} order by sys_cd limit 1")
        ok_nm = "ZZ110업" + TAG
        path = xlsx(os.path.join(TMP, "words.xlsx"), WORD_HDR, [
            [1, "1", ok_nm, "ZZU" + TAG, "Zzu" + TAG, "110 업로드", "N", "", "", "", sys_nm, "N"],
            [2, "1", "ZZ110숫자" + TAG, "1ZZ" + TAG, "Zzn" + TAG, "110 업로드", "N", "", "", "", "", "N"],
            [3, "1", syn, "ZZS" + TAG, "Zzs" + TAG, "110 업로드", "N", "", "", "", "", "N"],
        ])
        since = time.strftime("%Y%m%d%H%M%S")
        r = upload(sa, "/api/std/uploadWords", path, d)
        assert r.status_code == 200, r.text[:200]
        b = wait_bulk("WORD", since)
        det = psql(f"select string_agg(target_nm||':'||detail_type, ',' order by seq) from quality.tb_change_history_detail where change_id='{b['changeId']}'")
        got = psql(f"select coalesce(req_sys_cd,'(null)') from quality.tb_word where dict_id='{d}' and word_nm={q(ok_nm)}")
        bad = psql(f"select count(*) from quality.tb_word where dict_id='{d}' and word_nm in ({q('ZZ110숫자' + TAG)}, {q(syn)})")
        print(f"  일괄등록 이력: {b['summary']} 경로={b['source']} 건수={b['cnt']} 상세={det}\n  요청시스템 '{sys_nm}' → {got} (기대 {sys_cd}) / 숫자로 시작한 약어·유사어 행 등록={bad}")
        assert b["source"] == "BULK_UPLOAD" and b["cnt"] == 1 and "(성공:1, 건너뜀:0, 실패:2)" in b["summary"]
        assert det == f"{ok_nm}:SUCCESS" and got == sys_cd and bad == "0"
    step("S5-1. 단어 일괄등록 — 등록 경로 BULK_UPLOAD·상세 건, 요청시스템 이름→코드, 약어 규칙·유사어는 화면 등록과 같게 실패", _words)

    def _code():
        w4nm, w4ab = "ZZ110업단어" + TAG, "ZZW" + TAG
        word(sa, w4nm, w4ab, d)
        cnm = "ZZ110업코드" + TAG
        path = xlsx(os.path.join(TMP, "codes.xlsx"), TERM_HDR, [
            [1, "1", cnm, "110 업로드 코드", w4ab, FX["dom"], "", "", "", "Z1U" + TAG, "", "", "", "N", "", ""]])
        time.sleep(1.1)   # since 는 초 단위 — 바로 앞 S5-1 단어 일괄등록 이력과 같은 초에 들어가지 않게
        since = time.strftime("%Y%m%d%H%M%S")
        r = upload(sa, "/api/std/uploadCodeInfoList", path, d)
        assert r.status_code == 200, r.text[:200]
        b = wait_bulk("CODE", since)
        time.sleep(2)
        n = psql(f"select string_agg(target_type, ',') from quality.tb_change_history where change_type='BULK_INSERT' and change_user_id='space' "
                 f"and change_dt >= to_timestamp('{since}', 'YYYYMMDDHH24MISS')")
        print(f"  코드 일괄등록 이력 대상={n} / {b['summary']}")
        assert n == "CODE" and b["cnt"] == 1
    step("S5-2. 코드 일괄등록 — 이력은 CODE 한 건 (예전에는 TERM·CODE 두 건)", _code)


# ---------------------------------------------------------------- S6
def s6(sa, su):
    d = FX["dict"]

    def aprv(item_tp, item_id, nm, stat, rsn="110 시험", extra=None):
        body = {"reqTp": item_tp, "reqItemId": item_id, "reqItemNm": nm, "aprvStat": stat, "aprvStatUpdtRsn": rsn}
        body.update(extra or {})
        return post(sa, "/api/std/putStdAprvStat", [body])

    def _domain():
        dn = "ZZ110반D" + TAG
        r = post(su, "/api/std/createDomain", {"domainNm": dn, "domainGrpNm": FX["grp"], "domainClsfNm": FX["clsf"], "domainDesc": "110 시험",
                                              "dataType": "VARCHAR", "dataLen": 30, "dictId": d})
        did = psql(f"select domain_id from quality.tb_domain where domain_nm={q(dn)} and dict_id='{d}'")
        assert did, r
        psql(f"insert into quality.tb_terms (terms_id, terms_nm, terms_eng_abrv_nm, terms_desc, domain_nm, aprv_yn, dict_id) "
             f"values ('Z1T{TAG}', {q('ZZ110참조' + TAG)}, {q('ZZREF' + TAG)}, '110 시험', {q(dn)}, 'Y', '{d}')")
        r = aprv("DOMAIN", did, dn, 3)
        left = psql(f"select count(*) from quality.tb_domain where domain_id='{did}'")
        stats = psql(f"select count(*) from quality.tb_aprv_stats where req_item_id='{did}'")
        print(f"  쓰는 용어가 있는 도메인 반려 = {code(r)} {msg(r)} / 남음={left} 처리 행={stats}")
        assert code(r) == 500 and msg(r).startswith("이 도메인을 쓰는 용어가 있어 반려할 수 없습니다: ZZ110참조") and left == "1" and stats == "0"
        psql(f"delete from quality.tb_terms where terms_id='Z1T{TAG}'")
        long = "가나다라마바사아자차" * 8
        r = aprv("DOMAIN", did, dn, 3, rsn=long)
        rsn = psql(f"select length(aprv_stat_updt_rsn)||'|'||req_user_id from quality.tb_aprv_stats where req_item_id='{did}'")
        h = psql(f"select change_type||':'||target_type||':'||change_source||':'||(prev_value like '%domainNm={dn}%') "
                 f"from quality.tb_change_history where target_id='{did}'")
        print(f"  80자 사유로 반려 = {code(r)} {msg(r)} / 처리 행 사유 길이|요청자 = {rsn} / 변경 이력 = {h}")
        assert code(r) == 200 and rsn == "50|jyjang" and h == "DELETE:DOMAIN:ADMIN_REJECT:true"
    step("S6-1. 도메인 반려 — 쓰는 용어가 있으면 이름을 알려 거부, 50자 넘는 사유는 잘라 저장, 반려도 변경 이력(관리자 반려)", _domain)

    def _cascade():
        wn, ab = "ZZ110반어" + TAG, "ZZR" + TAG
        r = post(su, "/api/std/createWord", {"wordNm": wn, "wordEngAbrvNm": ab, "wordEngNm": ab.title(), "wordDesc": "110", "wordClsfYn": "N", "dictId": d})
        wid = psql(f"select word_id from quality.tb_word where word_nm={q(wn)} and dict_id='{d}'")
        assert wid, r
        tn = wn + "용"
        r = post(su, "/api/std/createTerms", {"termsNm": tn, "termsEngAbrvNm": ab, "domainNm": FX["dom"], "dictId": d,
                                             "wordList": [{"wordId": wid, "wordNm": wn, "wordOrd": 0}]})
        tid = psql(f"select terms_id from quality.tb_terms where terms_nm={q(tn)} and dict_id='{d}'")
        assert tid, r
        psql(f"update quality.tb_terms set cret_dt='20260101000000' where terms_id='{tid}'")
        r = aprv("WORD", wid, wn, 3)
        row = psql(f"select coalesce(req_user_id,'(null)')||'|'||req_cret_dt from quality.tb_aprv_stats where req_item_id='{tid}'")
        mine = [x for x in post(su, "/api/std/getMyRequestList", {"from": None, "to": None, "statusList": None, "limit": 0}) if x["reqItemId"] == tid]
        h = psql(f"select string_agg(target_type||':'||change_source, ',' order by target_type) from quality.tb_change_history "
                 f"where target_id in ('{wid}', '{tid}') and change_type='DELETE'")
        print(f"  단어 반려 = {code(r)} {msg(r)}\n  연쇄 반려 용어 처리 행 요청자|요청일 = {row} / 요청자 화면 = {[(x['reqItemNm'], x['aprvStat']) for x in mine]} / 이력 = {h}")
        assert code(r) == 200 and row == "jyjang|20260101000000" and len(mine) == 1 and mine[0]["aprvStat"] == 3
        assert h == "TERM:ADMIN_REJECT,WORD:ADMIN_REJECT"
    step("S6-2. 단어 반려의 연쇄 반려 용어 — 처리 행 요청자·요청일은 용어의 등록자·등록 시각, 요청자 화면에 나옴, 둘 다 반려 이력", _cascade)

    def _approve():
        wn, ab = "ZZ110승인어" + TAG, "ZZP" + TAG
        post(su, "/api/std/createWord", {"wordNm": wn, "wordEngAbrvNm": ab, "wordEngNm": ab.title(), "wordDesc": "110", "wordClsfYn": "N", "dictId": d})
        wid = psql(f"select word_id from quality.tb_word where word_nm={q(wn)} and dict_id='{d}'")
        assert wid
        psql(f"update quality.tb_word set cret_dt='20260102000000' where word_id='{wid}'")
        r = aprv("WORD", wid, wn, 2, extra={"reqUserId": "xxx", "reqCretDt": "20000101000000"})
        row = psql(f"select req_user_id||'|'||req_cret_dt from quality.tb_aprv_stats where req_item_id='{wid}'")
        again = aprv("WORD", wid, wn, 3)
        print(f"  승인 = {code(r)} {msg(r)} / 처리 행 요청자|요청일 = {row} (화면 값 xxx·20000101 무시)\n  (이미 고쳐짐 확인 109 #8) 승인된 단어 반려 = {code(again)} {msg(again)}")
        assert code(r) == 200 and msg(r) == "OK" and row == "jyjang|20260102000000"
        assert code(again) == 400 and "승인 대기 항목이 아닙니다" in msg(again)
    step("S6-3. 승인 — 처리 행 요청자·요청일은 항목 기준(화면이 보낸 값 무시), 승인된 항목은 다시 처리 안 됨", _approve)


# ---------------------------------------------------------------- S7
def s7(sa, su):
    def _disable():
        r1 = post(sa, "/api/dict/disable", {"dictId": "DEFAULT"})
        r2 = post(sa, "/api/dict/disable", {"dictId": "ZZ110없는사전" + TAG})
        use = psql("select use_yn from quality.tb_std_dict where dict_id='DEFAULT'")
        print(f"  기본 사전 중지 = {code(r1)} {msg(r1)} / 없는 사전 = {code(r2)} {msg(r2)} / 기본 사전 사용 = {use}")
        assert code(r1) == 500 and msg(r1) == "기본 사전은 사용 중지할 수 없습니다." and use == "Y"
        assert code(r2) == 500 and msg(r2).startswith("없는 사전입니다")
    step("S7-1. 표준사전 사용 중지 — 기본 사전·없는 사전은 실패 응답 (0건 갱신을 성공으로 보내지 않음)", _disable)


# ---------------------------------------------------------------- S8 화면
def s8(sa, su):
    from selenium import webdriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from edgeopt import headless

    d = FX["dict"]
    opts = webdriver.EdgeOptions()
    opts.add_argument("--log-level=3")
    drv = webdriver.Edge(options=headless(opts))
    drv.set_window_size(1600, 1000)
    W = lambda t=15: WebDriverWait(drv, t)

    def js(s, *a):
        return drv.execute_script(s, *a)

    def nav(nav_id, group_id=None):
        """좌측 메뉴 — 접힌 그룹(dsGroup·mmGroup)이면 먼저 펼친다"""
        if group_id and not drv.find_elements(By.ID, nav_id):
            hdr = drv.find_element(By.CSS_SELECTOR, f"#{group_id} .v-list-group__header")
            js("arguments[0].click();", hdr)
            time.sleep(1)
        el = W().until(EC.presence_of_element_located((By.ID, nav_id)))
        js("arguments[0].scrollIntoView({block:'center'}); arguments[0].click();", el)
        time.sleep(2)

    def swal_title(t=8, has=""):
        """보이는 SweetAlert 제목 (has 를 주면 그 글자가 든 제목이 뜰 때까지 — 앞 창이 닫히는 중일 수 있다)"""
        def _t(x):
            for e in x.find_elements(By.CSS_SELECTOR, ".swal2-popup .swal2-title"):
                if e.is_displayed() and e.text and has in e.text:
                    return e.text
            return None
        return W(t).until(_t)

    def swal_close(cancel=False):
        sel = ".swal2-cancel" if cancel else ".swal2-confirm"
        for b in drv.find_elements(By.CSS_SELECTOR, sel):
            if b.is_displayed():
                b.click()
                time.sleep(1)
                return

    def tab_titles():
        # 탭 글자는 "제목\nclose"(닫기 아이콘) — 제목만
        return [e.text.strip().split("\n")[0] for e in drv.find_elements(By.CSS_SELECTOR, ".v-tab") if e.text.strip()]

    try:
        drv.get(BASE)
        W().until(EC.presence_of_element_located((By.CSS_SELECTOR, "input[type='text']"))).send_keys("space")
        drv.find_element(By.CSS_SELECTOR, "input[type='password']").send_keys(PW)
        drv.find_element(By.CSS_SELECTOR, "button[type='submit'], .v-btn").click()
        W().until(lambda x: "/app/main" in x.current_url)
        time.sleep(3)
        js("window.sessionStorage.setItem('navidmeta.dictId', arguments[0]);", d)

        def _search():
            box = W().until(EC.presence_of_element_located((By.CSS_SELECTOR, "input[placeholder='통합 검색']")))
            box.send_keys(FX["w1nm"], Keys.ENTER)
            W().until(lambda x: "단어 (1건)" in x.find_element(By.ID, "tab_globalSearch").text)
            body = drv.find_element(By.ID, "tab_globalSearch").text
            dict_nm = "ZZ110사전" + TAG
            print(f"  통합 검색 탭: 사전 선택기 표시={'표준사전' in body and dict_nm in body} 단어 (1건)")
            assert "표준사전" in body and dict_nm in body
            # 컬럼 결과의 모델명 열
            box = drv.find_element(By.CSS_SELECTOR, "input[placeholder='통합 검색']")
            box.send_keys(Keys.CONTROL, "a")
            box.send_keys("ZZ110_T" + TAG, Keys.ENTER)
            W().until(lambda x: "테이블/컬럼 (2건)" in x.find_element(By.ID, "tab_globalSearch").text)
            ths = [e.text for e in drv.find_elements(By.CSS_SELECTOR, "#tab_globalSearch th")]
            cells = drv.find_element(By.ID, "tab_globalSearch").text
            print(f"  컬럼 결과 머리글={ths}")
            assert "데이터 모델명" in ths and FX["dmNm"] in cells and "시험테이블" in cells
            # 단어 결과에서 단어 화면으로 — 찾은 사전으로 열린다
            js("window.sessionStorage.setItem('navidmeta.dictId', 'DEFAULT');")
            box.send_keys(Keys.CONTROL, "a")
            box.send_keys(FX["w1nm"], Keys.ENTER)
            W().until(lambda x: "단어 (1건)" in x.find_element(By.ID, "tab_globalSearch").text)
            row = drv.find_element(By.XPATH, f"//div[@id='tab_globalSearch']//td[normalize-space()='{FX['w1nm']}']")
            row.click()
            W().until(lambda x: FX["w1nm"] in x.find_element(By.ID, "tab_word").text)
            wt = drv.find_element(By.ID, "tab_word").text
            print(f"  단어 화면: 시험 사전으로 열림={dict_nm in wt} 단어 보임={FX['w1nm'] in wt}")
            assert dict_nm in wt
        step("S8-1. 화면 통합 검색 — 사전 선택기, 컬럼의 데이터 모델명·테이블 한글명, 결과 → 찾은 사전으로 단어 화면", _search)

        def _upload():
            nm = "ZZ110화면업" + TAG
            path = xlsx(os.path.join(TMP, "ui_words.xlsx"), WORD_HDR, [[1, "1", nm, "ZZV" + TAG, "Zzv" + TAG, "110 화면 업로드", "N", "", "", "", "", "N"]])
            js("window.sessionStorage.setItem('navidmeta.dictId', arguments[0]);", d)
            inp = drv.find_element(By.ID, "inputUpload")
            js("arguments[0].style.display='block';", inp)   # 숨은 파일 칸에 파일을 넣으려면 보이게
            inp.send_keys(path)
            t = swal_title(60, "일괄등록 완료")
            time.sleep(4)   # 같은 요약이 두 번 오는지 볼 시간
            # 진행 창의 로그 한 줄 = 시각 span 하나를 가진 div
            logs = [e.text for e in drv.find_elements(By.XPATH, "//div[contains(@class,'v-dialog--active')]//div[span][not(.//div)]")
                    if "완료 -" in e.text and "[단어]" in e.text]
            got = psql(f"select count(*) from quality.tb_word where dict_id='{d}' and word_nm={q(nm)}")
            print(f"  업로드 결과 창 '{t}' / 진행 창의 '완료 -' 줄 {len(logs)}개 / 등록={got}")
            assert t == "단어 일괄등록 완료" and len(logs) == 1 and got == "1"
            swal_close()
            for b in drv.find_elements(By.XPATH, "//div[contains(@class,'v-dialog--active')]//button[normalize-space()='닫기']"):
                if b.is_displayed():
                    b.click()
            time.sleep(1)
        step("S8-2. 화면 단어 일괄등록 — 완료 요약이 요청한 창에 한 번만 (개인 큐)", _upload)

        def _approval():
            wn, ab = "ZZ110화면반" + TAG, "ZZX" + TAG
            post(su, "/api/std/createWord", {"wordNm": wn, "wordEngAbrvNm": ab, "wordEngNm": ab.title(), "wordDesc": "110", "wordClsfYn": "N", "dictId": d})
            wid = psql(f"select word_id from quality.tb_word where word_nm={q(wn)} and dict_id='{d}'")
            post(su, "/api/std/createTerms", {"termsNm": wn + "용", "termsEngAbrvNm": ab, "domainNm": FX["dom"], "dictId": d,
                                             "wordList": [{"wordId": wid, "wordNm": wn, "wordOrd": 0}]})
            dn = "ZZ110화면D" + TAG
            post(su, "/api/std/createDomain", {"domainNm": dn, "domainGrpNm": FX["grp"], "domainClsfNm": FX["clsf"], "domainDesc": "110",
                                              "dataType": "VARCHAR", "dataLen": 30, "dictId": d})
            nav("nav_approval", "mmGroup")
            tab = W().until(EC.presence_of_element_located((By.ID, "tab_approval")))
            W().until(lambda x: wn in tab.text and dn in tab.text)

            def row_btn(nm, label):
                return drv.find_element(By.XPATH, f"//div[@id='tab_approval']//tr[.//span[normalize-space()='{nm}']]//button[@title='{label}']")
            js("arguments[0].click();", row_btn(wn, "승인"))
            t1 = swal_title(has="승인하시겠습니까")
            swal_close(cancel=True)
            js("arguments[0].click();", row_btn(wn, "반려"))
            box = W().until(EC.presence_of_element_located((By.XPATH, "//div[@id='tab_approval']//input[@placeholder='사유']")))
            box.send_keys("가" * 60)
            vlen = len(box.get_attribute("value"))
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[@id='tab_approval']//button[@title='반려 확정']"))
            t2 = swal_title(has="연관")
            body2 = drv.find_element(By.CSS_SELECTOR, ".swal2-popup").text
            swal_close(cancel=True)
            still = psql(f"select count(*) from quality.tb_word where word_id='{wid}'")
            print(f"  행 [승인] → '{t1}' (취소) / 행 [반려] 사유 입력 길이={vlen} → '{t2}' 용어 안내={wn + '용' in body2} (취소, 단어 남음={still})")
            assert "승인하시겠습니까" in t1 and vlen == 50 and t2 == "연관 미승인 용어도 함께 삭제됩니다" and wn + "용" in body2 and still == "1"
            js("arguments[0].click();", row_btn(dn, "승인"))
            swal_title(has="승인하시겠습니까")
            swal_close()
            t3 = swal_title(has="완료")
            body3 = drv.find_element(By.CSS_SELECTOR, ".swal2-popup").text
            ok = psql(f"select aprv_yn from quality.tb_domain where domain_nm={q(dn)} and dict_id='{d}'")
            print(f"  행 [승인] 확인 → '{t3}' 본문에 OK 없음={'OK' not in body3} / 도메인 승인={ok}")
            assert t3 == "승인이 완료되었습니다." and "OK" not in body3 and ok == "Y"
            time.sleep(2)
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[@id='tab_approval']//button[starts-with(normalize-space(),'승인완료')]"))
            time.sleep(1)
            done_ok = row_btn(dn, "승인").get_attribute("disabled")
            done_rj = row_btn(dn, "반려").get_attribute("disabled")
            cb = drv.find_elements(By.XPATH, f"//div[@id='tab_approval']//tr[.//span[normalize-space()='{dn}']]//td[1]//*[contains(@class,'v-simple-checkbox')]")
            cb_off = bool(cb) and "v-simple-checkbox--disabled" in (cb[0].get_attribute("class") or "")
            print(f"  승인완료 행 [승인] disabled={done_ok} [반려] disabled={done_rj} 선택 칸 막힘={cb_off}")
            assert done_ok and done_rj and cb_off
        step("S8-3. 화면 표준 사전 변경 승인 — 행 승인 확인창, 행 반려 연관 용어 확인창·사유 50자, 성공 'OK' 경고 없음, 처리된 행 버튼·선택 막힘", _approval)

        def _history():
            psql("insert into quality.tb_change_history (change_id, change_type, target_type, target_nm, summary, change_user_id, change_dt, change_source) "
                 f"select 'Z1H{TAG}' || g, 'UPDATE', 'WORD', 'ZZ110이력{TAG}', '110 시험', 'space', now(), 'ADMIN_DIRECT' from generate_series(1, 200) g")
            nav("nav_changeHistory", "dsGroup")
            tab = W().until(EC.presence_of_element_located((By.ID, "tab_changeHistory")))
            time.sleep(2)
            has_src = bool(drv.find_elements(By.ID, "chgSourceFilter"))
            note = drv.find_elements(By.ID, "chgLimitNote")
            note_ok = bool(note) and note[0].is_displayed()
            from selenium.webdriver.common.action_chains import ActionChains
            lbl = drv.find_element(By.XPATH, "//div[@id='tab_changeHistory']//label[contains(normalize-space(.),'대상')]")
            slot = lbl.find_element(By.XPATH, "./ancestor::div[contains(@class,'v-input__slot')][1]")
            ActionChains(drv).move_to_element(slot).click().perform()
            time.sleep(1)
            items = [e.text for e in drv.find_elements(By.CSS_SELECTOR, ".v-menu__content.menuable__content__active .v-list-item__title")]
            print(f"  등록 경로 필터={has_src} / 200건 안내={note_ok} / 대상 항목={items}")
            assert has_src and note_ok and "도메인 그룹" in items and "도메인 분류" in items
            drv.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
        step("S8-4. 화면 변경 이력 — 등록 경로 필터, 대상에 도메인 그룹·분류, 200건이 차면 안내", _history)

        def _dict():
            nav("nav_std_dict", "mmGroup")
            tab = W().until(EC.presence_of_element_located((By.ID, "tab_std_dict")))
            row = drv.find_element(By.XPATH, "//div[@id='tab_std_dict']//tr[.//td[normalize-space()='기본 표준사전']]//td[1]//*[contains(@class,'v-simple-checkbox')]")
            js("arguments[0].click();", row)
            time.sleep(0.5)
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[@id='tab_std_dict']//button[normalize-space()='수정']"))
            sw = W().until(EC.presence_of_element_located((By.ID, "dictDefaultYn")))
            dis = sw.get_attribute("disabled")
            dlg = sw.find_element(By.XPATH, "./ancestor::div[contains(@class,'v-dialog')][1]")
            txt = dlg.get_attribute("textContent")
            print(f"  기본 사전 수정 창 — '기본 사전으로 지정' 스위치 disabled={dis} 안내={'다른 사전을 수정해 기본 사전으로 지정' in txt}")
            assert dis and "다른 사전을 수정해 기본 사전으로 지정" in txt
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[contains(@class,'v-dialog--active')]//button[normalize-space()='취소']"))
        step("S8-5. 화면 표준사전 관리 — 기본 사전 수정 창에서 '기본 사전으로 지정' 을 끌 수 없다", _dict)

        def _dash():
            nav("nav_dashboard")
            card = W().until(EC.presence_of_element_located((By.XPATH, "//div[@id='tab_dashboard']//div[contains(@class,'std-card__label') and normalize-space()='모델']/ancestor::div[contains(@class,'std-card')][1]")))
            js("arguments[0].click();", card)
            time.sleep(2)
            print(f"  탭 제목 = {tab_titles()}")
            assert "데이터 모델 현황" in tab_titles() and "현황" not in tab_titles()
        step("S8-6. 화면 대시보드 — '모델' 카드는 메뉴와 같은 '데이터 모델 현황' 탭", _dash)

        def _code_upload():
            grp = "Z1V" + TAG
            r = post(sa, "/api/std/createCode", {"termsNm": "ZZ110화면코드" + TAG, "termsEngAbrvNm": FX["w1ab"], "termsDesc": "110", "domainNm": FX["dom"],
                                                "codeGrp": grp, "codeTp": "SINGLE", "integSubGrps": [], "dictId": d,
                                                "wordList": [{"wordId": FX["w1"], "wordNm": FX["w1nm"], "wordOrd": 0}]})
            assert code(r) == 200, r
            path = xlsx(os.path.join(TMP, "ui_codedata.xlsx"), ["No", "코드그룹", "코드명", "코드영문명", "코드값", "코드값설명", "상위코드값"],
                        [[1, grp, "ZZ110화면값" + TAG, "VAL", "A1", "110 화면", ""]])
            js("window.sessionStorage.setItem('navidmeta.dictId', arguments[0]);", d)
            nav("nav_dsCode", "dsGroup")
            W().until(EC.presence_of_element_located((By.ID, "tab_dsCode")))
            time.sleep(1)
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[@id='tab_dsCode']//button[@title='코드 항목 관리']"))
            inp = W().until(EC.presence_of_element_located((By.ID, "inputCodeValUpload")))
            n_inputs = len(drv.find_elements(By.ID, "inputCodeValUpload"))
            js("arguments[0].style.display='block';", inp)
            inp.send_keys(path)
            t = swal_title(60, "일괄등록 완료")
            got = psql(f"select count(*) from quality.tb_code_data where dict_id='{d}' and code_grp='{grp}' and code_val='A1'")
            print(f"  코드 항목 관리 창 파일 칸 {n_inputs}개 / 결과 창 '{t}' / 등록={got}")
            assert n_inputs == 1 and t == "코드데이터 일괄등록 완료" and got == "1"
            time.sleep(3)
            drv.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
        step("S8-7. 화면 코드 항목(값) 일괄등록 — 파일 칸 id 하나, 바꾼 파일 칸의 파일로 올라감", _code_upload)

        def _clsf_items():
            nav("nav_word", "dsGroup")
            W().until(EC.presence_of_element_located((By.ID, "addWordBtn")))
            for _ in range(3):
                js("document.getElementById('addWordBtn').click();")
                time.sleep(1.5)
            # 화면 상태 확인용 — 단어 화면의 도메인 분류 후보 배열 (값을 바꾸지 않는다)
            n, uniq = js("let v = document.getElementById('addWordBtn').__vue__; while (v && v.$options.name !== 'DSWord') v = v.$parent;"
                         "return [v.domainClassificationItems.length, new Set(v.domainClassificationItems).size];")
            want = int(psql(f"select count(*) from quality.tb_domain_clsf where dict_id='{d}'"))
            print(f"  등록 창을 세 번 열어도 도메인 분류 후보 {n}개 (중복 없는 {uniq}개, 사전의 분류 {want}개)")
            assert n == uniq == want
            drv.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
        step("S8-8. 화면 단어 등록 창 — 열 때마다 도메인 분류 후보가 덧붙지 않는다", _clsf_items)

        def _upload_fail():
            bad = os.path.join(TMP, "broken_words.xlsx")
            with open(bad, "w", encoding="utf-8") as f:
                f.write("엑셀이 아닌 파일")
            drv.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
            time.sleep(1)
            inp = drv.find_element(By.ID, "inputUpload")
            js("arguments[0].style.display='block';", inp)
            t0 = time.time()
            inp.send_keys(bad)

            def _closed(x):
                for b in x.find_elements(By.XPATH, "//div[contains(@class,'v-dialog--active')]//button[normalize-space()='닫기']"):
                    if b.is_displayed():
                        return True
                return False
            W(30).until(_closed)
            took = time.time() - t0
            logs = [e.text for e in drv.find_elements(By.XPATH, "//div[contains(@class,'v-dialog--active')]//div[span][not(.//div)]")]
            fails = [x for x in logs if "[단어] 일괄등록 실패" in x]
            print(f"  엑셀이 아닌 파일 → {took:.1f}초 만에 진행 끝([닫기]) / 실패 줄 {len(fails)}개 / 시작 줄 오류 표시 없음={not any('ERROR' in x and '시작' in x for x in logs)}")
            assert took < 20 and len(fails) == 1
            for b in drv.find_elements(By.XPATH, "//div[contains(@class,'v-dialog--active')]//button[normalize-space()='닫기']"):
                if b.is_displayed():
                    b.click()
        step("S8-9. 화면 단어 일괄등록 실패 — '[단어] 일괄등록 실패: …' 한 줄로 바로 끝난다 (60초 대기 아님)", _upload_fail)
    finally:
        drv.quit()


def main(only):
    sa = login_api("space")
    su = login_api("jyjang")
    start = time.strftime("%Y%m%d%H%M%S")
    assert psql("select count(*) from quality.tb_aprv_line where use_yn='Y'") == "0", "켜진 결재선이 있다 — 시험이 표준 승인을 하므로 멈춘다"
    groups = [("S1", s1), ("S2", s2), ("S3", s3), ("S4", s4), ("S5", s5), ("S6", s6), ("S7", s7), ("S8", s8)]
    try:
        fixture(sa)
        for nm, fn in groups:
            if not only or nm in only:
                fn(sa, su)
    finally:
        print("\n=== 정리")
        for fn in reversed(cleanup):
            try:
                fn()
            except Exception as e:
                print("  정리 실패:", e)
        since = f"to_timestamp('{start}', 'YYYYMMDDHH24MISS')"
        psql(f"delete from quality.tb_notification where cret_dt >= '{start}'")
        psql(f"delete from quality.tb_change_history_detail where change_id in (select change_id from quality.tb_change_history "
             f"where change_dt >= {since} and change_type='BULK_INSERT' and change_user_id='space')")
        psql(f"delete from quality.tb_change_history where change_dt >= {since} and (target_nm like 'ZZ110%' or summary like '%ZZ110%' "
             f"or target_nm like 'Z1K%' or (change_type='BULK_INSERT' and change_user_id='space'))")
        left = psql(f"select (select count(*) from quality.tb_std_dict where dict_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_word where word_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_terms where terms_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_domain where domain_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_data_model where dm_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_aprv_stats where req_item_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_change_history where change_dt >= {since} and (target_nm like 'ZZ110%' or summary like '%ZZ110%'))||'/'||"
                    f"(select count(*) from quality.tb_struct_diag_history where diag_id like 'Z1S%')")
        print(f"  남은 시험 데이터 사전/단어/용어/도메인/모델/처리 행/변경 이력/구조 진단 = {left}")


if __name__ == "__main__":
    t0 = time.time()
    main(set(sys.argv[1:]))
    elapsed = time.time() - t0
    p = sum(1 for _, st in results if st == "PASS")
    f = sum(1 for _, st in results if st == "FAIL")
    print(f"\n{'=' * 60}\n결과: {p} PASS / {f} FAIL  ({elapsed:.0f}초)\n{'=' * 60}")
    for n, st in results:
        print(f"  [{st}] {n}")
    sys.exit(0 if f == 0 else 1)
