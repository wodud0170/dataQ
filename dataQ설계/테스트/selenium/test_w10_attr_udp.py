"""
104번 WP8-75 · 108 §4.23-4 (§4.27) — 컬럼 대상 사용자 정의 속성(UDP)

시험 모델 W10_UDP_<tag>(테이블 APP.W10_T, 컬럼 C_ID·C_NM)·숫자 속성(컬럼 대상, 필수 아님)·시험 일반 사용자는 API 로 만들고 끝에서 모두 지운다.
필수 선택 속성 '셀등급<tag>'(컬럼 대상)는 C1 에서 화면으로 만든다.

  C1. 관리 > 사용자 정의 속성 — 위 안내문 새 문장, [속성 추가] 대상 '컬럼'으로 바꾸면 순서 기본값(컬럼 속성 수 + 1)·필수 스위치·설명 안내 문구가 바뀜,
      기본 열 이름('디폴트값')은 저장 거부, 필수 선택 속성 저장 → 목록 '대상' 열 '컬럼', [수정] 창에서 대상 칸 비활성·안내
  C2. 데이터 모델 > 컬럼 — '디폴트' 뒤 속성 열(머리 '이름 *'·주황 머리·마우스를 올리면 설명), 새 행 필수 칸 빨강,
      비우고 저장 → '일부 저장 실패' "W10_T.W10_NEW — 필수 관리 항목을 넣어 주세요: …"(컬럼 안 생김), 값 고르고 저장 → DB 값
  C3. 기존 행 속성만 바꿔 저장(선택·숫자 칸 노랑, 수정중 1건) → 값 바뀜·변경 이력 그대로,
      영문명 바꾸고(속성은 그대로) 필수 속성이 빈 다른 행의 한글명만 바꿔 저장 → 값이 새 이름으로·한글명 저장(속성 안 보낸 수정은 필수 검사 안 함)
  C4. 엑셀 — 컬럼 양식 머리 끝에 컬럼 속성 이름(테이블 양식에는 없음), 업로드 미리보기 등록 행만 검사(필수 빠짐·숫자 아님 → 오류 행),
      데이터 다운로드 끝 열에 속성 값
  C5. 시험 일반 사용자 — 사용자 정의 속성 메뉴 없음, 그리드 속성 칸 편집·저장(값 저장), 테이블 추가 창에 컬럼 속성 칸 없음(테이블 속성 칸은 있음)
  C6. 선택값을 고쳐 저장된 값('중')이 선택값에서 빠지면 — 그리드 칸에 그 값이 보이고(목록 끝에도), 고치지 않은 행은 빨강 없음,
      같은 행 다른 속성만 고치면 그 칸 빨강·저장 거부 문구, 다른 값을 고르면 빨강이 풀리고 저장
"""
import base64, io, json, secrets, subprocess, sys, time, traceback
from edgeopt import headless
from dbcmd import PSQL
import requests
from openpyxl import Workbook, load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

BASE = "http://localhost:28091"
TAG = time.strftime("%H%M%S")
DM_NM = f"W10_UDP_{TAG}"
UDP_NM = f"셀등급{TAG}"          # 화면에서 만드는 필수 선택 속성 (컬럼 대상)
UDP_DESC = "셀 시험 등급 설명"
NUM_NM = f"셀수량{TAG}"          # API 로 만드는 숫자 속성 (컬럼 대상, 필수 아님)
USER = "w10udp" + TAG
ATTR_HEADERS = ["소유자", "테이블명(영문)", "테이블명(한글)", "컬럼명(영문)", "컬럼명(한글)", "데이터타입", "길이", "소수점자리", "컬럼 순서",
                "NULL여부", "PK여부", "FK여부", "디폴트값", "참조 테이블(한글)", "참조 컬럼(한글)", "삭제 규칙"]
results = []
FX = {}

# 탭 안 Vue 컴포넌트 — 자식부터 찾고 없으면 위로 (첫 요소가 keep-alive 로 잡히는 화면이 있다)
VM = ("const vm = (tab, name) => { const e = [...document.querySelectorAll('#' + tab + ' *')].find(x => x.__vue__); "
      "const down = v => { if (v.$options.name === name) return v; for (const ch of v.$children) { const r = down(ch); if (r) return r; } return null; }; "
      "let c = e && e.__vue__; while (c) { const d = down(c); if (d) return d; c = c.$parent; } return null; };")
CDEF = VM + "const c = vm('tab_udpDef', 'MMUdpDef');"
CCOL = VM + "const c = vm('tab_datamodelStatusColumn', 'DSDatamodelStatusColumn');"
CTBL = VM + "const c = vm('tab_datamodelStatusTable', 'DSDatamodelStatusTable');"
# 컬럼 그리드 칸 — 영문명 입력값이 key 인 행의, 머리 글자가 hdr 인 열
CELL = ("const cell = (key, hdr) => { const ths = [...document.querySelectorAll('#clTable_table thead tr:first-child th')]; "
        "const i = ths.findIndex(t => t.textContent.trim() === hdr); "
        "const tr = [...document.querySelectorAll('#clTable_table tbody tr')].find(r => [...r.querySelectorAll('input')].some(x => x.value === key)); "
        "return tr && i >= 0 ? tr.children[i] : null; };")


def step(name, fn):
    print(f"\n=== {name}")
    try:
        fn()
        results.append((name, "PASS"))
        print("  >> PASS")
    except Exception:
        traceback.print_exc()
        results.append((name, "FAIL"))


def psql(sql):
    r = subprocess.run(PSQL + ["-v", "ON_ERROR_STOP=1", "-At", "-c", sql],
                       capture_output=True, text=True, timeout=60, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        raise RuntimeError("psql: " + (r.stderr or "")[:400])
    return (r.stdout or "").strip()


def q(v):
    return "NULL" if v is None else "'" + str(v).replace("'", "''") + "'"


def b64(s):
    return base64.b64encode(s.encode("utf-8")).decode()


def api_login(uid, pw):
    s = requests.Session()
    r = s.post(BASE + "/login", data={"id": uid, "password": b64(pw)}, allow_redirects=False, timeout=10)
    assert r.status_code == 200, f"로그인 실패 {uid}: {r.status_code}"
    return s


def ok(r):
    j = r.json()
    assert j.get("resultCode") == 200, j
    return j


def setup():
    sa = FX["sa"] = api_login("space", "123")
    FX["upw"] = "Zz" + secrets.token_hex(6)
    ok(sa.post(BASE + "/api/login/createUser", json={"id": USER, "password": b64(FX["upw"]), "name": "W10 UDP 시험 " + USER,
                                                    "email": "", "admin": False}, timeout=10))
    sa.post(BASE + "/api/dm/createDataModel", json={"dataModelNm": DM_NM, "modelType": "LOGICAL", "ver": "1", "dictId": "DEFAULT"}, timeout=10)
    FX["dm"] = dm = psql(f"select dm_id from quality.tb_data_model where dm_nm={q(DM_NM)}")
    assert dm, "시험 모델이 안 만들어짐"
    ok(sa.post(BASE + "/api/dm/addObj", json={"dataModelId": dm, "objOwner": "APP", "objNm": "W10_T", "objNmKr": "시험테이블"}, timeout=10))
    ok(sa.post(BASE + "/api/dm/saveAttrs", json={"dataModelId": dm, "objNm": "W10_T", "objOwner": "APP", "attrs": [
        {"mode": "ADD", "attrNm": "C_ID", "attrNmKr": "셀아이디", "dataType": "VARCHAR", "dataLen": 10},
        {"mode": "ADD", "attrNm": "C_NM", "attrNmKr": "셀이름", "dataType": "VARCHAR", "dataLen": 50}]}, timeout=10))
    r = ok(sa.post(BASE + "/api/udp/save", json={"targetTp": "ATTR", "udpNm": NUM_NM, "dataTp": "NUMBER", "requiredYn": "N",
                                                 "sortOrder": 1, "useYn": "Y", "udpDesc": ""}, timeout=10))
    FX["num"] = r["contents"]


def cleanup():
    dm = FX.get("dm")
    if dm:
        for t in ("tb_udp_val", "tb_data_model_change_history", "tb_data_model_attr", "tb_data_model_index", "tb_data_model_constraint",
                  "tb_data_model_obj", "tb_data_model_clct", "tb_data_model_schema"):
            psql(f"delete from quality.{t} where dm_id={q(dm)}")
        psql(f"delete from quality.tb_data_model where dm_id={q(dm)}")
    psql(f"delete from quality.tb_udp_val where udp_id in (select udp_id from quality.tb_udp_def where udp_nm in ({q(UDP_NM)}, {q(NUM_NM)}))")
    psql(f"delete from quality.tb_udp_def where udp_nm in ({q(UDP_NM)}, {q(NUM_NM)})")
    psql(f"delete from quality.tb_user_role where user_id={q(USER)}")
    psql(f"delete from quality.tb_user where user_id={q(USER)}")
    psql(f"delete from quality.tb_event_log where job_id='authentication' and job_ex_user_id={q(USER)}")


def vals(attr):
    """컬럼 하나의 속성 값 {속성 이름: 값}"""
    rows = psql(f"select d.udp_nm || '=' || v.val from quality.tb_udp_val v join quality.tb_udp_def d on d.udp_id=v.udp_id "
                f"where v.dm_id={q(FX['dm'])} and v.obj_owner='APP' and v.obj_nm='W10_T' and v.attr_nm={q(attr)} order by 1")
    return dict(r.split("=", 1) for r in rows.splitlines() if r)


def hist_cnt():
    return int(psql(f"select count(*) from quality.tb_data_model_change_history where dm_id={q(FX['dm'])}"))


def attr_defs():
    """사용 중인 컬럼 대상 속성 (그리드·양식 순서)"""
    return FX["sa"].post(BASE + "/api/udp/defs", json={"useOnly": "Y", "targetTp": "ATTR"}, timeout=10).json()


def new_driver(user, pw):
    opts = webdriver.EdgeOptions()
    opts.add_argument("--log-level=3")
    d = webdriver.Edge(options=headless(opts))
    d.set_window_size(1600, 900)
    d.get(BASE)
    time.sleep(2)
    WebDriverWait(d, 10).until(EC.presence_of_element_located((By.CSS_SELECTOR, "input[type='text']"))).send_keys(user)
    d.find_element(By.CSS_SELECTOR, "input[type='password']").send_keys(pw)
    d.find_element(By.CSS_SELECTOR, ".login-btn, button.v-btn").click()
    WebDriverWait(d, 15).until(lambda x: "/signin" not in x.current_url)
    time.sleep(3)
    return d


class Ui:
    """한 브라우저의 화면 조작 도우미"""

    def __init__(self, d):
        self.d = d

    def js(self, script, *args):
        return self.d.execute_script(script, *args)

    def open_menu(self, group_id, nav_id):
        g = self.d.find_elements(By.CSS_SELECTOR, f"#{group_id} .v-list-group__header")
        if g and "v-list-group__header--active" not in (g[0].get_attribute("class") or ""):
            self.js("arguments[0].click();", g[0])
            time.sleep(1)
        self.js(f"document.getElementById('{nav_id}').click();")
        time.sleep(2.5)

    def swal(self):
        t = self.d.find_elements(By.CSS_SELECTOR, ".swal2-title")
        h = self.d.find_elements(By.CSS_SELECTOR, ".swal2-html-container")
        return (t[0].get_attribute("textContent") if t else "") + " | " + (h[0].get_attribute("textContent") if h else "")

    def wait_swal(self, text, timeout=10):
        WebDriverWait(self.d, timeout).until(lambda x: text in self.swal())
        return self.swal()

    def swal_ok(self):
        if self.d.find_elements(By.CSS_SELECTOR, ".swal2-confirm"):
            self.js("document.querySelector('.swal2-confirm').click();")
            time.sleep(0.8)

    def pick(self, inp, text):
        """v-select 를 열고 목록에서 text 를 고른다 (inp: 그 v-select 안의 요소 또는 그것을 담은 그리드 칸)"""
        self.js("const e = arguments[0]; (e.closest('.v-input') || e.querySelector('.v-input')).querySelector('.v-input__slot').click();", inp)
        WebDriverWait(self.d, 5).until(lambda x: x.execute_script(
            "return [...document.querySelectorAll('.menuable__content__active .v-list-item')].some(e => e.textContent.trim() === arguments[0]);", text))
        self.js("[...document.querySelectorAll('.menuable__content__active .v-list-item')].find(e => e.textContent.trim() === arguments[0]).click();", text)
        time.sleep(0.5)

    def open_columns(self, dm):
        self.open_menu("dmGroup", "nav_datamodelStatusColumn")
        self.js(CCOL + "c.selectedModelId = arguments[0]; c.onModelChange(arguments[0]);", dm)
        WebDriverWait(self.d, 15).until(lambda x: x.execute_script(CCOL + "return !c.loadTable && c.dmColumnAllItems.length > 0;"))
        time.sleep(1)

    def heads(self):
        return self.js("return [...document.querySelectorAll('#clTable_table thead tr:first-child th')].map(t => t.textContent.trim());")

    def cell(self, key, hdr):
        c = self.js(CELL + "return cell(arguments[0], arguments[1]);", key, hdr)
        assert c is not None, f"칸이 없음: {key} / {hdr}"
        return c

    def cell_cls(self, key, hdr):
        return self.js("return arguments[0].querySelector('.v-input').className;", self.cell(key, hdr))

    def save(self):
        b = self.d.find_element(By.ID, "btn-save-attrs")
        assert not b.get_attribute("disabled"), "저장 버튼 비활성"
        self.js("arguments[0].click();", b)


def main():
    setup()
    sa = FX["sa"]
    d = new_driver("space", "123")
    ui = Ui(d)
    try:
        def _c1():
            ui.open_menu("mmGroup", "nav_udpDef")
            guide = d.find_element(By.ID, "udpGuide").get_attribute("textContent")
            assert "대상이 컬럼인 항목은 데이터 모델 > 컬럼 목록의 열과 컬럼 엑셀 양식에 들어갑니다." in " ".join(guide.split()), guide
            heads = [h.get_attribute("textContent").strip() for h in d.find_elements(By.CSS_SELECTOR, "#udpTable thead th")]
            assert heads[1:4] == ["속성 이름", "대상", "유형"], heads
            n_obj = int(psql("select count(*) from quality.tb_udp_def where target_tp='OBJ' and coalesce(use_yn,'Y')='Y'"))
            n_attr = int(psql("select count(*) from quality.tb_udp_def where target_tp='ATTR' and coalesce(use_yn,'Y')='Y'"))

            def open_add():
                ui.js("document.getElementById('udpAdd').click();")
                time.sleep(1)
                assert d.find_element(By.ID, "udpSortOrder").get_attribute("value") == str(n_obj + 1), "테이블 순서 기본값"
                tgt = d.find_element(By.ID, "udpTargetTp")
                assert not tgt.get_attribute("disabled"), "추가 창에서 대상이 비활성"
                ui.pick(tgt, "컬럼")
                assert d.find_element(By.ID, "udpSortOrder").get_attribute("value") == str(n_attr + 1), "컬럼 순서 기본값"
                req = ui.js("return document.getElementById('udpRequired').closest('.v-input').querySelector('label').textContent.trim();")
                assert req == "필수 — 비우면 컬럼을 저장하지 않음", req
                assert d.find_element(By.ID, "udpDescGuide").text.strip() == "컬럼 목록에서 이 속성 열 머리에 마우스를 올리면 보입니다."
                # 테이블로 되돌리면 문구·순서도 되돌아간다
                ui.pick(tgt, "테이블")
                assert d.find_element(By.ID, "udpSortOrder").get_attribute("value") == str(n_obj + 1)
                assert ui.js("return document.getElementById('udpRequired').closest('.v-input').querySelector('label').textContent.trim();") \
                    == "필수 — 비우면 테이블을 저장하지 않음"
                ui.pick(tgt, "컬럼")

            # 기본 열 이름과 같은 이름은 거부
            open_add()
            d.find_element(By.ID, "udpNm").send_keys("디폴트값")
            ui.js("document.getElementById('udpSave').click();")
            msg = ui.wait_swal("저장 실패")
            print("  기본 열 이름:", msg)
            assert "엑셀 양식의 기본 열 이름과 같은 이름은 쓸 수 없습니다: 디폴트값" in msg
            ui.swal_ok()
            ui.js("[...document.querySelectorAll('.v-dialog--active button')].find(b => b.textContent.trim() === '취소').click();")
            time.sleep(0.8)

            open_add()
            d.find_element(By.ID, "udpNm").send_keys(UDP_NM)
            ui.pick(d.find_element(By.ID, "udpDataTp"), "선택")
            d.find_element(By.ID, "udpSelectVals").send_keys("상\n중\n하")
            ui.js("document.getElementById('udpRequired').closest('.v-input').querySelector('.v-input--selection-controls__ripple').click();")
            d.find_element(By.ID, "udpDesc").send_keys(UDP_DESC)
            time.sleep(0.3)
            ui.js("document.getElementById('udpSave').click();")
            time.sleep(2)
            row = psql(f"select udp_id || '|' || target_tp || '|' || data_tp || '|' || required_yn || '|' || sort_order || '|' || array_to_string(select_vals, ',') "
                       f"from quality.tb_udp_def where udp_nm={q(UDP_NM)}")
            print("  정의:", row)
            assert row.split("|", 1)[1] == f"ATTR|SELECT|Y|{n_attr + 1}|상,중,하", row
            FX["udp"] = row.split("|")[0]
            tr = [r for r in d.find_elements(By.CSS_SELECTOR, "#udpTable tbody tr") if UDP_NM in r.get_attribute("textContent")]
            assert tr, "목록에 없음"
            cells = [td.get_attribute("textContent").strip() for td in tr[0].find_elements(By.CSS_SELECTOR, "td")]
            print("  목록 행:", cells)
            assert cells[1] == UDP_NM and cells[2] == "컬럼" and cells[3] == "선택", cells
            # 수정 창 — 대상 비활성·안내
            ui.js("arguments[0].querySelector('.v-data-table__checkbox').click();", tr[0])
            time.sleep(0.4)
            ui.js("document.getElementById('udpEdit').click();")
            time.sleep(1)
            tgt = d.find_element(By.ID, "udpTargetTp")
            assert tgt.get_attribute("disabled"), "수정 창에서 대상이 바뀔 수 있음"
            info = ui.js("const v = arguments[0].closest('.v-input'); return [v.querySelector('.v-select__selection').textContent.trim(), "
                         "(v.querySelector('.v-messages__message') || {}).textContent];", tgt)
            print("  수정 창 대상:", info)
            assert info == ["컬럼", "저장한 뒤에는 대상을 바꿀 수 없습니다"], info
            ui.js("[...document.querySelectorAll('.v-dialog--active button')].find(b => b.textContent.trim() === '취소').click();")
            time.sleep(0.8)
        step("C1. 사용자 정의 속성 — 대상 '컬럼' 필수 선택 속성 추가(문구·순서 기본값·기본 열 이름 거부), 목록 대상 열, 수정 창 대상 비활성", _c1)

        def _c2():
            assert FX.get("udp"), "C1 의 속성이 없음"
            ui.open_columns(FX["dm"])
            h = ui.heads()
            names = [x["udpNm"] + (" *" if x["requiredYn"] == "Y" else "") for x in attr_defs()]
            i = h.index("디폴트")
            print("  머리:", h[i:i + len(names) + 2])
            assert h[i + 1:i + 1 + len(names)] == names and h[i + 1 + len(names)] == "컬럼 영문명 (물리)", h
            assert UDP_NM + " *" in names and NUM_NM in names
            th = ui.js("return [...document.querySelectorAll('#clTable_table thead tr:first-child th')].find(t => t.textContent.trim() === arguments[0]);", UDP_NM + " *")
            assert "hdr-udp" in th.get_attribute("class")
            assert ui.js("return arguments[0].querySelector('span[title]').getAttribute('title');", th) == UDP_DESC
            assert ui.js("return getComputedStyle(arguments[0]).backgroundColor;", th) == "rgb(255, 243, 224)"
            # 새 행 — 필수 칸 빨강, 비우고 저장하면 거부 (맨 아래 추가 — 맨 위면 기존 행 순서가 밀려 수정 행이 된다)
            ui.js(CCOL + "c.addTargetKey = c.objOptions.find(o => o.objNm === 'W10_T').key; c.onAddTargetChange(c.addTargetKey); c.addPosition = 'bottom';")
            time.sleep(0.5)
            ui.js("document.getElementById('btn-add-col-row').click();")
            time.sleep(0.8)
            ui.js(CCOL + "c.newRows[0].attrNmKr = '셀신규'; c.newRows[0].attrNm = 'W10_NEW';")
            time.sleep(0.5)
            cls = ui.cell_cls("W10_NEW", UDP_NM + " *")
            assert "inline-error" in cls, cls
            assert "inline-error" not in ui.cell_cls("W10_NEW", NUM_NM)
            assert ui.js("return arguments[0].querySelector('input').type;", ui.cell("W10_NEW", NUM_NM)) == "number"
            ui.save()
            msg = ui.wait_swal("일부 저장 실패")
            print("  비우고 저장:", msg)
            assert f"W10_T.W10_NEW — 필수 관리 항목을 넣어 주세요: {UDP_NM}" in msg, msg
            ui.swal_ok()
            assert psql(f"select count(*) from quality.tb_data_model_attr where dm_id={q(FX['dm'])} and attr_nm='W10_NEW'") == "0"
            ui.pick(ui.cell("W10_NEW", UDP_NM + " *"), "중")
            assert "inline-error" not in ui.cell_cls("W10_NEW", UDP_NM + " *")
            ui.save()
            print("  저장:", ui.wait_swal("저장 완료"))
            time.sleep(2.5)
            v = vals("W10_NEW")
            print("  값:", v)
            assert v == {UDP_NM: "중"}, v
            # 다시 읽은 그리드에 저장한 값 (선택 칸 글자)
            WebDriverWait(d, 10).until(lambda x: x.execute_script(CELL + "return !!cell('W10_NEW', arguments[0]);", UDP_NM + " *"))
            sel = ui.js("return arguments[0].querySelector('.v-select__selection').textContent.trim();", ui.cell("W10_NEW", UDP_NM + " *"))
            assert sel == "중", sel
        step("C2. 컬럼 그리드 — 속성 열 머리('이름 *'·주황·설명), 새 행 필수 빈 칸 빨강·저장 거부 문구, 값 고르고 저장 → DB 값", _c2)

        def _c3():
            assert FX.get("udp"), "C1 의 속성이 없음"
            h0 = hist_cnt()
            ui.pick(ui.cell("C_ID", UDP_NM + " *"), "상")
            inp = ui.cell("C_ID", NUM_NM).find_element(By.CSS_SELECTOR, "input")
            inp.send_keys("1500")
            time.sleep(0.4)
            assert "inline-dirty" in ui.cell_cls("C_ID", UDP_NM + " *") and "inline-dirty" in ui.cell_cls("C_ID", NUM_NM)
            # 같은 행 다른 칸(한글명)은 노랑이 아니다 — 속성만 고친 행
            assert "inline-dirty" not in ui.cell_cls("C_ID", "컬럼 한글명 (논리)")
            body = d.find_element(By.ID, "tab_datamodelStatusColumn").text
            assert "(수정중 1건)" in body, body[:400]
            ui.save()
            print("  속성만:", ui.wait_swal("저장 완료"))
            time.sleep(2.5)
            v = vals("C_ID")
            print("  값:", v, "이력", h0, "→", hist_cnt())
            assert v == {UDP_NM: "상", NUM_NM: "1500"} and hist_cnt() == h0, (v, h0)
            # 영문명 변경(속성 그대로) + 필수 속성이 빈 C_NM 의 한글명만 변경 — 둘 다 저장
            WebDriverWait(d, 10).until(lambda x: x.execute_script(CELL + "return !!cell('C_ID', '컬럼 영문명 (물리)');"))
            en = ui.cell("C_ID", "컬럼 영문명 (물리)").find_element(By.CSS_SELECTOR, "input")
            en.send_keys(Keys.CONTROL, "a")
            en.send_keys("C_ID2")
            kr = ui.cell("C_NM", "컬럼 한글명 (논리)").find_element(By.CSS_SELECTOR, "input")
            kr.send_keys(Keys.CONTROL, "a")
            kr.send_keys("셀이름변경")
            ui.js("document.body.click();")
            time.sleep(0.5)
            assert "inline-error" not in ui.cell_cls("C_NM", UDP_NM + " *"), "속성을 안 고친 행에 필수 빨강"
            ui.save()
            ui.wait_swal("영문명 변경 cascade 확인")
            ui.swal_ok()
            print("  이름 변경:", ui.wait_swal("저장 완료"))
            time.sleep(2.5)
            assert vals("C_ID") == {} and vals("C_ID2") == {UDP_NM: "상", NUM_NM: "1500"}, (vals("C_ID"), vals("C_ID2"))
            assert psql(f"select attr_nm_kr from quality.tb_data_model_attr where dm_id={q(FX['dm'])} and attr_nm='C_NM'") == "셀이름변경"
            assert vals("C_NM") == {}
        step("C3. 기존 행 — 속성만 바꿔 저장(노랑·수정중 1건·이력 그대로), 영문명 변경 → 값 새 이름, 필수 빈 행의 한글명만 수정은 저장", _c3)
    finally:
        d.quit()

    def _c4():
        defs = attr_defs()
        names = [x["udpNm"] for x in defs]
        x = sa.get(BASE + "/api/dm/uploadTemplate", params={"scope": "attrs"}, timeout=20)
        assert x.status_code == 200
        hdr = [c.value for c in next(load_workbook(io.BytesIO(x.content)).active.iter_rows(max_row=1))]
        print("  컬럼 양식 머리 끝:", hdr[len(ATTR_HEADERS):])
        assert hdr == ATTR_HEADERS + names and UDP_NM in names and NUM_NM in names, hdr
        t = sa.get(BASE + "/api/dm/uploadTemplate", params={"scope": "tables"}, timeout=20)
        thdr = [c.value for c in next(load_workbook(io.BytesIO(t.content)).active.iter_rows(max_row=1))]
        assert UDP_NM not in thdr and NUM_NM not in thdr, thdr
        # 업로드 미리보기 — 등록할 행만 검사 (이미 있는 C_NM 은 필수가 비어도 건너뜀)
        wb = Workbook()
        ws = wb.active
        ws.append(hdr)

        def row(col_en, col_kr, grade, num):
            r = {"소유자": "APP", "테이블명(영문)": "W10_T", "컬럼명(영문)": col_en, "컬럼명(한글)": col_kr, "데이터타입": "VARCHAR", "길이": 10,
                 UDP_NM: grade, NUM_NM: num}
            ws.append([r.get(hh) for hh in hdr])
        row("W10_UP1", "셀업로드1", None, None)
        row("W10_UP2", "셀업로드2", "상", "1,000")
        row("W10_UP3", "셀업로드3", "중", "abc")
        row("C_NM", "셀이름변경", None, None)
        buf = io.BytesIO()
        wb.save(buf)
        r = ok(sa.post(BASE + "/api/dm/uploadAttrs", data={"dataModelId": FX["dm"], "mode": "preview"},
                       files={"file": ("w10udp.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=30))
        p = json.loads(r["contents"])
        acts = {a["attrNm"]: (a["_action"], a.get("_msg")) for a in p["attrs"]}
        print("  미리보기:", acts)
        assert acts["W10_UP1"] == ("ERROR", f"필수 관리 항목을 넣어 주세요: {UDP_NM}"), acts
        assert acts["W10_UP2"][0] == "INSERT" and acts["W10_UP3"] == ("ERROR", f"'{NUM_NM}' 은 숫자여야 합니다 (abc)."), acts
        assert acts["C_NM"][0] == "SKIP", acts
        assert sorted(e["row"] for e in p["errors"]) == [2, 4], p["errors"]
        # 데이터 다운로드 — 끝 열에 속성 값
        x = sa.get(BASE + "/api/dm/downloadDataModelAttrs", params={"clctId": FX["dm"]}, timeout=30)
        assert x.status_code == 200
        rows = [[c for c in r] for r in load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True)]
        hi = [i for i, r in enumerate(rows) if r and r[0] == "소유자"][0]
        dh = list(rows[hi])
        print("  데이터 머리 끝:", dh[len(ATTR_HEADERS):])
        assert dh[:len(ATTR_HEADERS) + len(names)] == ATTR_HEADERS + names, dh
        byc = {r[3]: dict(zip(dh, r)) for r in rows[hi + 1:] if r and r[3]}
        assert byc["C_ID2"][UDP_NM] == "상" and byc["C_ID2"][NUM_NM] == "1500" and byc["W10_NEW"][UDP_NM] == "중", byc
        assert not byc["C_NM"][UDP_NM], byc["C_NM"]
    step("C4. 엑셀 — 컬럼 양식 머리 끝 속성 이름(테이블 양식엔 없음), 업로드 미리보기 등록 행만 검사, 데이터 다운로드 속성 값", _c4)

    def _c5():
        assert FX.get("udp"), "C1 의 속성이 없음"
        d2 = new_driver(USER, FX["upw"])
        u = Ui(d2)
        try:
            assert d2.execute_script("return !!document.getElementById('nav_udpDef');") is False, "일반 사용자에게 사용자 정의 속성 메뉴"
            u.open_columns(FX["dm"])
            assert UDP_NM + " *" in u.heads()
            u.pick(u.cell("C_NM", UDP_NM + " *"), "하")
            u.save()
            print("  일반 사용자 저장:", u.wait_swal("저장 완료"))
            time.sleep(2)
            assert vals("C_NM") == {UDP_NM: "하"}, vals("C_NM")
            assert psql(f"select aprv_status from quality.tb_data_model_attr where dm_id={q(FX['dm'])} and attr_nm='C_NM'") in ("APPROVED", ""), "속성만 저장했는데 컬럼 승인 상태가 바뀜"
            u.open_menu("dmGroup", "nav_datamodelStatusTable")
            WebDriverWait(d2, 10).until(lambda x: x.execute_script(CTBL + "return !!c;"))
            u.js(CTBL + "c.selectedModelId = arguments[0]; c.loadAreaOptions(); c.openAddObjDialog();", FX["dm"])
            time.sleep(1.5)
            assert not d2.find_elements(By.ID, "udp_" + FX["udp"]) and not d2.find_elements(By.ID, "udp_" + FX["num"]), "테이블 창에 컬럼 속성 칸"
            assert d2.find_elements(By.ID, "udp_UDP_RETENTION"), "테이블 창에 테이블 속성 칸이 없음"
        finally:
            d2.quit()
    step("C5. 시험 일반 사용자 — 사용자 정의 속성 메뉴 없음, 그리드 속성 칸 편집·저장, 테이블 창에 컬럼 속성 칸 없음", _c5)

    def _c6():
        assert FX.get("udp"), "C1 의 속성이 없음"
        # 관리 화면에서 선택값을 고쳐 W10_NEW 에 저장된 '중' 이 선택값에서 빠진다
        d0 = json.loads(psql(f"select json_build_object('udpId', udp_id, 'targetTp', target_tp, 'udpNm', udp_nm, 'dataTp', data_tp, "
                             f"'requiredYn', required_yn, 'sortOrder', sort_order, 'useYn', use_yn, 'udpDesc', udp_desc) "
                             f"from quality.tb_udp_def where udp_id={q(FX['udp'])}"))
        ok(sa.post(BASE + "/api/udp/save", json=dict(d0, selectVals=["상", "하", "특"]), timeout=10))
        assert vals("W10_NEW") == {UDP_NM: "중"}
        d3 = new_driver("space", "123")
        u = Ui(d3)
        try:
            u.open_columns(FX["dm"])
            g = UDP_NM + " *"
            sel = u.js("return arguments[0].querySelector('.v-select__selection').textContent.trim();", u.cell("W10_NEW", g))
            print("  선택값에 없는 저장 값:", sel, u.cell_cls("W10_NEW", g))
            assert sel == "중", sel
            assert "inline-error" not in u.cell_cls("W10_NEW", g), "고치지 않은 행에 빨강"
            # 같은 행의 다른 속성 칸만 고치면 그 칸이 빨강 (저장하면 서버가 거부)
            u.cell("W10_NEW", NUM_NM).find_element(By.CSS_SELECTOR, "input").send_keys("7")
            time.sleep(0.4)
            assert "inline-error" in u.cell_cls("W10_NEW", g), u.cell_cls("W10_NEW", g)
            u.save()
            msg = u.wait_swal("일부 저장 실패")
            print("  저장:", msg)
            assert f"W10_T.W10_NEW — '{UDP_NM}' 은 상·하·특 중 하나여야 합니다 (중)." in msg, msg
            u.swal_ok()
            # 목록에 선택값 뒤로 저장된 값 — 다른 값을 고르면 빨강이 풀리고 저장된다
            u.js("arguments[0].querySelector('.v-input__slot').click();", u.cell("W10_NEW", g))
            WebDriverWait(d3, 5).until(lambda x: x.execute_script("return document.querySelectorAll('.menuable__content__active .v-list-item').length > 0;"))
            opts = u.js("return [...document.querySelectorAll('.menuable__content__active .v-list-item')].map(e => e.textContent.trim());")
            print("  목록:", opts)
            assert opts == ["상", "하", "특", "중"], opts
            u.js("[...document.querySelectorAll('.menuable__content__active .v-list-item')].find(e => e.textContent.trim() === '특').click();")
            time.sleep(0.5)
            assert "inline-error" not in u.cell_cls("W10_NEW", g)
            u.save()
            print("  다시 저장:", u.wait_swal("저장 완료"))
            time.sleep(2)
            assert vals("W10_NEW") == {UDP_NM: "특", NUM_NM: "7"}, vals("W10_NEW")
        finally:
            d3.quit()
    step("C6. 선택값에서 빠진 저장 값 — 칸에 보임(고치지 않은 행은 빨강 없음), 같은 행 다른 속성만 고치면 빨강·저장 거부 문구, 다른 값 고르면 저장", _c6)


if __name__ == "__main__":
    t0 = time.time()
    try:
        main()
    finally:
        cleanup()
    p = sum(1 for _, st in results if st == "PASS")
    f = sum(1 for _, st in results if st == "FAIL")
    print(f"\n{'=' * 60}\n결과: {p} PASS / {f} FAIL  ({time.time() - t0:.0f}초)\n{'=' * 60}")
    for n, st in results:
        print(f"  [{st}] {n}")
    sys.exit(0 if f == 0 else 1)
