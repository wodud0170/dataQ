# -*- coding: utf-8 -*-
"""
110번 기존 기능 보완 2묶음(데이터 모델·모델 변경 승인 — 4.12·5.8 절) 조치 검증 — API + 화면

  S1. 수집·모델 관리 (4.12-13-a·14-a·14-b·36-b·38·41-a·41-c·41-d, 5.8-25) — 개발 DB PostgreSQL 에 시험 스키마를 만들어 수집
  S2. 컬럼 편집 (4.12-1·2·17·18·23·24·31, 5.8-28) — 소문자 컬럼명, 숨은 같은 이름 행, TMP_TBL 번호, 이름 변경 미리보기, 지운 API
  S3. 내려받기 (4.12-11·15-a·22·34) — 남의 미승인 행 제외, 소유자별 묶음, 엑셀 열·참조 한글명, XMI 한글명·길이·빈 테이블
  S4. 업로드·가져오기 (4.12-21-a·b·c·33-a·b) — 컬럼 순서, 소유자별 순서, FK 저장·해석 실패, 가져오기 소유자
  S5. 승인·반려·롤백 (5.8-13·14·15·16·17·20·26-b·27) — 행 상태 유지, 이름 변경 연쇄 원복, 전파 반려 원복, 롤백 원복, 모델별 신청
  S6. 기타 (5.8-6·30·31-b·32) — 진단 제외 통계·오류 메시지, 지운 업무영역 이름
  S7. 화면 — 셀레니움 Edge(기본 창 없음): 승인 화면 선택, 진단 제외 화면, 수집이력 모델 선택, 모델 관리(이름 공백·수집 확인창·진행 창)

필요:
  - q-center(28091)·q-executor(28098) 기동
  - 환경 변수 PGPASSWORD (개발 DB admin), DATAQ_TEST_PW (시험 계정 space·jyjang 비밀번호), 선택 PSQL (psql.exe 경로)
  - 인자로 돌릴 묶음을 고를 수 있다 (예: python verify_110_batch2.py S1 S5). 없으면 전부 (S2·S7 은 S1 의 수집 모델을 쓴다)
시험 데이터(시험 스키마·역할, 데이터 소스, 모델과 그 테이블·컬럼·제약·이력, 업무영역, 결재·알림)는 끝에서 지운다.
"""
import base64, io, json, os, secrets, subprocess, sys, tempfile, time, traceback
import requests
import openpyxl

BASE = "http://localhost:28091"
TAG = time.strftime("%H%M%S")
PSQL = os.environ.get("PSQL", r"C:\Program Files\PostgreSQL\13\bin\psql.exe")
PW = os.environ["DATAQ_TEST_PW"]
results = []
cleanup = []          # 끝에서 거꾸로 실행할 정리 함수
FX = {}               # 공통 시험 데이터
TMP = tempfile.mkdtemp(prefix="v110b2_")


def step(name, fn):
    print(f"\n=== {name}")
    try:
        fn()
        results.append((name, "PASS"))
        print("  >> PASS")
    except Exception:
        traceback.print_exc()
        results.append((name, "FAIL"))


def psql(sql):
    r = subprocess.run([PSQL, "-h", "localhost", "-p", "25433", "-U", "admin", "-d", "postgres", "-w", "-At", "-c", sql],
                       capture_output=True, text=True, timeout=60, encoding="utf-8", errors="replace",
                       env=dict(os.environ, PGCLIENTENCODING="UTF8"))
    if r.returncode != 0:
        raise RuntimeError("psql 실패: " + (r.stderr or "")[:300])
    return (r.stdout or "").strip()


def login_api(user, pwd=None):
    s = requests.Session()
    r = s.post(BASE + "/login", data={"id": user, "password": base64.b64encode((pwd or PW).encode()).decode()},
               allow_redirects=False, timeout=10)
    assert r.status_code == 200, f"로그인 실패 {user} {r.status_code}"
    return s


def post(s, path, body, timeout=30):
    r = s.post(BASE + path, json=body, timeout=timeout)
    try:
        return r.json()
    except Exception:
        return {"http": r.status_code, "text": r.text[:300]}


def get(s, path, params=None, timeout=30):
    r = s.get(BASE + path, params=params, timeout=timeout)
    try:
        return r.json()
    except Exception:
        return {"http": r.status_code, "text": r.text[:300]}


def contents(r):
    c = r.get("contents")
    return json.loads(c) if isinstance(c, str) and c else (c or {})


def code(r):
    return r.get("resultCode")


def msg(r):
    return r.get("resultMessage") or ""


def q(v):
    return "'" + str(v).replace("'", "''") + "'"


# ---------------------------------------------------------------- 모델·이력 헬퍼
def make_model(s, nm, ds=None):
    r = post(s, "/api/dm/createDataModel", {"dataModelNm": nm, "dataModelDsId": ds, "modelType": "PHYSICAL", "ver": "1",
                                            "dictId": "DEFAULT"})
    assert code(r) == 200, r
    dm = psql(f"select dm_id from quality.tb_data_model where dm_nm={q(nm)}")
    cleanup.append(lambda: drop_model(dm))
    return dm


def drop_model(dm):
    psql(f"delete from quality.tb_aprv_doc_step where doc_id in (select doc_id from quality.tb_aprv_doc where dm_id='{dm}')")
    psql(f"delete from quality.tb_aprv_doc where dm_id='{dm}'")
    for t in ("tb_data_model_change_history", "tb_data_model_attr", "tb_data_model_constraint", "tb_data_model_index",
              "tb_data_model_obj", "tb_data_model_clct", "tb_data_model_schema", "tb_udp_val", "tb_term_resolve_history"):
        psql(f"delete from quality.{t} where dm_id='{dm}'")
    psql(f"delete from quality.tb_struct_diag_detail where diag_id in (select diag_id from quality.tb_struct_diag_history where data_model_id='{dm}')")
    psql(f"delete from quality.tb_struct_diag_history where data_model_id='{dm}'")
    psql(f"delete from quality.tb_data_model where dm_id='{dm}'")


def add_obj(s, dm, owner, obj, kr="시험테이블"):
    return post(s, "/api/dm/addObj", {"dataModelId": dm, "objOwner": owner, "objNm": obj, "objNmKr": kr})


def save_attrs(s, dm, owner, obj, attrs):
    return post(s, "/api/dm/saveAttrs", {"dataModelId": dm, "objNm": obj, "objOwner": owner, "attrs": attrs})


def add_col(nm, kr, ln=10, **kw):
    r = {"mode": "ADD", "attrNm": nm, "attrNmKr": kr, "dataType": "VARCHAR", "dataLen": ln, "nullableYn": "Y"}
    r.update(kw)
    return r


def attr(dm, owner, obj, col, cols="aprv_status||'|'||use_yn||'|'||data_len"):
    return psql(f"select {cols} from quality.tb_data_model_attr where dm_id='{dm}' and coalesce(obj_owner,'')='{owner}' "
                f"and obj_nm='{obj}' and attr_nm='{col}'")


def seq_of(dm, change_type, obj=None, col=None, status=None):
    w = f"dm_id='{dm}' and change_type='{change_type}'"
    if obj:
        w += f" and obj_nm='{obj}'"
    if col:
        w += f" and attr_nm='{col}'"
    if status:
        w += f" and aprv_status='{status}'"
    v = psql(f"select max(change_seq) from quality.tb_data_model_change_history where {w}")
    assert v, f"이력 없음 {change_type} {obj} {col} {status}"
    return int(v)


def hist_status(seq):
    return psql(f"select aprv_status from quality.tb_data_model_change_history where change_seq={seq}")


# ---------------------------------------------------------------- 시험 DB (PostgreSQL)
def fixture_pg(sa):
    """개발 DB 에 시험 역할 1개와 스키마 3개(a·b — 같은 이름 테이블, c — 수집 중 실패), 데이터 소스 2개"""
    role = ("zz110r" + TAG).lower()
    pw = "Zz" + secrets.token_hex(6)
    a, b, c = (f"zz110a{TAG}", f"zz110b{TAG}", f"zz110c{TAG}")
    psql(f"create role {role} login password '{pw}'")
    for s in (a, b, c):
        psql(f"create schema {s} authorization {role}")

    def _drop():
        for s in (a, b, c):
            psql(f"drop schema if exists {s} cascade")
        psql(f"drop role if exists {role}")
    cleanup.append(_drop)
    psql(f"create table {a}.t_same (id int primary key, nm varchar(20)); alter table {a}.t_same owner to {role}")
    psql(f"create table {b}.t_same (id int primary key, nm varchar(20)); alter table {b}.t_same owner to {role}")
    psql(f"create table {a}.t_lower (cust_id int, cust_nm varchar(30)); alter table {a}.t_lower owner to {role}")
    psql(f"comment on column {a}.t_lower.cust_nm is '고객명'")
    psql(f"create table {c}.t_bad (c1 int); alter table {c}.t_bad owner to {role}")
    psql(f"comment on column {c}.t_bad.c1 is '{'가' * 300}'")   # ATTR_NM_KR(255) 를 넘어 컬럼 단계에서 실패
    FX.update(role=role, a=a, b=b, c=c)
    ds = {}
    for k, dbms in (("pg", "PostgreSQL"), ("alt", "Altibase")):
        dsn = f"ZZ110_{k}_{TAG}"
        r = post(sa, "/api/sysinfo/createDataSource", {"dsn": dsn, "dsTp": 0, "dbmsTp": dbms, "driverName": dbms,
                                                       "svrAddr": "127.0.0.1", "port": 25433, "userId": role, "pwd": pw, "dbName": "postgres"})
        assert code(r) == 200, r
        ds[k] = psql(f"select ds_id from ndata.tb_data_source where dsn='{dsn}'")
        cleanup.append(lambda i=ds[k]: psql(f"delete from ndata.tb_data_source where ds_id='{i}'"))
    psql(f"update ndata.tb_data_source set conn_test_yn='Y' where ds_id in ('{ds['pg']}', '{ds['alt']}')")
    FX["ds"] = ds


def collect(sa, dm, schemas, wait=True):
    assert post(sa, "/api/dm/saveDataModelSchemas", [{"dataModelId": dm, "schemaNm": s, "useYn": "Y"} for s in schemas]).get("resultCode") == 200
    before = int(psql(f"select count(*) from quality.tb_data_model_clct where dm_id='{dm}' and clct_end_dt is not null"))
    r = post(sa, "/api/dm/collectDataModel", {"dataModelId": dm, "dataModelDsId": FX["ds"]["pg"], "dataModelNm": "시험"}, timeout=60)
    if not wait:
        return r
    assert code(r) == 200, r
    for _ in range(60):
        time.sleep(2)
        if int(psql(f"select count(*) from quality.tb_data_model_clct where dm_id='{dm}' and clct_end_dt is not null")) > before:
            break
    return psql(f"select clct_cmptn_yn||'|'||coalesce(obj_cnt::text,'')||'|'||coalesce(attr_cnt::text,'') from quality.tb_data_model_clct "
                f"where dm_id='{dm}' order by clct_start_dt desc, clct_end_dt desc limit 1")


# ---------------------------------------------------------------- S1
def s1(sa, su):
    fixture_pg(sa)
    a, b, c = FX["a"], FX["b"], FX["c"]

    def _name():
        r1 = post(sa, "/api/dm/createDataModel", {"dataModelNm": "   ", "modelType": "PHYSICAL", "ver": "1", "dictId": "DEFAULT"})
        dm = make_model(sa, "ZZ110수집모델_" + TAG, FX["ds"]["pg"])
        FX["m1"], FX["m1Nm"] = dm, "ZZ110수집모델_" + TAG
        r2 = post(sa, "/api/dm/updateDataModel", {"dataModelId": dm, "dataModelNm": "", "dataModelDsId": FX["ds"]["pg"],
                                                  "ver": "1", "dictId": "DEFAULT"})
        nm = psql(f"select dm_nm from quality.tb_data_model where dm_id='{dm}'")
        cnt = psql("select count(*) from quality.tb_data_model where trim(dm_nm)=''")
        print(f"  공백 이름 등록={code(r1)} {msg(r1)} / 빈 이름 수정={code(r2)} {msg(r2)} / 이름 그대로={nm} / 빈 이름 모델={cnt}")
        assert code(r1) == 500 and msg(r1) == "데이터 모델명은 필수 입력값입니다." and code(r2) == 500 and nm == FX["m1Nm"] and cnt == "0"
    step("S1-1. 모델명 빈칸·공백 등록·수정 거부 (4.12-13-a)", _name)

    def _collect():
        dm = FX["m1"]
        res = collect(sa, dm, [a, b])
        rows = psql(f"select string_agg(obj_owner||'.'||obj_nm||'.'||attr_nm||':'||coalesce(array_to_string(word_lst,'/'),'(없음)'), ',' "
                    f"order by obj_owner, obj_nm, attr_nm) from quality.tb_data_model_attr where dm_id='{dm}'")
        print(f"  수집 결과 완료|테이블|컬럼={res}\n  컬럼·단어 목록={rows}")
        assert res == "Y|3|6", res
        assert "(없음)" not in rows and f"{a}.t_lower.cust_nm:CUST(고객)/NM(명)" in rows   # 4.12-41-d — 단어 목록이 저장된다
    step("S1-2. PostgreSQL 수집 — 수집이력에 그 수집의 테이블 3·컬럼 6, 단어 목록 저장 (4.12-14-a·41-d)", _collect)

    def _history():
        dm = FX["m1"]
        other = make_model(sa, FX["m1Nm"] + "X")          # 이름이 포함 관계인 다른 모델
        psql(f"insert into quality.tb_data_model_clct (dm_clct_id, dm_id, clct_type, clct_start_dt, clct_end_dt, clct_cmptn_yn, obj_cnt, attr_cnt) "
             f"values ('Z2C{TAG}', '{other}', 'DBMS', to_char(now(),'YYYYMMDDHH24MISS'), to_char(now(),'YYYYMMDDHH24MISS'), 'Y', 9, 99)")
        assert code(add_obj(sa, dm, a, "t_extra")) == 200           # 수집 뒤 모델에 테이블이 늘어도
        rows = post(sa, "/api/dm/getDataModelHistoryList", {"schId": dm})
        ids = {r["dataModelId"] for r in rows}
        cnts = [(r["objCnt"], r["attrCnt"]) for r in rows]
        old = post(sa, "/api/dm/getDataModelHistoryList", {"schNm": FX["m1Nm"]})
        print(f"  모델 ID 로 조회 {len(rows)}건 모델={len(ids)} 건수={cnts} / 이름 부분 일치 조회 {len(old)}건")
        assert ids == {dm} and cnts == [(3, 6)] and len(old) == 2
        psql(f"delete from quality.tb_data_model_obj where dm_id='{dm}' and obj_nm='t_extra'")
        psql(f"delete from quality.tb_data_model_change_history where dm_id='{dm}' and obj_nm='t_extra'")
    step("S1-3. 수집이력 — 모델 ID 로 거르고(이름 포함 관계 모델 제외), 건수는 그 수집 때 값 (4.12-14-a·14-b)", _history)

    def _unsupported():
        dm = make_model(sa, "ZZ110알티모델_" + TAG, FX["ds"]["alt"])
        r = post(sa, "/api/dm/collectDataModel", {"dataModelId": dm, "dataModelDsId": FX["ds"]["alt"]})
        clct = psql(f"select count(*) from quality.tb_data_model_clct where dm_id='{dm}'")
        print(f"  Altibase 수집 요청={code(r)} {msg(r)} / 수집 이력={clct}")
        assert code(r) == 400 and msg(r).startswith("Altibase 데이터 소스는 수집을 지원하지 않습니다") and clct == "0"
    step("S1-4. 수집 쿼리가 없는 DBMS 는 요청 때 거부·안내 (4.12-36-b)", _unsupported)

    def _failed():
        dm = make_model(sa, "ZZ110실패모델_" + TAG, FX["ds"]["pg"])
        res = collect(sa, dm, [c])
        end = psql(f"select count(*) from quality.tb_data_model_clct where dm_id='{dm}' and clct_end_dt is not null")
        r2 = collect(sa, dm, [c], wait=False)
        print(f"  컬럼 단계 실패 — 완료|테이블|컬럼={res} 끝난 시각 있음={end} / 바로 다시 수집 요청={code(r2)} {msg(r2)}")
        assert res.startswith("N|") and end == "1" and code(r2) == 200
        for _ in range(30):
            time.sleep(2)
            if psql(f"select count(*) from quality.tb_data_model_clct where dm_id='{dm}' and clct_end_dt is not null") == "2":
                break
    step("S1-5. 테이블 단계 커밋 뒤 실패한 수집도 끝난 시각·미완료(N)로 닫혀 바로 다시 수집할 수 있다 (4.12-38)", _failed)

    def _recollect():
        dm = FX["m1"]
        # 화면에서 한글명 수정·FK 생성, 표준 여부를 DB 에서 일부러 Y 로
        assert code(save_attrs(sa, dm, a, "t_lower", [{"mode": "UPDATE", "origAttrNm": "cust_nm", "attrNm": "cust_nm", "attrNmKr": "고객명수정",
                                                       "nullableYn": "Y", "pkYn": "N", "fkYn": "N"}])) == 200
        r = post(sa, "/api/dm/createFk", {"dataModelId": dm, "objOwner": a, "tableNm": "t_lower", "refOwner": a, "refTableNm": "t_same",
                                          "mappings": [{"ownAttrNm": "cust_id", "refAttrNm": "id"}]})
        assert code(r) == 200, r
        # nm 은 표준 용어가 아니다 — 일부러 Y 로 둔 뒤 다시 수집하면 다시 판정해 N
        psql(f"update quality.tb_data_model_attr set terms_stnd_yn='Y', domain_stnd_yn='Y' where dm_id='{dm}' and obj_owner='{b}' and obj_nm='t_same' and attr_nm='nm'")
        psql(f"alter table {a}.t_lower alter column cust_nm type varchar(40)")
        res = collect(sa, dm, [a, b])
        nm = attr(dm, a, "t_lower", "cust_nm", "attr_nm_kr||'|'||data_len")
        flag = attr(dm, b, "t_same", "nm", "terms_stnd_yn||domain_stnd_yn")
        fk = attr(dm, a, "t_lower", "cust_id", "fk_yn||'|'||coalesce(fk_parent_obj_owner,'')||'.'||coalesce(fk_parent_obj_nm,'')||'.'||coalesce(fk_parent_attr_nm,'')")
        print(f"  다시 수집={res} / cust_nm 한글명|길이={nm} / nm 표준 여부={flag} / cust_id FK={fk}")
        assert res.startswith("Y|") and nm == "고객명수정|40" and flag == "NN" and fk == f"Y|{a}.t_same.id"
    step("S1-6. 다시 수집 — 한글명 그대로, 타입·길이 갱신, 표준 여부 다시 판정, 화면에서 만든 FK 부모 그대로 (4.12-41-a·41-c)", _recollect)

    def _struct_off():
        dm = FX["m1"]
        r = post(sa, "/api/dm/setObjDiagTarget", {"dmId": dm, "objOwner": a, "objNm": "t_same", "diagType": "STRUCT", "targetYn": "N"})
        assert r.get("success") is True, r
        r = post(sa, "/api/std/structDiag/execute", {"dataModelId": dm}, timeout=60)
        assert code(r) == 200, r
        did = r.get("contents")
        for _ in range(60):
            time.sleep(2)
            if psql(f"select status from quality.tb_struct_diag_history where diag_id='{did}'") in ("DONE", "FAIL", "ERROR"):
                break
        st = psql(f"select status from quality.tb_struct_diag_history where diag_id='{did}'")
        rows = psql(f"select coalesce(string_agg(owner||'.'||table_nm||'.'||coalesce(column_nm,'')||':'||change_type, ',' order by seq),'') "
                    f"from quality.tb_struct_diag_detail where diag_id='{did}' and table_nm='t_same'")
        print(f"  {a}.t_same 구조 진단 OFF 뒤 진단={st} / t_same 차이={rows or '(없음)'}")
        assert st == "DONE" and rows == ""
    step("S1-7. 구조 진단 OFF 는 그 소유자의 테이블만 — 다른 소유자의 같은 이름 테이블은 삭제로 잡히지 않음 (5.8-25)", _struct_off)


# ---------------------------------------------------------------- S2
def s2(sa, su):
    dm, a = FX["m1"], FX["a"]

    def _lower():
        r = save_attrs(sa, dm, a, "t_lower", [{"mode": "UPDATE", "origAttrNm": "cust_nm", "attrNm": "CUST_NM", "attrNmKr": "고객이름",
                                               "nullableYn": "Y", "pkYn": "N", "fkYn": "N"}])
        nm = psql(f"select string_agg(attr_nm, ',') from quality.tb_data_model_attr where dm_id='{dm}' and obj_nm='t_lower' and lower(attr_nm)='cust_nm'")
        h = psql(f"select change_tier||'|'||attr_nm from quality.tb_data_model_change_history where dm_id='{dm}' and change_type='MODIFY_ATTR' "
                 f"and obj_nm='t_lower' order by change_seq desc limit 1")
        print(f"  한글명만 수정(영문명 대문자로 보냄)={code(r)} / 저장된 이름={nm} / 이력={h}")
        assert code(r) == 200 and nm == "cust_nm" and h == "TIER2|cust_nm"
        r1 = save_attrs(su, dm, a, "t_lower", [{"mode": "DELETE", "attrNm": "cust_nm"}])
        r2 = save_attrs(su, dm, a, "t_lower", [{"mode": "DELETE", "attrNm": "no_such_col"}])
        st = attr(dm, a, "t_lower", "cust_nm")
        h2 = psql(f"select count(*) from quality.tb_data_model_change_history where dm_id='{dm}' and change_type='DEL_ATTR' and attr_nm='no_such_col'")
        cnt = psql(f"select obj_attr_cnt from quality.tb_data_model_obj where dm_id='{dm}' and obj_nm='t_lower'")
        print(f"  소문자 컬럼 삭제(일반 사용자)={code(r1)} → {st} / 없는 컬럼 삭제={code(r2)} {msg(r2)} 이력={h2} / 테이블 컬럼 수={cnt}")
        assert code(r1) == 200 and st == "APPROVED|N|40" and code(r2) == 500 and msg(r2).startswith("삭제 대상 컬럼을 찾지 못했습니다") and h2 == "0"
        assert cnt == "1"   # 4.12-31 — 삭제 대기 컬럼은 세지 않는다
        FX["delLower"] = seq_of(dm, "DEL_ATTR", "t_lower", "cust_nm", "DRAFT")
    step("S2-1. 소문자(PostgreSQL) 컬럼 — 한글명 수정은 이름 그대로 TIER2, 삭제는 저장된 이름으로, 없으면 오류, 컬럼 수에서 삭제 대기 제외 (4.12-23·31)", _lower)

    def _hidden():
        assert code(save_attrs(su, dm, a, "t_same", [add_col("zz_hid", "숨은컬럼")])) == 200
        r = save_attrs(sa, dm, a, "t_same", [add_col("zz_hid", "관리자컬럼")])
        st = attr(dm, a, "t_same", "ZZ_HID", "aprv_status||'|'||requester_user_id||'|'||attr_nm_kr")
        print(f"  다른 사용자 승인 대기 행과 같은 이름 추가={code(r)} {msg(r)} / 행={st}")
        assert code(r) == 500 and msg(r).startswith("이미 있는 컬럼입니다") and st == "DRAFT|jyjang|숨은컬럼"
        # 반려된 같은 이름은 다시 추가할 수 있다
        seq = seq_of(dm, "ADD_ATTR", "t_same", "ZZ_HID", "DRAFT")
        assert code(post(sa, "/api/dmApproval/reject", {"changeSeqList": [seq], "aprvComment": "시험"})) == 200
        r2 = save_attrs(su, dm, a, "t_same", [add_col("zz_hid", "다시추가")])
        st2 = attr(dm, a, "t_same", "ZZ_HID", "aprv_status||'|'||attr_nm_kr")
        print(f"  반려 뒤 다시 추가={code(r2)} / 행={st2}")
        assert code(r2) == 200 and st2 == "DRAFT|다시추가"
    step("S2-2. 화면에 안 보이는 같은 이름 행(남의 승인 대기)을 덮어쓰지 않음, 반려된 행은 다시 추가 (4.12-24)", _hidden)

    def _tmp_tbl():
        dm2 = make_model(sa, "ZZ110논리모델_" + TAG)
        FX["m2"] = dm2
        assert code(add_obj(sa, dm2, "", "TMP_TBL_2")) == 200
        r = post(sa, "/api/dm/addObj", {"dataModelId": dm2, "objOwner": "", "objNmKr": "이름없는테이블"})
        names = psql(f"select string_agg(obj_nm, ',' order by obj_nm) from quality.tb_data_model_obj where dm_id='{dm2}'")
        print(f"  TMP_TBL_2 가 있을 때 이름 없이 추가={code(r)} {msg(r)} / 테이블={names}")
        assert code(r) == 200 and names == "TMP_TBL_2,TMP_TBL_3"
    step("S2-3. 이름 없는 테이블 추가는 기존 TMP_TBL_N 최대+1 (4.12-18)", _tmp_tbl)

    def _preview():
        p = get(sa, "/api/dm/previewObjRename", {"dataModelId": dm, "objOwner": a, "origObjNm": "t_lower", "newObjNm": "t_same"})
        p0 = get(sa, "/api/dm/previewObjRename", {"dataModelId": dm, "origObjNm": "t_lower", "newObjNm": "t_same"})
        print(f"  소유자 있음 → {p} / 소유자 없음(빈 소유자로 셈) → {p0}")
        assert p["attrCnt"] == 1 and p["constraintCnt"] == 1 and p["conflict"] is True
        assert p0["attrCnt"] == 0 and p0["conflict"] is False
    step("S2-4. 테이블 이름 변경 미리보기가 소유자로 센다 (4.12-17)", _preview)

    def _swap():
        r = post(sa, "/api/dm/swapAttrOrd", {"dataModelId": dm, "objOwner": FX["b"], "objNm": "t_same", "attrNm": "nm", "direction": "UP"})
        js = psql(f"select (before_json::json->>'id')||'|'||(before_json::json->>'nm')||'|'||(after_json::json->>'id')||'|'||(after_json::json->>'nm') "
                  f"from quality.tb_data_model_change_history where dm_id='{dm}' and change_type='SWAP_ATTR_ORD' order by change_seq desc limit 1")
        print(f"  순서 바꾸기={code(r)} / 이력 JSON id·nm 전·후={js}")
        assert code(r) == 200 and js == "1|2|2|1"
    step("S2-5. 컬럼 순서 바꾸기 이력이 JSON {컬럼: 순서} (5.8-28)", _swap)

    def _removed():
        out = []
        for path, kind in (("/api/dm/resolveStandard?termsNm=x", "GET"), ("/api/dm/getConstraintsByTable?dataModelId=x&objNm=x", "GET"),
                           ("/api/dm/addAttr", "POST"), ("/api/dm/updateAttr", "POST"), ("/api/dm/deleteAttr", "POST"),
                           ("/api/dm/getDataModelClctList", "POST"), ("/api/dm/getDataModelStatsByClctId?clctId=x", "GET"),
                           ("/api/dm/getDataModelObjListByObjNm", "POST"), ("/api/dm/getDataModelAttrListByRetreiveCond", "POST")):
            r = sa.get(BASE + path, timeout=15) if kind == "GET" else sa.post(BASE + path, json={}, timeout=15)
            # 없는 GET 경로는 화면(SPA) 첫 페이지 HTML 로 돌아온다 — API 응답(JSON)이 아니면 없는 것
            gone = r.status_code in (404, 405) or "text/html" in (r.headers.get("content-type") or "")
            out.append((path.split("?")[0].rsplit("/", 1)[1], r.status_code if not gone else "없음"))
        before = attr(dm, FX["b"], "t_same", "nm", "attr_nm||'|'||coalesce(attr_nm_kr,'')||'|'||data_len")
        r = post(sa, "/api/dm/resolveAttrs", {"dataModelId": dm, "dryRun": False,
                                              "attrs": [{"objOwner": FX["b"], "objNm": "t_same", "attrNm": "nm", "attrNmKr": "명"}]})
        after = attr(dm, FX["b"], "t_same", "nm", "attr_nm||'|'||coalesce(attr_nm_kr,'')||'|'||data_len")
        print(f"  지운 API → {out}\n  resolveAttrs(dryRun=false) tried={r.get('tried')} 모델 그대로={before == after}")
        assert all(st == "없음" for _, st in out) and before == after
    step("S2-6. 호출부 없는 API 9개 없음, 한글명 기준 표준화는 후보만 돌려준다 (4.12-1·2·8-g)", _removed)


# ---------------------------------------------------------------- S3
def s3(sa, su):
    dm = make_model(sa, "ZZ110내려받기모델_" + TAG)
    FX["m4"] = dm

    def _setup():
        assert code(add_obj(sa, dm, "X", "T1", "엑스테이블")) == 200
        assert code(add_obj(sa, dm, "Y", "T1", "와이테이블")) == 200
        assert code(add_obj(sa, dm, "X", "T_EMPTY", "빈테이블")) == 200
        assert code(save_attrs(sa, dm, "X", "T1", [add_col("C_A", "에이", attrOrder=1), add_col("C_B", "비", 20, attrOrder=2)])) == 200
        assert code(save_attrs(sa, dm, "Y", "T1", [add_col("C_C", "씨", attrOrder=1), add_col("C_REF", "참조", attrOrder=2)])) == 200
        r = post(sa, "/api/dm/createFk", {"dataModelId": dm, "objOwner": "Y", "tableNm": "T1", "refOwner": "X", "refTableNm": "T1",
                                          "mappings": [{"ownAttrNm": "C_REF", "refAttrNm": "C_A"}]})
        assert code(r) == 200, r
        assert code(save_attrs(su, dm, "X", "T1", [add_col("ZZ_SU", "남의대기")])) == 200
    step("S3-0. 준비 — 소유자 X·Y 의 같은 이름 테이블 T1, 컬럼 없는 테이블, FK, 일반 사용자 승인 대기 컬럼", _setup)

    def _ddl():
        t = sa.get(BASE + "/api/dm/downloadDdl", params={"dataModelId": dm, "dbType": "Oracle"}, timeout=30).text
        blocks = [blk for blk in t.split("CREATE TABLE ")[1:] if blk.startswith("T1")]
        cols = [sorted(c for c in ("C_A", "C_B", "C_C", "C_REF", "ZZ_SU") if c + " " in blk) for blk in blocks]
        tsu = su.get(BASE + "/api/dm/downloadDdl", params={"dataModelId": dm, "dbType": "Oracle"}, timeout=30).text
        print(f"  관리자 DDL 의 T1 블록 컬럼={cols} / 일반 사용자 DDL 에 자기 대기 컬럼={'ZZ_SU' in tsu}")
        assert sorted(cols) == [["C_A", "C_B"], ["C_C", "C_REF"]] and "ZZ_SU" in tsu
    step("S3-1. DDL 내려받기 — (소유자, 테이블)별로 묶고 남의 승인 대기 컬럼은 빼며 본인 것은 담는다 (4.12-15-a·11)", _ddl)

    def _excel():
        rt = sa.get(BASE + "/api/dm/downloadDataModelObjs", params={"clctId": dm}, timeout=30)
        wt = openpyxl.load_workbook(io.BytesIO(rt.content)).active
        hdr_t = [c.value for c in wt[1]][:6]
        ra = sa.get(BASE + "/api/dm/downloadDataModelAttrs", params={"clctId": dm}, timeout=30)
        wa = openpyxl.load_workbook(io.BytesIO(ra.content)).active
        rows = [[c.value for c in r] for r in wa.iter_rows(min_row=2)]
        names = sorted(r[3] for r in rows)
        ref = next(r for r in rows if r[3] == "C_REF")
        print(f"  테이블 머리글={hdr_t}\n  컬럼 행={names} / C_REF 참조 테이블·컬럼(한글)={ref[13]}·{ref[14]}")
        assert hdr_t == ["소유자", "업무영역", "테이블명(한글)", "테이블명(영문)", "테이블스페이스", "설명"]
        assert "ZZ_SU" not in names and ref[13] == "엑스테이블" and ref[14] == "에이"
    step("S3-2. 엑셀 내려받기 — 테이블은 업로드 양식 열(업무영역·테이블스페이스), 컬럼 참조는 한글명, 남의 대기 행 제외 (4.12-22·11)", _excel)

    def _xmi():
        x = sa.get(BASE + "/api/dm/exportXmi", params={"dataModelId": dm}, timeout=30).text
        r = sa.post(BASE + "/api/dm/parseXmi", files={"file": ("e.xmi", x.encode("utf-8"), "application/xml")}, timeout=30).json()
        tb = sorted(t["objNm"] + ":" + t["objNmKr"] for t in r["tables"])
        cb = next(c for c in r["columns"] if c["attrNm"] == "C_B")
        print(f"  XMI 테이블={tb} / 남의 대기 컬럼={'ZZ_SU' in x} / 다시 읽은 C_B={cb['attrNmKr']} {cb['dataType']}({cb['dataLen']})")
        assert "T_EMPTY:빈테이블" in tb and "T1:엑스테이블" in tb and "ZZ_SU" not in x
        assert cb["attrNmKr"] == "비" and cb["dataType"] == "VARCHAR" and cb["dataLen"] == 20
    step("S3-3. XMI 추출 — 컬럼 없는 테이블 포함, 한글명·타입·길이를 실어 다시 가져오면 그대로 (4.12-34·11)", _xmi)


# ---------------------------------------------------------------- S4
def s4(sa, su):
    dm = FX["m4"]

    def xlsx(rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["소유자", "테이블명(영문)", "테이블명(한글)", "컬럼명(영문)", "컬럼명(한글)", "데이터타입", "길이", "소수점자리", "컬럼 순서",
                   "NULL여부", "PK여부", "FK여부", "디폴트값", "참조 테이블(한글)", "참조 컬럼(한글)", "삭제 규칙"])
        for r in rows:
            ws.append(r)
        p = os.path.join(TMP, f"a{len(os.listdir(TMP))}.xlsx")
        wb.save(p)
        return p

    def up(path, mode):
        with open(path, "rb") as f:
            r = sa.post(BASE + "/api/dm/uploadAttrs", files={"file": ("attrs.xlsx", f)}, data={"dataModelId": dm, "mode": mode}, timeout=60)
        return r.json()

    good = [["X", "T1", None, "C_NEW1", "새컬럼1", "VARCHAR", 10, None, None, "Y", "N", "N"],
            ["Y", "T1", None, "C_NEW2", "새컬럼2", "VARCHAR", 10, None, 7, "Y", "N", "N"],
            ["Y", "T1", None, "C_FK", "에프케이", "VARCHAR", 10, None, None, "Y", "N", "Y", None, "와이테이블", "씨", "CASCADE"]]

    def _preview_errors():
        bad = good + [["X", "T1", None, "C_DUPORD", "순서겹침", "VARCHAR", 10, None, 1, "Y", "N", "N"],
                      ["Y", "T1", None, "C_FK2", "에프케이2", "VARCHAR", 10, None, None, "Y", "N", "Y", None, "와이테이블", "없는컬럼"]]
        r = up(xlsx(bad), "preview")
        errs = [e["message"] for e in contents(r)["errors"]]
        print(f"  미리보기 오류={errs}")
        assert any("순서 1 는 이미 'C_A' 가 사용 중" in e for e in errs) and any("참조 컬럼이 모델·파일에 없습니다" in e for e in errs) and len(errs) == 2
    step("S4-1. 컬럼 업로드 미리보기 — 순서 겹침·FK 참조 해석 실패를 오류로 (4.12-21-b·21-c)", _preview_errors)

    def _commit():
        r = up(xlsx(good), "commit")
        assert code(r) == 200 and not contents(r)["errors"], r
        o1 = attr(dm, "X", "T1", "C_NEW1", "attr_ord")
        o2 = attr(dm, "Y", "T1", "C_NEW2", "attr_ord")
        fk = attr(dm, "Y", "T1", "C_FK", "fk_yn||'|'||fk_parent_obj_owner||'.'||fk_parent_obj_nm||'.'||fk_parent_attr_nm")
        cons = psql(f"select string_agg(constraint_type||':'||column_nm||'>'||ref_owner||'.'||ref_table_nm||'.'||ref_column_nm||':'||coalesce(delete_rule,''), ',') "
                    f"from quality.tb_data_model_constraint where dm_id='{dm}' and obj_owner='Y' and column_nm='C_FK'")
        print(f"  X.T1.C_NEW1 순서={o1} (기존 C_A·C_B·남의 ZZ_SU 다음) / Y.T1.C_NEW2 순서={o2} / C_FK={fk} / 제약조건={cons}")
        assert o1 == "4" and o2 == "7" and fk == "Y|Y.T1.C_C" and cons == "R:C_FK>Y.T1.C_C:CASCADE"
    step("S4-2. 컬럼 업로드 등록 — 소유자별 다음 순서, 파일의 컬럼 순서, FK 제약조건·부모 저장 (4.12-21-a·b·c)", _commit)

    def _import_owner():
        dm5 = make_model(sa, "ZZ110가져오기모델_" + TAG)
        assert code(add_obj(sa, dm5, "OTHER", "ZZ_IMP", "다른소유자")) == 200
        imp = {"dataModelId": dm5, "tables": [{"objNm": "ZZ_IMP", "objNmKr": "가져오기", "objAttrCnt": 2},
                                              {"objNm": "ZZ_PAR", "objNmKr": "부모", "objAttrCnt": 1}],
               "columns": [{"objNm": "ZZ_IMP", "attrNm": "ZZ_A", "attrNmKr": "에이", "dataType": "VARCHAR", "dataLen": 10, "nullableYn": "Y",
                            "pkYn": "N", "fkYn": "N", "attrOrder": 1},
                           {"objNm": "ZZ_IMP", "attrNm": "ZZ_P", "attrNmKr": "부모키", "dataType": "VARCHAR", "dataLen": 10, "nullableYn": "Y",
                            "pkYn": "N", "fkYn": "Y", "fkParentObjNm": "ZZ_PAR", "fkParentAttrNm": "ZZ_K", "attrOrder": 2},
                           {"objNm": "ZZ_PAR", "attrNm": "ZZ_K", "attrNmKr": "키", "dataType": "VARCHAR", "dataLen": 10, "nullableYn": "N",
                            "pkYn": "Y", "fkYn": "N", "attrOrder": 1}]}
        r = post(sa, "/api/dm/importXmiModel", imp)
        owners = psql(f"select string_agg(obj_owner||'.'||obj_nm||'.'||attr_nm||':'||coalesce(fk_parent_obj_owner,'-'), ',' order by attr_nm) "
                      f"from quality.tb_data_model_attr where dm_id='{dm5}'")
        clct = psql(f"select obj_cnt||'|'||attr_cnt from quality.tb_data_model_clct where dm_id='{dm5}'")
        print(f"  XMI 가져오기={r.get('success')} / 컬럼 소유자·FK 부모 소유자={owners} / 가져오기 이력 건수={clct}")
        assert r.get("success") is True and owners == "XMI.ZZ_IMP.ZZ_A:-,XMI.ZZ_PAR.ZZ_K:-,XMI.ZZ_IMP.ZZ_P:XMI" and clct == "2|3"
    step("S4-3. XMI 가져오기 — 컬럼은 테이블과 같은 소유자(XMI, 다른 소유자의 같은 이름 테이블에 안 붙음), FK 부모 소유자 XMI (4.12-33-a·b·14-a)", _import_owner)


# ---------------------------------------------------------------- S5
def s5(sa, su):
    dm = make_model(sa, "ZZ110승인모델_" + TAG)
    FX["m6"] = dm

    def _setup():
        assert code(add_obj(sa, dm, "", "P", "부모")) == 200
        assert code(add_obj(sa, dm, "", "C", "자식")) == 200
        assert code(save_attrs(sa, dm, "", "P", [add_col("P_ID", "아이디", nullableYn="N"), add_col("P_NM", "이름")])) == 200
        assert code(save_attrs(sa, dm, "", "C", [add_col("C_ID", "아이디"), add_col("C_PID", "부모아이디")])) == 200
        assert code(post(sa, "/api/dm/createPk", {"dataModelId": dm, "objOwner": "", "tableNm": "P", "pkName": "PK_P", "columns": ["P_ID"]})) == 200
        r = post(sa, "/api/dm/createFk", {"dataModelId": dm, "objOwner": "", "tableNm": "C", "refOwner": "", "refTableNm": "P",
                                          "fkName": "FK_C_P", "mappings": [{"ownAttrNm": "C_PID", "refAttrNm": "P_ID"}]})
        assert code(r) == 200, r
    step("S5-0. 준비 — 관리자 승인 테이블 P(PK)·C(FK→P)", _setup)

    def _modify_reject():
        r = save_attrs(su, dm, "", "P", [{"mode": "UPDATE", "origAttrNm": "P_NM", "attrNm": "P_NM", "attrNmKr": "이름", "dataType": "VARCHAR",
                                          "dataLen": 20, "nullableYn": "Y", "pkYn": "N", "fkYn": "N"}])
        seq = seq_of(dm, "MODIFY_ATTR", "P", "P_NM", "DRAFT")
        assert code(r) == 200 and code(post(su, "/api/dmApproval/submit", {"changeSeqList": [seq]})) == 200
        r = post(sa, "/api/dmApproval/reject", {"changeSeqList": [seq], "aprvComment": "시험"})
        st = attr(dm, "", "P", "P_NM")
        seen = any(x["objNm"] == "P" and x["attrNm"] == "P_NM" for x in get(sa, "/api/dm/getDataModelAttrListByClctId", {"clctId": dm}))
        print(f"  길이만 바꾼 수정 반려={code(r)} / 행={st} / 관리자 컬럼 목록에 보임={seen}")
        assert code(r) == 200 and st == "APPROVED|Y|10" and seen
    step("S5-1. 타입·길이만 바꾼 컬럼 수정 반려 — 값만 되돌리고 승인 컬럼 그대로 (5.8-13)", _modify_reject)

    def _rename_reject():
        r = save_attrs(su, dm, "", "P", [{"mode": "UPDATE", "origAttrNm": "P_ID", "attrNm": "P_KEY", "attrNmKr": "아이디", "dataType": "VARCHAR",
                                          "dataLen": 10, "nullableYn": "N", "pkYn": "Y", "fkYn": "N"}])
        assert code(r) == 200, r
        mid = psql(f"select (select string_agg(column_nm, ',') from quality.tb_data_model_constraint where dm_id='{dm}' and constraint_nm='PK_P')||'|'||"
                   f"(select string_agg(ref_column_nm, ',') from quality.tb_data_model_constraint where dm_id='{dm}' and constraint_nm='FK_C_P')")
        seq = seq_of(dm, "MODIFY_ATTR", "P", "P_KEY", "DRAFT")
        r = post(sa, "/api/dmApproval/reject", {"changeSeqList": [seq], "aprvComment": "시험"})
        refs = psql(f"select (select string_agg(attr_nm||':'||aprv_status, ',' order by attr_nm) from quality.tb_data_model_attr where dm_id='{dm}' and obj_nm='P')||'|'||"
                    f"(select string_agg(column_nm, ',') from quality.tb_data_model_constraint where dm_id='{dm}' and constraint_nm='PK_P')||'|'||"
                    f"(select string_agg(column_nm, ',') from quality.tb_data_model_index where dm_id='{dm}' and index_nm='PK_P')||'|'||"
                    f"(select string_agg(ref_column_nm, ',') from quality.tb_data_model_constraint where dm_id='{dm}' and constraint_nm='FK_C_P')||'|'||"
                    f"(select fk_parent_attr_nm from quality.tb_data_model_attr where dm_id='{dm}' and obj_nm='C' and attr_nm='C_PID')")
        print(f"  이름 변경 뒤 PK·FK 참조={mid} / 반려={code(r)} → 컬럼|PK|PK 인덱스|FK 참조|FK 부모={refs}")
        assert mid == "P_KEY|P_KEY" and code(r) == 200 and refs == "P_ID:APPROVED,P_NM:APPROVED|P_ID|P_ID|P_ID|P_ID"
    step("S5-2. 컬럼 영문명 변경 반려 — 인덱스·제약조건·참조 FK·FK 부모 이름도 되돌림 (5.8-14)", _rename_reject)

    def _del_reject():
        assert code(save_attrs(su, dm, "", "C", [add_col("C_NEW", "새컬럼")])) == 200
        assert code(post(su, "/api/dm/deleteObj", {"dataModelId": dm, "objOwner": "", "objNm": "C"})) == 200
        seq = seq_of(dm, "DEL_OBJ", "C", status="DRAFT")
        r = post(sa, "/api/dmApproval/reject", {"changeSeqList": [seq], "aprvComment": "시험"})
        st = attr(dm, "", "C", "C_NEW", "aprv_status||'|'||use_yn")
        ob = psql(f"select aprv_status||'|'||use_yn from quality.tb_data_model_obj where dm_id='{dm}' and obj_nm='C'")
        cnt = psql(f"select obj_attr_cnt from quality.tb_data_model_obj where dm_id='{dm}' and obj_nm='C'")
        print(f"  테이블 삭제 반려={code(r)} / 테이블={ob} / 추가 신청 중이던 C_NEW={st} / 컬럼 수={cnt}")
        assert code(r) == 200 and ob == "APPROVED|Y" and st == "DRAFT|Y" and cnt == "3"
        FX["cNew"] = seq_of(dm, "ADD_ATTR", "C", "C_NEW", "DRAFT")
    step("S5-3. 대기 삭제 반려 복구는 USE_YN 만 — 추가 신청 중인 컬럼은 DRAFT 그대로, 컬럼 수 다시 셈 (5.8-15·20)", _del_reject)

    def _cascade_revert():
        assert code(save_attrs(su, dm, "", "P", [add_col("P_X", "추가부모")])) == 200
        r = post(su, "/api/dm/createFk", {"dataModelId": dm, "objOwner": "", "tableNm": "C", "refOwner": "", "refTableNm": "P",
                                          "fkName": "FK_C_X", "mappings": [{"ownAttrNm": "C_ID", "refAttrNm": "P_X"}]})
        assert code(r) == 200, r
        assert code(save_attrs(su, dm, "", "C", [{"mode": "UPDATE", "origAttrNm": "C_ID", "attrNm": "C_ID", "attrNmKr": "아이디", "dataType": "VARCHAR",
                                                  "dataLen": 30, "nullableYn": "Y", "pkYn": "N", "fkYn": "Y"}])) == 200
        mod = seq_of(dm, "MODIFY_ATTR", "C", "C_ID", "DRAFT")
        add = seq_of(dm, "ADD_ATTR", "P", "P_X", "DRAFT")
        r = post(sa, "/api/dmApproval/reject", {"changeSeqList": [add], "aprvComment": "시험"})
        st = attr(dm, "", "C", "C_ID", "aprv_status||'|'||data_len")
        print(f"  부모 컬럼 추가 반려={code(r)} / 전파로 반려된 자식 컬럼 수정 이력={hist_status(mod)} / 자식 컬럼={st}")
        assert code(r) == 200 and hist_status(mod) == "REJECTED" and st == "APPROVED|10"
    step("S5-4. 전파로 반려된 컬럼 수정 이력도 값을 되돌림 (5.8-16)", _cascade_revert)

    def _rollback():
        assert code(save_attrs(su, dm, "", "P", [{"mode": "UPDATE", "origAttrNm": "P_NM", "attrNm": "P_NM", "attrNmKr": "이름", "dataType": "VARCHAR",
                                                  "dataLen": 30, "nullableYn": "Y", "pkYn": "N", "fkYn": "N"}])) == 200
        assert code(save_attrs(su, dm, "", "P", [add_col("P_Y", "롤백추가")])) == 200
        assert code(save_attrs(su, dm, "", "C", [{"mode": "DELETE", "attrNm": "C_PID"}])) == 200
        seqs = [seq_of(dm, "MODIFY_ATTR", "P", "P_NM", "DRAFT"), seq_of(dm, "ADD_ATTR", "P", "P_Y", "DRAFT"),
                seq_of(dm, "DEL_ATTR", "C", "C_PID", "DRAFT"), FX["cNew"]]
        mid = attr(dm, "", "P", "P_NM") + "/" + attr(dm, "", "C", "C_PID")
        r = post(su, "/api/dmApproval/rollbackDraft", {"changeSeqList": seqs})
        after = (attr(dm, "", "P", "P_NM") + "/" + (attr(dm, "", "P", "P_Y") or "(없음)") + "/" + attr(dm, "", "C", "C_PID")
                 + "/" + (attr(dm, "", "C", "C_NEW") or "(없음)"))
        left = psql(f"select count(*) from quality.tb_data_model_change_history where change_seq in ({','.join(map(str, seqs))})")
        cnt = psql(f"select string_agg(obj_nm||':'||obj_attr_cnt, ',' order by obj_nm) from quality.tb_data_model_obj where dm_id='{dm}'")
        print(f"  롤백 전 P_NM/C_PID={mid}\n  롤백={code(r)} {r.get('contents')} → P_NM/P_Y/C_PID/C_NEW={after} / 남은 이력={left} / 컬럼 수={cnt}")
        # P 는 P_ID·P_NM (S5-4 에서 반려된 P_X 는 세지 않는다)
        assert code(r) == 200 and after == "APPROVED|Y|10/(없음)/APPROVED|Y|10/(없음)" and left == "0" and cnt == "C:2,P:2"
    step("S5-5. DRAFT 롤백 — 수정 값·대기 삭제·추가 행을 되돌리고 이력을 지움, 컬럼 수 다시 셈 (5.8-17·20)", _rollback)

    def _rollback_guard():
        assert code(add_obj(su, dm, "", "Q", "롤백테이블")) == 200
        assert code(save_attrs(su, dm, "", "Q", [add_col("Q1", "큐", nullableYn="N")])) == 200
        assert code(post(su, "/api/dm/createPk", {"dataModelId": dm, "objOwner": "", "tableNm": "Q", "pkName": "PK_Q", "columns": ["Q1"]})) == 200
        s_obj, s_col, s_pk = seq_of(dm, "ADD_OBJ", "Q"), seq_of(dm, "ADD_ATTR", "Q", "Q1"), seq_of(dm, "ADD_PK", "Q")
        r1 = post(su, "/api/dmApproval/rollbackDraft", {"changeSeq": s_obj})
        r0 = post(sa, "/api/dmApproval/rollbackDraft", {"changeSeq": s_pk})
        r2 = post(su, "/api/dmApproval/rollbackDraft", {"changeSeqList": [s_obj, s_col, s_pk]})
        left = psql(f"select (select count(*) from quality.tb_data_model_obj where dm_id='{dm}' and obj_nm='Q')||'/'||"
                    f"(select count(*) from quality.tb_data_model_attr where dm_id='{dm}' and obj_nm='Q')||'/'||"
                    f"(select count(*) from quality.tb_data_model_constraint where dm_id='{dm}' and table_nm='Q')||'/'||"
                    f"(select count(*) from quality.tb_data_model_change_history where dm_id='{dm}' and obj_nm='Q')")
        print(f"  테이블 추가만 롤백={code(r1)} {msg(r1)}\n  남(관리자)의 롤백={code(r0)} {msg(r0)}\n  셋 함께={code(r2)} → 테이블/컬럼/제약/이력={left}")
        assert code(r1) == 500 and "그 건부터 롤백하세요" in msg(r1) and code(r0) == 500 and "본인의 DRAFT 만" in msg(r0)
        assert code(r2) == 200 and left == "0/0/0/0"
    step("S5-6. 롤백 — 같은 테이블에 고르지 않은 대기 PK 가 있으면 거부, 남의 DRAFT 거부, 함께 고르면 테이블째 되돌림 (5.8-17)", _rollback_guard)

    def _split():
        dm7 = make_model(sa, "ZZ110승인모델2_" + TAG)
        assert code(add_obj(sa, dm7, "", "R", "다른모델")) == 200
        assert code(save_attrs(su, dm, "", "P", [add_col("P_S1", "신청1")])) == 200
        assert code(save_attrs(su, dm7, "", "R", [add_col("R_S1", "신청2")])) == 200
        seqs = [seq_of(dm, "ADD_ATTR", "P", "P_S1"), seq_of(dm7, "ADD_ATTR", "R", "R_S1")]
        r = post(su, "/api/dmApproval/submit", {"changeSeqList": seqs})
        c = contents(r)
        subs = psql(f"select count(distinct submission_id)||'/'||count(distinct dm_id) from quality.tb_data_model_change_history where change_seq in ({seqs[0]},{seqs[1]})")
        cards = [s for s in post(sa, "/api/dmApproval/submissions", {}) if s["submissionId"] in c.get("submissionIds", [])]
        print(f"  두 모델 한 번에 신청={code(r)} submitted={c.get('submitted')} 신청 수={len(c.get('submissionIds', []))} / 신청·모델={subs}"
              f"\n  승인 화면 카드 모델={[x['dmNm'] for x in cards]} oldestStatus 키={any('oldestStatus' in x for x in cards)}")
        assert code(r) == 200 and c["submitted"] == 2 and subs == "2/2" and sorted(x["dmNm"] for x in cards) == sorted(["ZZ110승인모델_" + TAG, "ZZ110승인모델2_" + TAG])
        assert not any("oldestStatus" in x for x in cards)
        h = post(su, "/api/dmApproval/history", {"dmId": "", "changeUserId": "jyjang", "aprvStatus": "SUBMITTED"})
        print(f"  내 변경 신청 SUBMITTED 조회 {len(h)}건, 모두 SUBMITTED={all(x['aprvStatus'] == 'SUBMITTED' for x in h)}")
        assert len(h) >= 2 and all(x["aprvStatus"] == "SUBMITTED" for x in h)
        FX["subSeqs"] = seqs
    step("S5-7. 여러 모델을 한 번에 신청하면 모델마다 신청이 나뉘고 카드에 각 모델명 (5.8-27·6·26-b)", _split)


# ---------------------------------------------------------------- S6
def s6(sa, su):
    dm = FX["m6"]

    def _stats():
        assert code(add_obj(su, dm, "", "ZZ_SUOBJ", "남의대기테이블")) == 200
        st_a = get(sa, "/api/dm/diagTargetStats", {"dmId": dm})
        st_u = get(su, "/api/dm/diagTargetStats", {"dmId": dm})
        lst_a = len(get(sa, "/api/dm/getDataModelObjListByClctId", {"clctId": dm}))
        lst_u = len(get(su, "/api/dm/getDataModelObjListByClctId", {"clctId": dm}))
        print(f"  관리자 통계 테이블={st_a.get('totalobj')}(목록 {lst_a}) / 일반 사용자 통계 테이블={st_u.get('totalobj')}(목록 {lst_u}) / 품질 키={'objqualoff' in st_a}")
        assert st_a["totalobj"] == lst_a and st_u["totalobj"] == lst_u and st_u["totalobj"] == st_a["totalobj"] + 1 and "objqualoff" not in st_a
    step("S6-1. 진단 제외 통계가 목록과 같은 행만 센다 (5.8-30·9)", _stats)

    def _error_msg():
        r1 = su.post(BASE + "/api/dm/setObjDiagTarget", json={"dmId": dm, "objOwner": "", "objNm": "P", "diagType": "STND", "targetYn": "N"}, timeout=15)
        r2 = sa.post(BASE + "/api/dm/setObjDiagTarget", json={"dmId": dm, "objOwner": "", "objNm": "P", "diagType": "XX", "targetYn": "N"}, timeout=15)
        print(f"  권한 없음 → {r1.status_code} {r1.json()} / 잘못된 유형 → {r2.status_code} {r2.json()}")
        assert r1.status_code == 403 and r1.json()["message"].startswith("진단 대상을 변경할 권한이 없습니다")
        assert r2.status_code == 400 and r2.json()["message"] == "diagType 잘못됨: XX"
    step("S6-2. 진단 제외 변경 거부 사유가 응답 본문 message 로 (5.8-32)", _error_msg)

    def _biz():
        r = post(sa, "/api/area/biz/save", {"bizAreaNm": "ZZ110업무_" + TAG, "bizAreaDesc": "110 시험", "sortOrder": 0})
        bid = contents(r)["bizAreaId"]
        cleanup.append(lambda: psql(f"delete from quality.tb_biz_area where biz_area_id='{bid}'"))
        r = post(sa, "/api/dm/updateObj", {"dataModelId": dm, "origObjNm": "P", "objNm": "P", "objOwner": "", "origObjOwner": "",
                                           "objNmKr": "부모", "bizAreaId": bid})
        assert code(r) == 200, r
        nm1 = next(o for o in get(sa, "/api/dm/getDataModelObjListByClctId", {"clctId": dm}) if o["objNm"] == "P").get("bizAreaNm")
        r2 = post(su, "/api/area/biz/delete", {"bizAreaId": bid})
        assert code(post(sa, "/api/area/biz/delete", {"bizAreaId": bid})) == 200
        nm2 = next(o for o in get(sa, "/api/dm/getDataModelObjListByClctId", {"clctId": dm}) if o["objNm"] == "P").get("bizAreaNm")
        print(f"  지정 뒤 업무영역명={nm1} / 일반 사용자 삭제={code(r2)} {msg(r2)} / 삭제 뒤 업무영역명={nm2}")
        assert nm1 == "ZZ110업무_" + TAG and code(r2) == 500 and nm2 is None
    step("S6-3. 지운 업무영역 이름은 테이블 목록에 안 보임 (5.8-31-b, 화면 실패 안내용 응답 5.8-11-c)", _biz)


# ---------------------------------------------------------------- S7
def s7(sa, su):
    from selenium import webdriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from edgeopt import headless

    opts = webdriver.EdgeOptions()
    opts.add_argument("--log-level=3")
    drv = webdriver.Edge(options=headless(opts))
    drv.set_window_size(1600, 1000)
    W = lambda t=15: WebDriverWait(drv, t)

    def js(s, *a):
        return drv.execute_script(s, *a)

    def nav(nav_id, group_id=None):
        if group_id and not drv.find_elements(By.ID, nav_id):
            hdr = drv.find_element(By.CSS_SELECTOR, f"#{group_id} .v-list-group__header")
            js("arguments[0].click();", hdr)
            time.sleep(1)
        el = W().until(EC.presence_of_element_located((By.ID, nav_id)))
        js("arguments[0].scrollIntoView({block:'center'}); arguments[0].click();", el)
        time.sleep(2)

    def swal_text(t=10, has=""):
        def _t(x):
            for e in x.find_elements(By.CSS_SELECTOR, ".swal2-popup"):
                if e.is_displayed() and has in e.text:
                    return e.text
            return None
        return W(t).until(_t)

    def swal_click(sel=".swal2-confirm"):
        for b in drv.find_elements(By.CSS_SELECTOR, sel):
            if b.is_displayed():
                b.click()
                time.sleep(1)
                return

    def login(user):
        drv.delete_all_cookies()
        drv.get(BASE)
        W().until(EC.presence_of_element_located((By.CSS_SELECTOR, "input[type='text']"))).send_keys(user)
        drv.find_element(By.CSS_SELECTOR, "input[type='password']").send_keys(PW)
        drv.find_element(By.CSS_SELECTOR, "button[type='submit'], .v-btn").click()
        W().until(lambda x: "/app/main" in x.current_url)
        time.sleep(3)

    try:
        login("space")

        def _approval():
            seqs = FX["subSeqs"]
            # 첫 모델 신청에 한 건을 더 넣고 그중 한 건을 반려 — 처리된 건이 섞인 신청
            dm = FX["m6"]
            sub = psql(f"select submission_id from quality.tb_data_model_change_history where change_seq={seqs[0]}")
            assert code(save_attrs(su, dm, "", "P", [add_col("P_S2", "신청3")])) == 200
            s2 = seq_of(dm, "ADD_ATTR", "P", "P_S2")
            psql(f"update quality.tb_data_model_change_history set aprv_status='SUBMITTED', submission_id='{sub}' where change_seq={s2}")
            assert code(post(sa, "/api/dmApproval/reject", {"changeSeqList": [seqs[0]], "aprvComment": "시험"})) == 200
            nav("nav_dm_approval", "mmGroup")
            card = W().until(EC.presence_of_element_located((By.XPATH, f"//div[contains(@class,'submission-row')][contains(., 'ZZ110승인모델_{TAG}')]")))
            card.click()
            time.sleep(2)
            rows = drv.find_elements(By.CSS_SELECTOR, "#dmAprvItems tbody tr")
            info = []
            for r in rows:
                cb = r.find_elements(By.CSS_SELECTOR, ".v-simple-checkbox")
                if not cb:
                    continue
                dis = "v-simple-checkbox--disabled" in cb[0].get_attribute("class")
                info.append(("REJECTED" in r.text, dis))
            # 전체 선택 → 처리 안 된 건만
            js("arguments[0].click();", drv.find_element(By.XPATH, "//span[normalize-space()='전체 선택']/preceding::input[@type='checkbox'][1]"))
            time.sleep(1)
            btn = drv.find_element(By.XPATH, "//button[contains(., '선택 승인')]").text
            print(f"  행(반려됨, 체크 비활성)={info} / 전체 선택 뒤 버튼={btn!r}")
            assert (True, True) in info and (False, False) in info and "(1)" in btn
        step("S7-1. 데이터 모델 변경 승인 — 처리된 건은 체크 불가, 전체 선택도 대기 건만 (5.8-12-a)", _approval)

        def _diag_target():
            nav("nav_diagTargetMgmt", "dmGroup")
            t = drv.find_element(By.ID, "tab_diagTargetMgmt").text
            print(f"  진단 제외 관리 화면에 '품질' 글자={'품질' in t}")
            assert "품질" not in t and "구조대상" in t
        step("S7-2. 진단 제외 관리 — 품질 진단 칩·버튼·열 없음 (5.8-9)", _diag_target)

        def _history():
            nav("nav_datamodelHistory", "dmGroup")
            box = W().until(EC.presence_of_element_located((By.CSS_SELECTOR, "#tab_datamodelHistory input[placeholder='모델 선택']")))
            box.click()
            box.send_keys(FX["m1Nm"])
            time.sleep(1)
            opt = W().until(EC.presence_of_element_located((By.XPATH, f"//div[contains(@class,'v-list-item__title')][normalize-space()='{FX['m1Nm']}']")))
            opt.click()
            drv.find_element(By.XPATH, "//div[@id='tab_datamodelHistory']//button[contains(., '조회')]").click()
            time.sleep(2)
            body = drv.find_element(By.ID, "dmHistory_table").text
            print(f"  수집이력 표:\n{body}")
            assert FX["m1Nm"] + "X" not in body and FX["m1Nm"] in body
        step("S7-3. 데이터 모델 수집이력 — 고른 모델만 (이름이 포함 관계인 다른 모델 이력 없음) (4.12-14-b)", _history)

        def _collection():
            nav("nav_datamodelCollection", "dmGroup")
            cell = W().until(EC.presence_of_element_located((By.XPATH, f"//div[@id='tab_datamodelCollection']//span[normalize-space()='{FX['m1Nm']}']")))
            js("arguments[0].scrollIntoView({block:'center'}); arguments[0].click();", cell)
            time.sleep(3)
            # 수정 창 — 이름을 공백으로
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[@id='tab_datamodelCollection']//button[normalize-space()='수정']"))
            time.sleep(2)
            nm = W().until(EC.visibility_of_element_located((By.CSS_SELECTOR, "input[name='update_dataModelNm']")))
            nm.send_keys(Keys.CONTROL, "a")
            nm.send_keys(Keys.DELETE)
            nm.send_keys("   ")
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[contains(@class,'v-dialog--active')]//button[normalize-space()='수정']"))
            t1 = swal_text(has="필수")
            swal_click()
            js("document.querySelectorAll('.v-dialog--active button').forEach(b => { if (b.innerText.trim() === '취소') b.click(); });")
            time.sleep(1)
            # 수집 확인창 → 진행 창 완료 (RELOAD 가 요청한 창으로 온다)
            js("arguments[0].click();", drv.find_element(By.XPATH, "//div[@id='tab_datamodelCollection']//button[contains(., '수집')]"))
            t2 = swal_text(has="수집 대상 스키마")
            swal_click()
            W(90).until(lambda x: "수집이 완료되었습니다" in x.page_source or "수집에 실패했습니다" in x.page_source)
            done = "수집이 완료되었습니다" in drv.page_source
            print(f"  이름 공백 수정 → {t1.splitlines()[0]!r}\n  수집 확인창 한글명 안내={'한글명은 바꾸지 않습니다' in t2} / 진행 창 완료={done}")
            assert "데이터 모델명은 필수 입력값입니다" in t1 and "한글명은 바꾸지 않습니다" in t2 and done
        step("S7-4. 데이터 모델 관리 — 이름 공백 수정 거부, 수집 확인창 문구, 진행 창이 요청한 창의 RELOAD 로 완료 (4.12-13-a·41-a·37-b)", _collection)

        def _column():
            login("jyjang")
            nav("nav_datamodelStatusColumn", "dmGroup")
            tab = "#tab_datamodelStatusColumn"
            box = W().until(EC.presence_of_element_located((By.CSS_SELECTOR, f"{tab} input[placeholder='모델 선택']")))
            mnm = "ZZ110내려받기모델_" + TAG
            box.click()
            box.send_keys(mnm)
            W().until(EC.presence_of_element_located((By.XPATH, f"//div[contains(@class,'v-list-item__title')][normalize-space()='{mnm}']"))).click()
            time.sleep(3)
            # 영문명·한글명은 입력칸 값이다
            draft = [[i.get_attribute("value") for i in r.find_elements(By.CSS_SELECTOR, "input") if i.get_attribute("value")]
                     for r in drv.find_elements(By.CSS_SELECTOR, f"{tab} tr.row-draft")]
            draft_txt = " ".join(" ".join(v) for v in draft)
            # 업로드 미리보기 — 오류 사유
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["소유자", "테이블명(영문)", "테이블명(한글)", "컬럼명(영문)", "컬럼명(한글)", "데이터타입", "길이", "소수점자리", "컬럼 순서"])
            ws.append(["X", "T1", None, "C_SCR", "화면시험", "VARCHAR", 10, None, 1])
            p = os.path.join(TMP, "screen.xlsx")
            wb.save(p)
            drv.find_element(By.CSS_SELECTOR, f"{tab} input[type='file']").send_keys(p)
            W().until(lambda x: "컬럼 엑셀 업로드 미리보기" in x.page_source)
            time.sleep(1)
            dlg = next(e.text for e in drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active") if "미리보기" in e.text)
            js("document.querySelectorAll('.v-dialog--active button').forEach(b => { if (b.innerText.trim() === '취소') b.click(); });")
            time.sleep(1)
            # 붙여넣기 — 다른 탭에 있을 때는 이 화면에 행이 생기지 않는다
            tgt = drv.find_element(By.CSS_SELECTOR, f"{tab} input[placeholder='테이블 선택 (소유자.테이블)']")
            tgt.click()
            time.sleep(1)
            opts_ = [e for e in drv.find_elements(By.CSS_SELECTOR, ".menuable__content__active .v-list-item") if "X.T1" in e.text]
            opts_[0].click()
            time.sleep(1)
            paste = ("const dt = new DataTransfer(); dt.setData('text/plain', arguments[0]);"
                     "document.body.dispatchEvent(new ClipboardEvent('paste', {clipboardData: dt, bubbles: true}));")
            nav("nav_datamodelStatusTable", "dmGroup")
            js(paste, "다른탭붙임1\n다른탭붙임2")
            nav("nav_datamodelStatusColumn", "dmGroup")
            time.sleep(1)
            other = drv.page_source.count("다른탭붙임")
            js(paste, "이탭붙임1")
            time.sleep(1)
            here = "이탭붙임1" in " ".join(e.get_attribute("value") or "" for e in drv.find_elements(By.CSS_SELECTOR, f"{tab} input"))
            print(f"  일반 사용자 승인 대기 행(row-draft)={draft} / 업로드 미리보기 사유 표시={'순서 1 는 이미' in dlg}"
                  f"\n  다른 탭에서 붙여넣기로 생긴 행={other} / 이 탭에서 붙여넣기 행 생김={here}")
            assert "ZZ_SU" in draft_txt and "순서 1 는 이미 'C_A'" in dlg and other == 0 and here
        step("S7-5. 컬럼 화면 — 미승인 행 색, 업로드 미리보기 오류 사유, 다른 탭에서 붙여넣기 무시 (4.12-25·20·29)", _column)
    finally:
        drv.quit()


def main(only):
    sa = login_api("space")
    su = login_api("jyjang")
    start = time.strftime("%Y%m%d%H%M%S")
    assert psql("select count(*) from quality.tb_aprv_line where use_yn='Y'") == "0", "켜진 결재선이 있다 — 시험이 모델 승인을 하므로 멈춘다"
    groups = [("S1", s1), ("S2", s2), ("S3", s3), ("S4", s4), ("S5", s5), ("S6", s6), ("S7", s7)]
    try:
        for nm, fn in groups:
            if not only or nm in only:
                fn(sa, su)
    finally:
        print("\n=== 정리")
        for fn in reversed(cleanup):
            try:
                fn()
            except Exception as e:
                print("  정리 실패:", e)
        psql(f"delete from quality.tb_notification where cret_dt >= '{start}'")
        left = psql(f"select (select count(*) from quality.tb_data_model where dm_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_data_model_attr a where not exists (select 1 from quality.tb_data_model m where m.dm_id=a.dm_id))||'/'||"
                    f"(select count(*) from quality.tb_data_model_change_history h where h.change_dt >= '{start}' "
                    f"and not exists (select 1 from quality.tb_data_model m where m.dm_id=h.dm_id))||'/'||"
                    f"(select count(*) from ndata.tb_data_source where dsn like 'ZZ110%')||'/'||"
                    f"(select count(*) from quality.tb_biz_area where biz_area_nm like 'ZZ110%')||'/'||"
                    f"(select count(*) from pg_namespace where nspname like 'zz110%')||'/'||"
                    f"(select count(*) from pg_roles where rolname like 'zz110%')")
        print(f"  남은 시험 데이터 모델/고아 컬럼/고아 이력/데이터 소스/업무영역/스키마/역할 = {left}")


if __name__ == "__main__":
    t0 = time.time()
    main(set(sys.argv[1:]))
    elapsed = time.time() - t0
    p = sum(1 for _, st in results if st == "PASS")
    f = sum(1 for _, st in results if st == "FAIL")
    print(f"\n{'=' * 60}\n결과: {p} PASS / {f} FAIL  ({elapsed:.0f}초)\n{'=' * 60}")
    for n, st in results:
        print(f"  [{st}] {n}")
    sys.exit(0 if f == 0 else 1)
