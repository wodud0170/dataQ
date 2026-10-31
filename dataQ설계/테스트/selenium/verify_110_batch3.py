# -*- coding: utf-8 -*-
"""
110번 기존 기능 보완 3묶음(표준화 진단·구조 진단·스케줄 6.10 / 용어 추천·변환 이력 7.5) + 114 §2 권고 5건 조치 검증 — API + 화면 (2026-10-31, PC1)

  S2. 진단 서버(q-executor) 내림·거부·재기동 — 6.10-5a·5b(진단 서버 연결 실패·거부 시 Job·구조 진단 이력 ERROR),
      6.10-7a·7b(즉시 실행 거부 시 LOG 마감, launcher 가 이미 마감한 LOG 는 덮지 않음), 6.10-6(기동 때 남은 READY·RUNNING 정리)
      ※ 진짜 q-executor 를 내리고 28098 에 시험용 가짜 서버를 띄웠다가 다시 진짜를 올린다.
  R.  예약 실행기(DiagSchedulerRunner) — 재기동 직후 심은 CRON 스케줄·오래된 LOG 를 실행기가 처리한 결과
      6.10-9(건너뛴 슬롯 소비·기동 때 RUNNING), 6.10-12a·12b(겹치는 진단이면 SKIPPED), 6.10-13a(자동 실행자=등록자),
      6.10-2(스케줄 사전), 6.10-10(수집 이력 없는 모델도 진단), 6.10-8a·8b(타임아웃 마감·기저 작업 진행 중이면 제외)
  S1. 표준화 진단·구조 진단 — 6.10-15(결과 상세는 Job 사전 계보만), 6.10-20a(수집일시), 6.10-23(진행 100컬럼마다),
      6.10-24b·25a·26·27c(구조 진단: 인덱스·제약 비교, 소유자 대소문자, 스키마 목록, 수집건 인자 없음), 6.10-27a·27b·29·31a·31b·7.5-3e(지운 API·화면)
  S3. 스케줄 API — 권한, 6.10-11a·11b(등록 검사), 6.10-2(사전), 6.10-12a(즉시 실행 409), 6.10-13b·9·10(즉시 실행자·RUNNING·수집 없는 모델)
  S4. 용어 추천·변환 이력 API — 7.5-6·7·8(일괄 등록), 7.5-9c, 7.5-11(엑셀 읽기 실패 400), 7.5-13c(한글명 변경 기준)
  S5. 114 §2 — 2-4(BINARY·VARBINARY 길이 필수), 2-2(남의 승인 대기 영문약어는 등록 때 거부), 2-5·결정 기록
  S6. 화면 — 셀레니움 Edge(창 없이): 진단 결과(6.10-14·15·17·18·19b·19c·19d·20b), 진단 실행(6.10-21·22a·22b), 구조 진단(6.10-28),
      스케줄 관리·실행 이력(6.10-2·11b·33), 한글컬럼 일괄 표준화(7.5-5b·9a·9b·10·11), 한글 변환 이력(7.5-13a·13b·13d),
      내 변경 신청(114 2-1), 데이터 모델 관리 수정 창(114 2-3), 도메인 길이(114 2-4), 콘솔 오류

필요:
  - q-center(28091)·q-executor(28098) 기동. S2 는 q-executor 를 내렸다가 `java -jar target/q-executor.jar` 로 다시 올린다
    (환경 변수 QEXEC_LOG = q-executor 표준 출력 파일, 기본 C:\\Temp\\dataq\\logs\\executor.out — 6.10-6·23 이 이 로그를 읽는다. DATAQ_TMP 기본 C:\\Temp\\dataq)
  - 환경 변수 DATAQ_TEST_PW (관리자 space 비밀번호). 일반 사용자·스케줄 관리자는 시험 안에서 만들고 지운다
  - 개발 DB 는 dbcmd.PSQL (PC1 dataq-db 컨테이너, PC2 는 DATAQ_PSQL=native + PGPASSWORD)
  - 인자: 돌릴 묶음(예: S1 S4 — 없으면 전부, R 은 S2 뒤에만 의미가 있다), --keep (정리를 건너뛰고 남긴다 — 화면 검수용, 뒤에 --purge 로 지운다)
시험 데이터(시험 사전과 단어·용어·도메인, 사용자·역할, 시험 스키마·역할, 데이터 소스, 시스템, 모델과 그 수집·진단·구조 진단, 스케줄과 LOG,
변환 이력, 일괄 등록 이력, 가짜 기저 작업)는 끝에서 지운다. 이름은 모두 ZZB3·zzb3 로 시작한다.
"""
import base64, json, os, re, secrets, subprocess, sys, tempfile, threading, time, traceback
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import requests
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dbcmd import PSQL  # noqa: E402

BASE = "http://localhost:28091"
EXEC = "http://127.0.0.1:28098"
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
DATAQ_TMP = os.environ.get("DATAQ_TMP", r"C:\Temp\dataq")
EXEC_LOG = os.environ.get("QEXEC_LOG", os.path.join(DATAQ_TMP, "logs", "executor.out"))
PW = os.environ["DATAQ_TEST_PW"]
TAG = time.strftime("%H%M%S")
TMP = tempfile.mkdtemp(prefix="v110b3_")
results = []
cleanup = []          # 끝에서 거꾸로 실행할 정리 함수
FX = {}               # 공통 시험 데이터
KEEP = "--keep" in sys.argv


def step(name, fn):
    print(f"\n=== {name}")
    try:
        fn()
        results.append((name, "PASS"))
        print("  >> PASS")
    except Exception:
        traceback.print_exc()
        results.append((name, "FAIL"))


_off = time.localtime().tm_gmtoff
# 서버(JVM)는 이 PC 의 시간대로 CURRENT_TIMESTAMP 를 쓴다. psql 세션도 같은 시간대로 맞춰야 now() 로 넣은 시각·경과 시간이 어긋나지 않는다
# (dataq-db 컨테이너 기본은 UTC — 맞추지 않으면 now() 로 넣은 LOG 가 9시간 전으로 보인다)
TZ_SQL = f"SET TIME ZONE INTERVAL '{'+' if _off >= 0 else '-'}{abs(_off) // 3600:02d}:{abs(_off) % 3600 // 60:02d}' HOUR TO MINUTE;\n"


def psql(sql):
    r = subprocess.run(PSQL + ["-At", "-q", "-v", "ON_ERROR_STOP=1"], input=(TZ_SQL + sql).encode("utf-8"),
                       capture_output=True, timeout=60)
    if r.returncode != 0:
        raise RuntimeError("psql 실패: " + r.stderr.decode("utf-8", "replace")[:300])
    return r.stdout.decode("utf-8", "replace").strip()


def q(v):
    return "NULL" if v is None else "'" + str(v).replace("'", "''") + "'"


def login_api(user, pwd=None):
    s = requests.Session()
    r = s.post(BASE + "/login", data={"id": user, "password": base64.b64encode((pwd or PW).encode()).decode()},
               allow_redirects=False, timeout=10)
    assert r.status_code == 200, f"로그인 실패 {user} {r.status_code}"
    return s


def post(s, path, body, timeout=30):
    r = s.post(BASE + path, json=body, timeout=timeout)
    try:
        return r.json()
    except Exception:
        return {"http": r.status_code, "text": r.text[:300]}


def get(s, path, params=None, timeout=30):
    r = s.get(BASE + path, params=params, timeout=timeout)
    try:
        return r.json()
    except Exception:
        return {"http": r.status_code, "text": r.text[:300]}


def code(r):
    return r.get("resultCode") if isinstance(r, dict) else None


def msg(r):
    return (r.get("resultMessage") or "") if isinstance(r, dict) else ""


def contents(r):
    c = r.get("contents")
    return json.loads(c) if isinstance(c, str) and c[:1] in "{[" else c


def wait_until(fn, timeout=60, every=1.0):
    end = time.time() + timeout
    v = fn()
    while not v and time.time() < end:
        time.sleep(every)
        v = fn()
    return v


def now_db():
    return psql("select to_char(now(), 'YYYY-MM-DD HH24:MI:SS')")


# ---------------------------------------------------------------- 시험 데이터 헬퍼
def make_user(sa, uid, perms=None, upw=None):
    upw = upw or secrets.token_hex(6)
    r = post(sa, "/api/login/createUser", {"id": uid, "password": base64.b64encode(upw.encode()).decode(),
                                           "name": "110 3묶음 시험 " + uid, "email": "", "admin": False})
    assert code(r) == 200, r
    cleanup.append(lambda: psql(f"delete from quality.tb_user_role where user_id='{uid}'; delete from quality.tb_user where user_id='{uid}'"))
    if perms:
        r = post(sa, "/api/role/save", {"roleNm": "ZZB3역할_" + uid, "roleDesc": "110 3묶음 시험", "perms": perms})
        assert code(r) == 200, r
        rid = r.get("contents")
        cleanup.append(lambda: psql(f"delete from quality.tb_user_role where role_id='{rid}'; delete from quality.tb_role_perm where role_id='{rid}'; "
                                    f"delete from quality.tb_role where role_id='{rid}'"))
        assert code(post(sa, "/api/role/assign", {"userId": uid, "roleId": rid, "sysCd": "*"})) == 200
    return upw


def make_dict(sa, nm, last_clsf, parent=None):
    r = post(sa, "/api/dict/save", {"dictNm": nm, "dictDesc": "110 3묶음 시험", "termLastWordClsfYn": last_clsf, "upperYn": "Y",
                                    "wordEngNmDupYn": "Y", "termEngMaxLen": "", "oldStdYn": "N", "dictTp": "ORG",
                                    "parentDictId": parent or ""})
    assert code(r) == 200, r
    d = contents(r)["dictId"]
    cleanup.append(lambda: drop_dict(d))
    return d


def drop_dict(d):
    psql(f"""delete from quality.tb_change_history where change_id in (select change_id from quality.tb_change_history_detail where target_nm like 'ZZB3%')
               or target_nm like 'ZZB3%' or summary like '%ZZB3%';
             delete from quality.tb_change_history_detail where target_nm like 'ZZB3%';
             delete from quality.tb_aprv_stats where req_item_nm like 'ZZB3%';
             delete from quality.tb_term_resolve_history where resolved_terms_id in (select terms_id from quality.tb_terms where dict_id='{d}');
             delete from quality.tb_terms_words where terms_id in (select terms_id from quality.tb_terms where dict_id='{d}');
             delete from quality.tb_terms where dict_id='{d}';
             delete from quality.tb_word where dict_id='{d}';
             delete from quality.tb_domain where dict_id='{d}';
             delete from quality.tb_domain_clsf where dict_id='{d}';
             delete from quality.tb_domain_grp where dict_id='{d}';
             update quality.tb_std_dict set parent_dict_id=null where parent_dict_id='{d}';
             delete from quality.tb_std_dict where dict_id='{d}';""")


def make_ds(sa, k, dbms, port=25433, user="x", pwd="x"):
    dsn = f"ZZB3_{k}_{TAG}"
    r = post(sa, "/api/sysinfo/createDataSource", {"dsn": dsn, "dsTp": 0, "dbmsTp": dbms, "driverName": dbms, "svrAddr": "127.0.0.1",
                                                   "port": port, "userId": user, "pwd": pwd, "dbName": "postgres"})
    assert code(r) == 200, r
    ds = psql(f"select ds_id from ndata.tb_data_source where dsn='{dsn}'")
    cleanup.append(lambda: psql(f"delete from ndata.tb_data_source where ds_id='{ds}'"))
    psql(f"update ndata.tb_data_source set conn_test_yn='Y' where ds_id='{ds}'")
    return ds


def make_model(sa, nm, ds=None, sys_cd=None, dict_id="DEFAULT"):
    r = post(sa, "/api/dm/createDataModel", {"dataModelNm": nm, "dataModelDsId": ds, "dataModelSysCd": sys_cd,
                                            "modelType": "PHYSICAL", "ver": "1", "dictId": dict_id})
    dm = psql(f"select dm_id from quality.tb_data_model where dm_nm={q(nm)}")
    assert code(r) == 200 and dm, r
    cleanup.append(lambda: drop_model(dm))
    return dm


def drop_model(dm):
    psql(f"""delete from quality.tb_aprv_doc_step where doc_id in (select doc_id from quality.tb_aprv_doc where dm_id='{dm}');
             delete from quality.tb_aprv_doc where dm_id='{dm}';
             delete from quality.tb_diag_schedule_log where schedule_id in (select schedule_id from quality.tb_diag_schedule where data_model_id='{dm}');
             delete from quality.tb_diag_schedule where data_model_id='{dm}';
             delete from quality.tb_diag_result where diag_job_id in (select diag_job_id from quality.tb_diag_job where dm_id='{dm}');
             delete from quality.tb_diag_job where dm_id='{dm}';
             delete from quality.tb_struct_diag_detail where diag_id in (select diag_id from quality.tb_struct_diag_history where data_model_id='{dm}');
             delete from quality.tb_struct_diag_index_detail where diag_id in (select diag_id from quality.tb_struct_diag_history where data_model_id='{dm}');
             delete from quality.tb_struct_diag_constraint_detail where diag_id in (select diag_id from quality.tb_struct_diag_history where data_model_id='{dm}');
             delete from quality.tb_struct_diag_history where data_model_id='{dm}';
             delete from quality.tb_gov_meta_send where dm_id='{dm}';
             delete from quality.tb_batch_run where dm_id='{dm}';
             delete from quality.tb_data_model_change_history where dm_id='{dm}';
             delete from quality.tb_data_model_attr where dm_id='{dm}';
             delete from quality.tb_data_model_constraint where dm_id='{dm}';
             delete from quality.tb_data_model_index where dm_id='{dm}';
             delete from quality.tb_data_model_obj where dm_id='{dm}';
             delete from quality.tb_data_model_clct where dm_id='{dm}';
             delete from quality.tb_data_model_schema where dm_id='{dm}';
             delete from quality.tb_udp_val where dm_id='{dm}';
             delete from quality.tb_term_resolve_history where dm_id='{dm}';
             delete from quality.tb_data_model where dm_id='{dm}';""")


def add_attrs(dm, owner, obj, cols):
    """모델에 승인된 테이블·컬럼을 바로 넣는다. cols: (이름, 한글명, 타입, 길이, 소수)"""
    psql(f"insert into quality.tb_data_model_obj (dm_id, obj_nm, obj_nm_kr, obj_owner, use_yn, aprv_status) "
         f"select '{dm}', '{obj}', '시험테이블', '{owner}', 'Y', 'APPROVED' "
         f"where not exists (select 1 from quality.tb_data_model_obj where dm_id='{dm}' and obj_owner='{owner}' and obj_nm='{obj}')")
    vals = []
    base = int(psql(f"select coalesce(max(attr_ord),0) from quality.tb_data_model_attr where dm_id='{dm}' and obj_nm='{obj}'"))
    for i, (nm, kr, tp, ln, dec) in enumerate(cols, base + 1):
        vals.append(f"('{dm}', '{owner}', '{obj}', '{nm}', {q(kr)}, '{tp}', {ln if ln is not None else 'NULL'}, "
                    f"{dec if dec is not None else 'NULL'}, {i}, 'Y', 'Y', 'APPROVED')")
    psql("insert into quality.tb_data_model_attr (dm_id, obj_owner, obj_nm, attr_nm, attr_nm_kr, data_type, data_len, data_decimal_len, "
         "attr_ord, nullable_yn, use_yn, aprv_status) values " + ",".join(vals))


def make_schedule(s, nm, diag_type, dm, use_yn="N", dict_id=None, **kw):
    body = {"scheduleNm": nm, "diagType": diag_type, "dataModelId": dm, "scheduleType": "SIMPLE", "repeatCycle": "DAILY",
            "repeatTime": "02:00", "useYn": use_yn, "dictId": dict_id}
    body.update(kw)
    r = post(s, "/api/diag/schedule/create", body)
    assert code(r) == 200, r
    sid = r.get("contents")
    cleanup.append(lambda: psql(f"delete from quality.tb_diag_schedule_log where schedule_id='{sid}'; "
                                f"delete from quality.tb_diag_schedule where schedule_id='{sid}'"))
    return sid


def job_row(job, cols="status||'|'||coalesce(dm_clct_id,'-')||'|'||coalesce(dict_id,'-')||'|'||coalesce(cret_user_id,'-')"):
    return psql(f"select {cols} from quality.tb_diag_job where diag_job_id='{job}'")


def wait_job(job, timeout=90):
    return wait_until(lambda: psql(f"select status from quality.tb_diag_job where diag_job_id='{job}' and status not in ('READY','RUNNING')"), timeout)


def job_ids(dm):
    return set(psql(f"select diag_job_id from quality.tb_diag_job where dm_id='{dm}'").split())


def struct_ids(dm):
    return set(psql(f"select diag_id from quality.tb_struct_diag_history where data_model_id='{dm}'").split())


def one_new(before, after):
    new = after - before
    assert len(new) == 1, f"새로 생긴 행 {new}"
    return new.pop()


def sched_last(sid):
    return psql(f"select coalesce(last_exec_status,'-')||'|'||coalesce(last_exec_log_id,'-')||'|'||coalesce(to_char(last_exec_dt,'YYYY-MM-DD HH24:MI:SS'),'-') "
                f"from quality.tb_diag_schedule where schedule_id='{sid}'")


def logs_of(sid):
    v = psql(f"select log_id||'|'||coalesce(exec_status,'')||'|'||coalesce(trigger_type,'')||'|'||coalesce(error_msg,'')||'|'||coalesce(diag_job_id,'')"
             f"||'|'||coalesce(struct_diag_id,'') from quality.tb_diag_schedule_log where schedule_id='{sid}' order by exec_dt")
    return [r.split("|") for r in v.splitlines()] if v else []


def is_gone(r):
    """지운 API — 매핑이 없으면 GET 은 화면(index.html)으로, POST 는 405·404"""
    return r.status_code in (404, 405) or "text/html" in (r.headers.get("content-type") or "")


# ---------------------------------------------------------------- q-executor 내리기·올리기, 가짜 진단 서버
def exec_pid():
    r = subprocess.run(["powershell", "-NoProfile", "-Command",
                        "(Get-NetTCPConnection -LocalPort 28098 -State Listen -ErrorAction SilentlyContinue).OwningProcess"],
                       capture_output=True, text=True, timeout=30)
    return (r.stdout or "").strip().splitlines()[0] if (r.stdout or "").strip() else None


def exec_up():
    try:
        requests.get(EXEC + "/", timeout=2)
        return True
    except Exception:
        return False


def stop_executor():
    pid = exec_pid()
    if pid:
        subprocess.run(["powershell", "-NoProfile", "-Command", f"Stop-Process -Id {pid} -Force"], capture_output=True, timeout=30)
    assert wait_until(lambda: not exec_up(), 30), "q-executor 가 내려가지 않음"


def start_executor():
    err = EXEC_LOG[:-4] + ".err" if EXEC_LOG.endswith(".out") else EXEC_LOG + ".err"
    ps = (f"$env:TMP='{DATAQ_TMP}'; $env:TEMP='{DATAQ_TMP}'; $env:ND_HOME='{ROOT}'; "
          f"Start-Process -FilePath java -ArgumentList '-jar','target\\q-executor.jar' -WorkingDirectory '{ROOT}\\q-executor' "
          f"-RedirectStandardOutput '{EXEC_LOG}' -RedirectStandardError '{err}' -WindowStyle Hidden")
    # 출력을 파이프로 받으면 java 가 그 핸들을 물려받아 끝나지 않는다 — 버린다
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                   stderr=subprocess.DEVNULL, timeout=60)
    assert wait_until(exec_up, 150, 2), "q-executor 가 올라오지 않음"


FAKE = {"msg": "", "hook": None, "calls": []}


class FakeExecutor(BaseHTTPRequestHandler):
    def do_POST(self):
        if "chunked" in (self.headers.get("Transfer-Encoding") or ""):
            raw = b""
            while True:
                n = int(self.rfile.readline().strip() or b"0", 16)
                if n == 0:
                    self.rfile.readline()
                    break
                raw += self.rfile.read(n)
                self.rfile.readline()
        else:
            raw = self.rfile.read(int(self.headers.get("Content-Length") or 0))
        body = json.loads(raw or b"{}")
        FAKE["calls"].append((self.path, body))
        if FAKE["hook"]:
            FAKE["hook"](self.path, body)
        data = json.dumps({"resultCode": 500, "resultMessage": FAKE["msg"], "contents": None}, ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json;charset=UTF-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, *a):
        pass


# ---------------------------------------------------------------- 공통 시험 데이터
def setup(sa):
    # 사용자 — 일반 사용자(권한 없음), 스케줄 관리자(SCHEDULE_ADMIN)
    FX["su"], FX["sm"] = "zzb3u" + TAG, "zzb3m" + TAG
    FX["suPw"] = make_user(sa, FX["su"], upw=os.environ.get("DATAQ_B3_USER_PW"))   # --keep 뒤 화면 검수 때 일반 사용자로 들어갈 비밀번호 (없으면 임의)
    FX["smPw"] = make_user(sa, FX["sm"], ["SCHEDULE_ADMIN"])
    # 사전 — A(형식단어 강제 안 함), B(강제, A 와 무관 — 같은 영문약어 용어를 다른 이름으로), C(강제, 상위 A)
    FX["DA"] = make_dict(sa, "ZZB3사전A_" + TAG, "N")
    FX["DB"] = make_dict(sa, "ZZB3사전B_" + TAG, "Y")
    FX["DC"] = make_dict(sa, "ZZB3사전C_" + TAG, "Y", FX["DA"])
    DA, DB, DC = FX["DA"], FX["DB"], FX["DC"]
    sql = []
    for d, k in ((DA, "A"), (DB, "B"), (DC, "C")):
        sql.append(f"insert into quality.tb_domain_grp (domain_grp_id, domain_grp_nm, comm_stnd_yn, dict_id) values ('ZB3G{k}{TAG}', 'ZZB3그룹', 'N', '{d}')")
        sql.append(f"insert into quality.tb_domain_clsf (domain_clsf_id, domain_clsf_nm, domain_grp_nm, comm_stnd_yn, dict_id) "
                   f"values ('ZB3K{k}{TAG}', 'ZZB3분류{k}', 'ZZB3그룹', 'N', '{d}')")

    def dom(i, d, k, nm, tp, ln, dec=None):
        sql.append(f"insert into quality.tb_domain (domain_id, domain_nm, domain_grp_nm, domain_clsf_nm, domain_desc, data_type, data_len, data_decimal_len, aprv_yn, dict_id) "
                   f"values ('ZB3D{i}{TAG}', {q(nm)}, 'ZZB3그룹', 'ZZB3분류{k}', '110 3묶음 시험', '{tp}', {ln if ln is not None else 'NULL'}, {dec if dec is not None else 'NULL'}, 'Y', '{d}')")
    dom(1, DA, "A", "ZZB3문자30", "VARCHAR", 30)
    dom(2, DA, "A", "ZZB3금액", "NUMBER", 10, 2)
    dom(3, DA, "A", "ZZB3일시A", "DATETIME", None)
    dom(4, DC, "C", "ZZB3일시C", "DATETIME", None)
    dom(5, DB, "B", "ZZB3문자30", "VARCHAR", 99)

    def term(i, d, nm, eng, dm):
        sql.append(f"insert into quality.tb_terms (terms_id, terms_nm, terms_eng_abrv_nm, terms_desc, domain_nm, aprv_yn, dict_id, cret_user_id, cret_dt) "
                   f"values ('ZB3T{i:02d}{TAG}', {q(nm)}, '{eng}', '110 3묶음 시험', {q(dm)}, 'Y', '{d}', 'space', to_char(now(),'YYYYMMDDHH24MISS'))")
    term(1, DA, "ZZB3고객명", "C_KR", "ZZB3문자30")
    term(2, DA, "ZZB3금액", "C_NUM", "ZZB3금액")
    term(3, DA, "ZZB3일시A", "C_DT", "ZZB3일시A")
    term(4, DA, "ZZB3정상일", "C_OK1", "ZZB3문자30")
    term(5, DA, "ZZB3정상이", "C_OK2", "ZZB3문자30")
    term(6, DC, "ZZB3가까운일시", "C_DT", "ZZB3일시C")
    term(7, DB, "ZZB3남의고객명", "C_KR", "ZZB3문자30")
    term(8, DB, "ZZB3남의용어", "ZQX_ZQY", "ZZB3문자30")

    def word(i, d, nm, eng, clsf="N", clsf_nm=None, aprv="Y", user="space"):
        sql.append(f"insert into quality.tb_word (word_id, word_nm, word_eng_abrv_nm, word_eng_nm, word_desc, word_clsf_yn, domain_clsf_nm, aprv_yn, dict_id, cret_user_id, cret_dt) "
                   f"values ('ZB3W{i:02d}{TAG}', {q(nm)}, '{eng}', '{eng}NAME', '110 3묶음 시험', '{clsf}', {q(clsf_nm)}, '{aprv}', '{d}', '{user}', to_char(now(),'YYYYMMDDHH24MISS'))")
    word(1, DA, "ZZB3가", "ZZBA")
    word(2, DA, "ZZB3나", "ZZBB", "Y", "ZZB3분류A")
    word(3, DB, "ZZB3다", "ZZBC")
    word(4, DA, "ZZB3대기", "ZZBP", aprv="N", user="zzb3other")
    psql(";\n".join(sql))
    FX["W1"], FX["W2"], FX["W3"] = f"ZB3W01{TAG}", f"ZB3W02{TAG}", f"ZB3W03{TAG}"

    # 데이터 소스 — 실제 PostgreSQL(개발 DB 의 시험 역할·스키마 2개) 하나, 화면 시험용(접속하지 않음) 하나
    role = ("zzb3r" + TAG).lower()
    pw = "Zz" + secrets.token_hex(6)
    a, b = f"zzb3a{TAG}", f"zzb3b{TAG}"
    psql(f"create role {role} login password '{pw}'")
    for s in (a, b):
        psql(f"create schema {s} authorization {role}")

    def _drop_pg():
        for s in (a, b):
            psql(f"drop schema if exists {s} cascade")
        psql(f"drop role if exists {role}")
    cleanup.append(_drop_pg)
    psql(f"""create table {a}.t1 (id int primary key, nm varchar(20)); create index idx_t1_nm on {a}.t1 (nm);
             create table {a}.t_off (c1 int); create table {b}.t2 (id int primary key, cd varchar(10));
             alter table {a}.t1 owner to {role}; alter table {a}.t_off owner to {role}; alter table {b}.t2 owner to {role};""")
    FX.update(role=role, sa_=a, sb_=b)
    FX["dsPg"] = make_ds(sa, "pg", "PostgreSQL", 25433, role, pw)
    FX["dsUi"] = make_ds(sa, "ui", "PostgreSQL", 1)
    # 시스템 — sys1 은 화면용 데이터 소스를 가진다, sys2 는 없다
    for k, lst in (("sys1", [FX["dsUi"]]), ("sys2", [])):
        nm = f"ZZB3시스템{k[-1]}_{TAG}"
        r = post(sa, "/api/sysinfo/createSysInfo", {"sysNm": nm, "sysTp": 1, "sysDsLst": lst, "sysDesc": "110 3묶음 시험"})
        assert code(r) == 200, r
        cd = psql(f"select sys_cd from quality.tb_sys_info where sys_nm={q(nm)}")
        cleanup.append(lambda cd=cd: psql(f"delete from quality.tb_sys_info where sys_cd='{cd}'"))
        FX[k], FX[k + "Nm"] = cd, nm

    # 모델
    FX["mUNm"] = "ZZB3진단모델_" + TAG
    FX["mU"] = make_model(sa, FX["mUNm"], FX["dsUi"], dict_id=DA)
    add_attrs(FX["mU"], "zzown", "ZZ_T1", [("C_KR", "틀린한글", "VARCHAR", 30, None), ("C_NUM", "ZZB3금액", "NUMBER", 8, 2),
                                           ("C_DT", "틀린일시", "DATETIME", None, None), ("ZQX_ZQY", "ZZB3남의용어", "VARCHAR", 10, None),
                                           ("C_OK1", "ZZB3정상일", "VARCHAR", 30, None), ("C_OK2", "ZZB3정상이", "VARCHAR", 30, None)])
    mu = FX["mU"]
    psql(f"""insert into quality.tb_data_model_clct (dm_clct_id, dm_id, clct_start_dt, clct_end_dt, clct_cmptn_yn, cret_user_id, clct_type) values
             ('ZB3C1{TAG}', '{mu}', '20261001000000', '20261001000100', 'Y', 'space', 'DBMS'),
             ('ZB3C2{TAG}', '{mu}', '20261002000000', '20261002000100', 'Y', 'space', 'DBMS'),
             ('ZB3C3{TAG}', '{mu}', '20261003000000', '20261003000100', 'N', 'space', 'DBMS')""")
    FX["mBigNm"] = "ZZB3대량모델_" + TAG
    FX["mBig"] = make_model(sa, FX["mBigNm"], dict_id=DA)
    psql(f"insert into quality.tb_data_model_obj (dm_id, obj_nm, obj_nm_kr, obj_owner, use_yn, aprv_status) values ('{FX['mBig']}', 'ZZ_BIG', '대량', '', 'Y', 'APPROVED');"
         f"insert into quality.tb_data_model_attr (dm_id, obj_owner, obj_nm, attr_nm, attr_nm_kr, data_type, data_len, attr_ord, nullable_yn, use_yn, aprv_status) "
         f"select '{FX['mBig']}', '', 'ZZ_BIG', 'ZZCOL_' || lpad(g::text, 3, '0'), '대량' || g, 'VARCHAR', 10, g, 'Y', 'Y', 'APPROVED' from generate_series(1, 250) g")
    FX["mSNm"] = "ZZB3구조모델_" + TAG
    FX["mS"] = make_model(sa, FX["mSNm"], FX["dsPg"], dict_id=DA)
    FX["mLNm"] = "ZZB3논리모델_" + TAG
    FX["mL"] = make_model(sa, FX["mLNm"], dict_id=DA)
    psql(f"update quality.tb_data_model set dict_id=null where dm_id='{FX['mL']}'")
    for k in ("mR", "mAuto", "mBusy", "mOv", "mOrph", "mStale", "mDraft"):
        FX[k + "Nm"] = f"ZZB3{k}_{TAG}"
        # mOv 는 BOTH 스케줄을 걸어야 해서 데이터 소스가 있어야 한다 (6.10-11a)
        FX[k] = make_model(sa, FX[k + "Nm"], FX["dsUi"] if k == "mOv" else None, dict_id=DA)
    FX["mSysNm"] = "ZZB3시스템모델_" + TAG
    FX["mSys"] = make_model(sa, FX["mSysNm"], FX["dsUi"], FX["sys1"], DA)
    add_attrs(FX["mDraft"], "", "ZZ_DRAFT", [("C_BASE", "기본", "VARCHAR", 10, None)])

    # 구조 진단 모델 — 실제 이름(소문자)으로 수집한 뒤 스키마 목록만 대문자로 바꾼다 (모델 소유자와 스키마 목록의 대소문자가 다른 경우, 6.10-25a)
    assert code(post(sa, "/api/dm/saveDataModelSchemas", [{"dataModelId": FX["mS"], "schemaNm": s, "useYn": "Y"} for s in (a, b)])) == 200
    r = post(sa, "/api/dm/collectDataModel", {"dataModelId": FX["mS"], "dataModelDsId": FX["dsPg"], "dataModelNm": FX["mSNm"]}, timeout=60)
    assert code(r) == 200, r
    done = wait_until(lambda: psql(f"select clct_cmptn_yn from quality.tb_data_model_clct where dm_id='{FX['mS']}' and clct_end_dt is not null"), 90, 2)
    owners = psql(f"select string_agg(distinct obj_owner, ',') from quality.tb_data_model_obj where dm_id='{FX['mS']}'")
    print(f"  구조 진단 모델 수집 완료={done} / 모델 소유자={owners}")
    assert done == "Y" and owners == f"{a},{b}"
    psql(f"update quality.tb_data_model_schema set schema_nm=upper(schema_nm) where dm_id='{FX['mS']}'")
    psql(f"update quality.tb_data_model_obj set struct_diag_target_yn='N' where dm_id='{FX['mS']}' and obj_nm='t_off'")
    print(f"  사전 A·B·C={DA},{DB},{DC} / 사용자 {FX['su']}·{FX['sm']} / 모델 {len([k for k in FX if k.startswith('m') and not k.endswith('Nm')])}개")


# ---------------------------------------------------------------- S2 진단 서버 내림·거부·재기동
def s2(sa, su, sm):
    mr, ms = FX["mR"], FX["mS"]
    FX["SR"] = make_schedule(sa, "ZZB3즉시실행_" + TAG, "STANDARD", mr, dict_id=FX["DB"])

    def _down():
        stop_executor()
        jb, hb = job_ids(mr), struct_ids(ms)
        r1 = post(sa, "/api/diag/startDiag", {"dataModelId": mr, "dictId": FX["DA"]})
        j = one_new(jb, job_ids(mr))
        r2 = post(sa, "/api/std/structDiag/execute", {"dataModelId": ms})
        h = one_new(hb, struct_ids(ms))
        hst = psql(f"select status from quality.tb_struct_diag_history where diag_id='{h}'")
        r3 = post(sa, "/api/diag/schedule/runNow", {"scheduleId": FX["SR"]})
        lg = logs_of(FX["SR"])[-1]
        print(f"  표준 진단 → {code(r1)} {msg(r1)[:60]} / Job {job_row(j, 'status')}\n  구조 진단 → {code(r2)} {msg(r2)[:60]} / 이력 {hst}"
              f"\n  즉시 실행 → {code(r3)} / LOG {lg[1]} {lg[3][:50]} / 스케줄 {sched_last(FX['SR'])}")
        assert code(r1) == 500 and msg(r1).startswith("진단 서버 연결 실패: ") and job_row(j, "status") == "ERROR"
        assert code(r2) == 500 and msg(r2).startswith("진단 서버 연결 실패: ") and hst == "ERROR"
        assert lg[1] == "ERROR" and lg[3].startswith("[CONFIG] q-executor 호출 실패") and sched_last(FX["SR"]).startswith("ERROR|" + lg[0])
    step("S2-1. 진단 서버가 내려가 있으면 Job·구조 진단 이력을 ERROR 로, '진단 서버 연결 실패' (6.10-5a)", _down)

    srv = ThreadingHTTPServer(("127.0.0.1", 28098), FakeExecutor)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        def _reject():
            FAKE["msg"] = "[TEST] 진단 서버가 거부함"
            out = []
            for _ in range(2):   # 두 번째도 409 가 아니어야 한다 (첫 거부 Job·이력이 READY 로 남지 않음)
                jb, hb = job_ids(mr), struct_ids(ms)
                r1 = post(sa, "/api/diag/startDiag", {"dataModelId": mr, "dictId": FX["DA"]})
                j = one_new(jb, job_ids(mr))
                r2 = post(sa, "/api/std/structDiag/execute", {"dataModelId": ms, "clctId": "ZZB3무시"})
                h = one_new(hb, struct_ids(ms))
                out.append((code(r1), msg(r1), job_row(j, "status"), code(r2), msg(r2),
                            psql(f"select status from quality.tb_struct_diag_history where diag_id='{h}'")))
            calls = [p for p, _ in FAKE["calls"]]
            print(f"  거부 응답 2회 = {out}\n  가짜 서버가 받은 요청 = {calls}")
            assert all(o == (500, "[TEST] 진단 서버가 거부함", "ERROR", 500, "[TEST] 진단 서버가 거부함", "ERROR") for o in out)
            # 6.10-27c — 구조 진단 실행 요청에 수집건(clctId) 이 실리지 않는다
            sbody = [b for p, b in FAKE["calls"] if p.endswith("/api/structDiag/run")]
            assert sbody and all("clctId" not in b and b.get("dataModelId") == ms for b in sbody), sbody
        step("S2-2. 진단 서버가 거부하면 Job·구조 진단 이력을 ERROR 로, 응답 사유 그대로 (6.10-5b, 6.10-27c 요청에 clctId 없음)", _reject)

        def _ui_reject():
            ui = UI.get()
            ui.login("space")
            ui.open("structDiag", "진단 실행")
            ui.js(VM + "const c = vm('tab_structDiag', 'DSStructDiag'); c.selectedModel = arguments[0];", ms)
            time.sleep(1)
            ui.click_text("#tab_structDiag button", "진단 실행")
            snack = ui.snack("진단 요청 실패")
            time.sleep(1.5)
            hist = ui.js(VM + "const c = vm('tab_structDiag', 'DSStructDiag'); return c.historyList.filter(h => h.dataModelId === arguments[0]).map(h => h.status);", ms)
            ui.shot("s2_structdiag_reject")
            print(f"  스낵바={snack!r} / 이력 상태={hist}")
            assert snack == "진단 요청 실패: [TEST] 진단 서버가 거부함" and hist and hist[0] == "ERROR"
        step("S2-3. 구조 진단 화면 — 진단 서버 거부 사유를 스낵바에, 이력 다시 읽기 (6.10-28)", _ui_reject)

        def _runnow():
            sid = FX["SR"]
            FAKE["msg"] = "[CONFIG] 시험 거부"
            r = post(sa, "/api/diag/schedule/runNow", {"scheduleId": sid})
            lg = logs_of(sid)[-1]
            last = sched_last(sid)
            print(f"  즉시 실행 거부 → {code(r)} {msg(r)!r} / LOG {lg[1]} {lg[3]!r} / 스케줄 {last}")
            assert code(r) == 500 and msg(r) == "[CONFIG] 시험 거부" and lg[1] == "ERROR" and lg[3] == "[EXECUTOR] [CONFIG] 시험 거부"
            assert last.startswith("ERROR|" + lg[0])
            # launcher 가 이미 마감한 LOG — center 가 덮지 않고 스케줄 LAST_EXEC 도 그대로
            FAKE["hook"] = lambda p, b: psql(f"update quality.tb_diag_schedule_log set exec_status='ERROR', error_msg='[CONFIG] 시험 launcher 마감', "
                                            f"exec_end_dt=now() where log_id={q(b.get('logId'))}")
            r = post(sa, "/api/diag/schedule/runNow", {"scheduleId": sid})
            FAKE["hook"] = None
            lg2 = logs_of(sid)[-1]
            last2 = sched_last(sid)
            print(f"  launcher 가 먼저 마감 → {code(r)} / LOG {lg2[1]} {lg2[3]!r} / 스케줄 {last2}")
            assert code(r) == 500 and lg2[0] != lg[0] and lg2[3] == "[CONFIG] 시험 launcher 마감" and last2.startswith("ERROR|" + lg[0])
        step("S2-4. 즉시 실행 — 진단 서버 거부 시 실패 응답·RUNNING LOG 마감·마지막 실행 갱신, launcher 가 마감한 LOG 는 덮지 않음 (6.10-7a)", _runnow)

        def _notfound():
            sid = FX["SR"]
            FAKE["msg"] = "[DATA_NOT_FOUND] 스케줄 없음: " + sid
            r = post(sa, "/api/diag/schedule/runNow", {"scheduleId": sid})
            lg = logs_of(sid)[-1]
            running = psql(f"select count(*) from quality.tb_diag_schedule_log where schedule_id='{sid}' and exec_status='RUNNING'")
            print(f"  [DATA_NOT_FOUND] → {code(r)} {msg(r)!r} / LOG {lg[1]} {lg[3]!r} / RUNNING 남은 LOG {running}")
            assert code(r) == 500 and lg[1] == "ERROR" and lg[3] == "[EXECUTOR] [DATA_NOT_FOUND] 스케줄 없음: " + sid and running == "0"
        step("S2-5. 즉시 실행 — 진단 서버가 [DATA_NOT_FOUND] 로 거부하면 RUNNING LOG 를 바로 마감 (6.10-7b)", _notfound)
    finally:
        srv.shutdown()
        srv.server_close()

    # 기동 정리 대상 — 죽은 executor 가 남긴 것처럼
    mo = FX["mOrph"]
    o = {k: f"ZB3O{k}{TAG}" for k in ("j1", "j2", "s1", "g1", "g2", "b1")}
    FX["orph"] = o
    psql(f"""insert into quality.tb_diag_job (diag_job_id, dm_id, status, cret_dt, cret_user_id, dict_id) values
               ('{o['j1']}', '{mo}', 'READY', to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}'),
               ('{o['j2']}', '{mo}', 'RUNNING', to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}');
             insert into quality.tb_struct_diag_history (diag_id, data_model_id, status, cret_user_id) values ('{o['s1']}', '{ms}', 'RUNNING', 'space');
             insert into quality.tb_gov_meta_send (send_id, dm_id, trigger_type, status, start_dt) values
               ('{o['g1']}', '{mo}', 'SCHEDULE', 'RUNNING', to_char(now(),'YYYYMMDDHH24MISS')),
               ('{o['g2']}', '{mo}', 'MANUAL', 'RUNNING', to_char(now(),'YYYYMMDDHH24MISS'));
             insert into quality.tb_batch_run (run_id, run_tp, status, start_dt, dm_id) values ('{o['b1']}', 'UNUSED_STD', 'RUNNING', to_char(now(),'YYYYMMDDHH24MISS'), '{mo}');""")

    def _restart():
        start_executor()
        FX["T0"] = time.time()
        setup_runner(sa)      # 실행기 첫 평가(기동 30초 뒤) 전에 심는다
        st = psql(f"""select string_agg(x, ',' order by n) from (
                        select 1 n, status||coalesce('/'||left(end_dt,8),'') x from quality.tb_diag_job where diag_job_id='{o['j1']}' union all
                        select 2, status||coalesce('/'||left(end_dt,8),'') from quality.tb_diag_job where diag_job_id='{o['j2']}' union all
                        select 3, status from quality.tb_struct_diag_history where diag_id='{o['s1']}' union all
                        select 4, status from quality.tb_gov_meta_send where send_id='{o['g1']}' union all
                        select 5, status from quality.tb_gov_meta_send where send_id='{o['g2']}' union all
                        select 6, status||'/'||coalesce(error_msg,'') from quality.tb_batch_run where run_id='{o['b1']}') t""")
        logline = [ln for ln in open(EXEC_LOG, encoding="utf-8", errors="replace") if "기동 정리" in ln]
        r = post(sa, "/api/diag/startDiag", {"dataModelId": mo, "dictId": FX["DA"]})
        print(f"  재기동 뒤 Job·Job·구조·범정부(예약)·범정부(화면)·배치 = {st}\n  로그={logline[-1].strip()[-110:] if logline else None}"
              f"\n  같은 모델 진단 시작 → {code(r)} {msg(r)[:40]}")
        today = time.strftime("%Y%m%d")
        assert st == f"ERROR/{today},ERROR/{today},ERROR,ERROR,RUNNING,ERROR/q-executor 재기동으로 중단됨", st
        assert logline and code(r) == 200
        wait_job(r.get("contents"))
        r2 = requests.post(EXEC + "/api/diag/schedule/runNow", json={"scheduleId": "ZZB3없음", "logId": "ZZB3없음"}, timeout=10).json()
        print(f"  진단 서버에 없는 스케줄 즉시 실행 → {r2.get('resultCode')} {r2.get('resultMessage')!r}")
        assert r2.get("resultCode") == 500 and "[DATA_NOT_FOUND] 스케줄 없음" in (r2.get("resultMessage") or "")
    step("S2-6. q-executor 재기동 — 남은 READY·RUNNING Job·구조 진단·예약 범정부 메타·배치를 ERROR 로, 화면 검증은 그대로 (6.10-6, 6.10-7b 진단 서버 응답)", _restart)


# ---------------------------------------------------------------- R 예약 실행기
def setup_runner(sa):
    """CRON '매년 1월 1일 0시' + 지난 실행 2025-01-01 → 첫 평가에서 도래"""
    cron = {"scheduleType": "CRON", "cronExpr": "0 0 0 1 1 *"}
    R = FX.setdefault("R", {})
    R["K1"] = make_schedule(sa, "ZZB3건너뜀_진행중Job_" + TAG, "STANDARD", FX["mBusy"], "Y", **cron)
    R["K2"] = make_schedule(sa, "ZZB3건너뜀_BOTH실행중_" + TAG, "STANDARD", FX["mOv"], "Y", **cron)
    R["K2b"] = make_schedule(sa, "ZZB3겹침BOTH_" + TAG, "BOTH", FX["mOv"], "N", dict_id=None)
    R["K3"] = make_schedule(sa, "ZZB3자동실행_" + TAG, "STANDARD", FX["mAuto"], "Y", dict_id=FX["DB"], **cron)
    # 자기 실행이 진행 중인 스케줄 — 건너뛰어도 마지막 실행은 그 실행(RUNNING) 그대로
    R["K4"] = make_schedule(sa, "ZZB3건너뜀_자기실행중_" + TAG, "STANDARD", FX["mOv"], "Y", **cron)
    R["S8"] = make_schedule(sa, "ZZB3타임아웃_" + TAG, "STANDARD", FX["mStale"])
    for k in ("K1", "K2", "K3"):
        psql(f"update quality.tb_diag_schedule set last_exec_dt='2025-01-01 00:00:00', last_exec_status='DONE', last_exec_log_id='ZB3PREV{k}' "
             f"where schedule_id='{R[k]}'")
    jb, js, hs = f"ZB3JB{TAG}", f"ZB3JS{TAG}", f"ZB3HS{TAG}"
    R.update(jb=jb, js=js, hs=hs, l2=f"ZB3L2{TAG}", l4=f"ZB3L4{TAG}", l8a=f"ZB3L8A{TAG}", l8b=f"ZB3L8B{TAG}", l8c=f"ZB3L8C{TAG}")
    psql(f"""update quality.tb_diag_schedule set last_exec_dt='2025-01-01 00:00:00', last_exec_status='RUNNING', last_exec_log_id='{R['l4']}'
               where schedule_id='{R['K4']}';
             insert into quality.tb_diag_schedule_log (log_id, schedule_id, exec_dt, exec_status, trigger_type, diag_type, schedule_nm_snapshot) values
               ('{R['l4']}', '{R['K4']}', now(), 'RUNNING', 'AUTO', 'STANDARD', 'ZZB3건너뜀_자기실행중');""")
    psql(f"""insert into quality.tb_diag_job (diag_job_id, dm_id, status, cret_dt, cret_user_id, dict_id) values
               ('{jb}', '{FX['mBusy']}', 'READY', to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}'),
               ('{js}', '{FX['mStale']}', 'RUNNING', to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}');
             insert into quality.tb_struct_diag_history (diag_id, data_model_id, status, cret_user_id) values ('{hs}', '{FX['mStale']}', 'RUNNING', 'space');
             insert into quality.tb_diag_schedule_log (log_id, schedule_id, exec_dt, exec_status, trigger_type, diag_type, schedule_nm_snapshot) values
               ('{R['l2']}', '{R['K2b']}', now(), 'RUNNING', 'MANUAL', 'BOTH', 'ZZB3겹침BOTH');
             insert into quality.tb_diag_schedule_log (log_id, schedule_id, exec_dt, exec_status, trigger_type, diag_type, diag_job_id, struct_diag_id, schedule_nm_snapshot) values
               ('{R['l8a']}', '{R['S8']}', now() - interval '2 hours', 'RUNNING', 'AUTO', 'STANDARD', null, null, 'ZZB3타임아웃'),
               ('{R['l8b']}', '{R['S8']}', now() - interval '2 hours', 'RUNNING', 'AUTO', 'STANDARD', '{js}', null, 'ZZB3타임아웃'),
               ('{R['l8c']}', '{R['S8']}', now() - interval '2 hours', 'RUNNING', 'AUTO', 'STRUCT', null, '{hs}', 'ZZB3타임아웃');""")


def r_check(sa, su, sm):
    R = FX["R"]
    t0 = FX["T0"]

    def _skip_busy():
        # 첫 SKIPPED 가 생긴 뒤 평가가 한 번 더 지나도 SKIPPED 는 하나 (슬롯 소비)
        first = wait_until(lambda: [l for l in logs_of(R["K1"]) if l[1] == "SKIPPED"], max(5, t0 + 150 - time.time()), 3)
        assert first, "SKIPPED 가 생기지 않음"
        age = float(psql(f"select extract(epoch from now() - exec_dt) from quality.tb_diag_schedule_log where log_id='{first[0][0]}'"))
        time.sleep(max(0, 70 - age))     # 실행기 평가(1분 주기)가 한 번 더 지나도록
        lg = logs_of(R["K1"])
        last = sched_last(R["K1"])
        jobs = psql(f"select count(*) from quality.tb_diag_job where dm_id='{FX['mBusy']}'")
        print(f"  LOG={[(l[1], l[2], l[3]) for l in lg]}\n  스케줄(상태|로그|마지막 실행)={last} / 모델 Job 수={jobs}")
        assert len(lg) == 1 and lg[0][1] == "SKIPPED" and lg[0][2] == "AUTO" and lg[0][3] == "[SKIPPED] 같은 모델에 진행 중인 진단이 있음"
        st, lid, dt = last.split("|")
        # 리뷰 지적 — 지난 실행 상태(DONE)가 건너뛴 시각과 섞여 '최종 실행' 이 완료처럼 보이지 않게, 상태·로그 ID 도 이 SKIPPED 로
        assert st == "SKIPPED" and lid == lg[0][0] and dt >= time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(t0)) and jobs == "1"
    step("R-1. 예약 실행 — 같은 모델에 화면 진단(READY Job) 진행 중이면 SKIPPED, 그 슬롯을 소비해 다음 평가에서 또 건너뛰지 않음, 마지막 실행은 SKIPPED·그 LOG (6.10-12a·6.10-9)", _skip_busy)

    def _skip_overlap():
        lg = wait_until(lambda: [l for l in logs_of(R["K2"]) if l[1] == "SKIPPED"], 30, 3)
        jobs = psql(f"select count(*) from quality.tb_diag_job where dm_id='{FX['mOv']}'")
        last = sched_last(R["K2"])
        print(f"  LOG={[(l[1], l[3]) for l in logs_of(R['K2'])]} / 모델 Job 수={jobs} / 스케줄(상태|로그|시각)={last}")
        assert lg and lg[0][3] == "[SKIPPED] 동일 모델+유형 RUNNING 존재" and jobs == "0" and len(logs_of(R["K2"])) == 1
        assert last.split("|")[:2] == ["SKIPPED", lg[0][0]]
    step("R-2. 예약 실행 — 같은 모델에 BOTH 스케줄 실행이 RUNNING 이면 STANDARD 스케줄은 SKIPPED, 마지막 실행도 SKIPPED (6.10-12b·9)", _skip_overlap)

    def _skip_self():
        lg = wait_until(lambda: [l for l in logs_of(R["K4"]) if l[1] == "SKIPPED"], 30, 3)
        last = sched_last(R["K4"])
        print(f"  LOG={[(l[0][-6:], l[1], l[3]) for l in logs_of(R['K4'])]} / 스케줄(상태|로그|시각)={last}")
        st, lid, dt = last.split("|")
        assert lg and lg[0][3] == "[SKIPPED] 동일 모델+유형 RUNNING 존재"
        assert st == "RUNNING" and lid == R["l4"] and dt >= time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(t0))
    step("R-2b. 예약 실행 — 자기 실행이 진행 중이라 건너뛰면 슬롯만 소비하고 마지막 실행은 그 실행(RUNNING·그 LOG) 그대로 (6.10-9)", _skip_self)

    def _auto():
        lg = wait_until(lambda: [l for l in logs_of(R["K3"]) if l[1] in ("DONE", "ERROR")], 90, 3)
        lg = logs_of(R["K3"])
        job = lg[0][4] if lg else None
        jr = job_row(job) if job else None
        print(f"  LOG={[(l[1], l[2], l[3]) for l in lg]} / Job(상태|수집|사전|실행자)={jr} / 스케줄={sched_last(R['K3'])}")
        assert len(lg) == 1 and lg[0][1] == "DONE" and lg[0][2] == "AUTO"
        assert jr == f"DONE|-|{FX['DB']}|space"
    step("R-3. 예약 자동 실행 — 실행자는 스케줄 등록자, 사전은 스케줄 사전, 수집 이력 없는 모델도 진단 (6.10-13a·6.10-2·6.10-10)", _auto)

    def _stale():
        wait_until(lambda: psql(f"select exec_status from quality.tb_diag_schedule_log where log_id='{R['l8a']}'") == "ERROR",
                   max(5, t0 + 170 - time.time()), 3)
        st = psql(f"select string_agg(log_id||':'||exec_status||':'||coalesce(error_msg,''), ' / ' order by log_id) from quality.tb_diag_schedule_log "
                  f"where schedule_id='{R['S8']}'")
        last = sched_last(R["S8"])
        print(f"  LOG={st}\n  스케줄={last}")
        assert f"{R['l8a']}:ERROR:[TIMEOUT] 진단 실행이 비정상 종료됨 (타임아웃 초과)" in st
        assert f"{R['l8b']}:RUNNING:" in st and f"{R['l8c']}:RUNNING:" in st
        assert last.startswith(f"ERROR|{R['l8a']}|")
    step("R-4. 타임아웃 마감 — 60분 넘은 RUNNING LOG 를 ERROR 로 닫고 마지막 실행도 갱신, 기저 Job·구조 진단이 진행 중이면 제외 (6.10-8a·8b)", _stale)


# ---------------------------------------------------------------- S1 표준화 진단·구조 진단
def s1(sa, su, sm):
    mu = FX["mU"]

    def _clct():
        r = post(sa, "/api/diag/startDiag", {"dataModelId": mu, "dictId": FX["DC"]})
        assert code(r) == 200, r
        FX["J1"] = r.get("contents")
        wait_job(FX["J1"])
        r = post(sa, "/api/diag/startDiag", {"dataModelId": FX["mBig"]})
        FX["JBig"] = r.get("contents")
        wait_job(FX["JBig"], 120)
        j1 = job_row(FX["J1"])
        jb = job_row(FX["JBig"])
        lst = {j["diagJobId"]: j for j in post(sa, "/api/diag/getDiagJobList", {}) if j["diagJobId"] in (FX["J1"], FX["JBig"])}
        print(f"  수집 이력 있는 모델 Job={j1} 수집일시={lst[FX['J1']].get('clctDt')} / 없는 모델 Job={jb} 수집일시={lst[FX['JBig']].get('clctDt')}")
        assert j1 == f"DONE|ZB3C2{TAG}|{FX['DC']}|space" and lst[FX["J1"]]["clctDt"] == "2026-10-02 00:01:00"
        assert jb.startswith("DONE|-|") and lst[FX["JBig"]].get("clctDt") is None
    step("S1-1. 진단 Job 에 모델의 최근 완료 수집 (실패 수집 제외, 없으면 비움) (6.10-20a)", _clct)

    def _detail():
        rows = requests.get(BASE + "/api/diag/getDiagResultDetail", params={"diagJobId": FX["J1"]}, cookies=sa.cookies, timeout=30).json()
        by = {}
        for r in rows:
            by.setdefault(r["attrNm"], []).append(r)
        view = {k: (v[0].get("stdTermsNm"), v[0].get("stdDataType"), v[0].get("stdDataLen"), v[0].get("stdDataDecimalLen"),
                    v[0].get("matchedTermsNm"), v[0].get("diagTypes")) for k, v in by.items()}
        for k, v in sorted(view.items()):
            print(f"  {k:8s} 표준용어·타입·길이·소수·한글명 매칭·이슈 = {v}")
        none = requests.get(BASE + "/api/diag/getDiagResultDetail", params={"diagJobId": "ZZB3없는Job"}, cookies=sa.cookies, timeout=30).json()
        assert len(rows) == 6 and all(len(v) == 1 for v in by.values()), "행 수"
        assert view["C_KR"][:3] == ("ZZB3고객명", "VARCHAR", 30) and view["C_KR"][5] == "TERM_KR_NM_MISMATCH"
        assert view["C_DT"][0] == "ZZB3가까운일시" and view["C_NUM"][1:4] == ("NUMBER", 10, 2) and "DATA_LEN_MISMATCH" in view["C_NUM"][5]
        assert view["ZQX_ZQY"][0] is None and view["ZQX_ZQY"][4] is None and view["ZQX_ZQY"][5] == "TERM_NOT_EXIST"
        assert view["C_OK1"][5] is None and none == []
    step("S1-2. 진단 결과 상세 — Job 사전과 상위 사전에서만, 가까운 사전 우선, 컬럼당 1행, 없는 Job 은 빈 목록 (6.10-15)", _detail)

    def _progress():
        jid = FX["JBig"]
        vals = []
        for ln in open(EXEC_LOG, encoding="utf-8", errors="replace"):
            if "updateDiagJobProgress" in ln and "Parameters:" in ln and jid in ln:
                m = re.search(r"Parameters: (\d+)\(Integer\), (\d+)\(Integer\)", ln)
                if m:
                    vals.append(int(m.group(2)))
        print(f"  250컬럼(모두 이슈) 진단의 진행 갱신(처리 컬럼 수) = {vals}")
        assert vals == [0, 100, 200, 250]
    step("S1-3. 표준화 진단 진행 갱신은 컬럼 100개마다 (이슈 50건 배치와 따로) (6.10-23)", _progress)

    def _struct():
        ms, a, b = FX["mS"], FX["sa_"], FX["sb_"]

        def run():
            r = post(sa, "/api/std/structDiag/execute", {"dataModelId": ms, "clctId": "ZZB3무시"})
            assert code(r) == 200, r
            did = r.get("contents")
            wait_until(lambda: psql(f"select status from quality.tb_struct_diag_history where diag_id='{did}' and status not in ('READY','RUNNING')"), 60, 2)
            h = psql(f"select status||'|'||coalesce(schema_nm,'')||'|'||total_tables||'|'||total_indexes||'|'||total_constraints||'|'||"
                     f"added_tables||','||added_columns||','||modified_columns||','||deleted_tables||','||deleted_columns||'|'||"
                     f"added_indexes||','||modified_indexes||','||deleted_indexes||'|'||added_constraints||','||modified_constraints||','||deleted_constraints "
                     f"from quality.tb_struct_diag_history where diag_id='{did}'")
            det = psql(f"select string_agg(x, ' ; ') from (select table_nm||'.'||coalesce(column_nm,'')||':'||change_type x from quality.tb_struct_diag_detail where diag_id='{did}' "
                       f"union all select table_nm||'#'||index_nm||':'||change_type from quality.tb_struct_diag_index_detail where diag_id='{did}' "
                       f"union all select table_nm||'@'||constraint_nm||':'||change_type from quality.tb_struct_diag_constraint_detail where diag_id='{did}') t")
            return did, h, det
        d1, h1, det1 = run()
        psql(f"alter table {a}.t1 add column nm2 varchar(5); drop index {a}.idx_t1_nm; alter table {a}.t_off add column c2 int;")
        d2, h2, det2 = run()
        print(f"  수집 직후(상태|스키마|테이블|인덱스|제약|컬럼 변경|인덱스 변경|제약 변경) = {h1}\n  변경 뒤 = {h2}\n  변경 상세 = {det2}")
        sch = f"{a.upper()},{b.upper()}"
        assert h1 == f"DONE|{sch}|2|3|2|0,0,0,0,0|0,0,0|0,0,0", h1
        assert h2.split("|")[0:2] == ["DONE", sch] and h2.split("|")[5:] == ["0,1,0,0,0", "0,0,1", "0,0,0"], h2
        assert "t1.nm2:ADDED" in det2 and "t1#idx_t1_nm:DELETED" in det2 and "t_off" not in det2
    step("S1-4. 구조 진단 — 스키마 목록(대문자)과 소유자(소문자)가 달라도 인덱스·제약 비교가 맞음, 진단 제외 테이블 제외, 스키마 목록 기록 (6.10-24b·25a·26·27c)", _struct)

    def _gone():
        out = {
            "center compareSchema": is_gone(sa.post(BASE + "/api/std/structDiag/compareSchema", json={"dataModelId": FX["mS"]}, timeout=10)),
            "executor compareSchema": is_gone(requests.post(EXEC + "/api/structDiag/compareSchema", json={"dataModelId": FX["mS"]}, timeout=10)),
            "getDiagResultList": is_gone(sa.get(BASE + "/api/diag/getDiagResultList", params={"diagJobId": FX["J1"]}, timeout=10)),
            "schedule/{id}": is_gone(sa.get(BASE + "/api/diag/schedule/" + FX["SR"], params={"scheduleId": FX["SR"]}, timeout=10)),
            "schedule/logs/{id}": is_gone(sa.get(BASE + "/api/diag/schedule/logs/ZZB3", params={"logId": "ZZB3"}, timeout=10)),
            "structDiag/latestStatus": is_gone(sa.get(BASE + "/api/std/structDiag/latestStatus", params={"dsId": FX["dsPg"]}, timeout=10)),
            "executor analyzeTermsBatch": is_gone(requests.post(EXEC + "/api/std/analyzeTermsBatch", json={"termNames": ["ZZB3가"]}, timeout=10)),
        }
        kept = {
            "getDiagResultSummary": "json" in sa.get(BASE + "/api/diag/getDiagResultSummary", params={"diagJobId": FX["J1"]}, timeout=10).headers.get("content-type", ""),
            "schedule/list": "json" in sa.get(BASE + "/api/diag/schedule/list", timeout=10).headers.get("content-type", ""),
            "schedule/logs": "json" in sa.get(BASE + "/api/diag/schedule/logs", params={"scheduleId": FX["SR"]}, timeout=10).headers.get("content-type", ""),
            "structDiag/history": "json" in sa.get(BASE + "/api/std/structDiag/history", params={"dataModelId": FX["mS"]}, timeout=10).headers.get("content-type", ""),
            "center analyzeTermsBatch": isinstance(post(sa, "/api/std/analyzeTermsBatch", {"termNames": ["ZZB3가ZZB3나"], "dictId": FX["DA"]}), list),
        }
        print(f"  지운 API(없음=True) {out}\n  남긴 API(JSON=True) {kept}")
        assert all(out.values()) and all(kept.values())
    step("S1-5. 지운 API — 스키마 비교(center·executor), 결과 목록, 스케줄 단건·LOG 단건, 최신 구조 진단 상태, executor 용어 분석 (6.10-27a·27b·29, 7.5-3e)", _gone)

    def _screens():
        comp = os.path.join(ROOT, "q-center", "vue", "front", "src", "components")
        files = {f: os.path.exists(os.path.join(comp, f)) for f in ("DSStructDiagHistory.vue", "DSSchemaCompare.vue")}
        nd = open(os.path.join(ROOT, "q-center", "vue", "front", "src", "views", "content", "NdContent.vue"), encoding="utf-8").read()
        html = requests.get(BASE + "/", timeout=10).text
        app = re.search(r'src=/?(static/js/app\.[0-9a-f]+\.js)', html) or re.search(r'src="/?(static/js/app\.[0-9a-f]+\.js)"', html)
        js = requests.get(BASE + "/" + app.group(1), timeout=30).text
        print(f"  파일 있음={files} / NdContent schemaCompare={'schemaCompare' in nd} / 배포 화면에 schemaCompare={'schemaCompare' in js} "
              f"compareSchema={'compareSchema' in js} / DSStructDiagHistory={'DSStructDiagHistory' in js}")
        assert not any(files.values()) and "schemaCompare" not in nd and "schemaCompare" not in js and "compareSchema" not in js
    step("S1-6. 메뉴 없는 화면 정리 — 구조 진단 이력·스키마 비교 화면과 schemaCompare 탭 분기 없음 (6.10-31a·31b)", _screens)


# ---------------------------------------------------------------- S3 스케줄 API
def s3(sa, su, sm):
    s_su = login_api(FX["su"], FX["suPw"])
    s_sm = login_api(FX["sm"], FX["smPw"])
    FX["s_su"], FX["s_sm"] = s_su, s_sm
    base = {"scheduleNm": "ZZB3검사_" + TAG, "diagType": "STANDARD", "dataModelId": FX["mR"], "scheduleType": "SIMPLE",
            "repeatCycle": "DAILY", "repeatTime": "02:00", "useYn": "N"}

    def mk(**kw):
        b = dict(base)
        b.update(kw)
        r = post(sa, "/api/diag/schedule/create", b)
        if code(r) == 200:
            sid = r.get("contents")
            cleanup.append(lambda: psql(f"delete from quality.tb_diag_schedule where schedule_id='{sid}'"))
        return r

    def _perm():
        r1 = post(s_su, "/api/diag/schedule/create", base)
        r2 = post(s_su, "/api/diag/schedule/runNow", {"scheduleId": FX["SR"]})
        lst = get(s_su, "/api/diag/schedule/list")
        print(f"  일반 사용자 등록 → {code(r1)} {msg(r1)} / 즉시 실행 → {code(r2)} / 목록 {type(lst).__name__} {len(lst)}건")
        assert code(r1) == 403 and msg(r1) == "권한이 없습니다 (스케줄 관리 권한 필요)" and code(r2) == 403 and isinstance(lst, list)
    step("S3-1. 스케줄 — 권한 없는 사용자는 등록·즉시 실행 403, 목록은 본다", _perm)

    def _time():
        cases = [("24:00", {}), ("29:59", {}), ("23:59", {}), ("00:00", {}),
                 ("02:00", {"repeatCycle": "WEEKLY", "repeatDayOfWeek": 0}), ("02:00", {"repeatCycle": "WEEKLY", "repeatDayOfWeek": 8}),
                 ("02:00", {"repeatCycle": "WEEKLY", "repeatDayOfWeek": 7}),
                 ("02:00", {"repeatCycle": "MONTHLY", "repeatDayOfMonth": 0}), ("02:00", {"repeatCycle": "MONTHLY", "repeatDayOfMonth": 29}),
                 ("02:00", {"repeatCycle": "MONTHLY", "repeatDayOfMonth": 28})]
        got = []
        for t, extra in cases:
            r = mk(repeatTime=t, **extra)
            got.append((t, extra.get("repeatDayOfWeek", extra.get("repeatDayOfMonth")), code(r), msg(r) if code(r) != 200 else ""))
        sid = FX["SR"]
        up = post(sa, "/api/diag/schedule/update", {**base, "scheduleId": sid, "repeatTime": "24:30", "dictId": FX["DB"]})
        for g in got:
            print(f"  {g}")
        print(f"  수정 24:30 → {code(up)} {msg(up)}")
        exp = [400, 400, 200, 200, 400, 400, 200, 400, 400, 200]
        assert [g[2] for g in got] == exp
        assert got[0][3] == "실행 시각(HH:mm) 형식 오류" and got[4][3] == got[5][3] == "요일은 1(월)~7(일) 이어야 합니다."
        assert got[7][3] == got[8][3] == "일자는 1~28 이어야 합니다." and code(up) == 400 and msg(up) == "실행 시각(HH:mm) 형식 오류"
    step("S3-2. 스케줄 등록·수정 검사 — 실행 시각 00~23시, 요일 1~7, 일자 1~28 (6.10-11b)", _time)

    def _struct_need_ds():
        r1 = mk(diagType="STRUCT", dataModelId=FX["mL"])
        r2 = mk(diagType="BOTH", dataModelId=FX["mL"], dictId=FX["DA"])
        r3 = mk(diagType="STRUCT", dataModelId="ZZB3없는모델")
        r4 = mk(diagType="STRUCT", dataModelId=FX["mS"])
        print(f"  논리 모델 STRUCT → {code(r1)} {msg(r1)} / BOTH → {code(r2)} / 없는 모델 → {code(r3)} {msg(r3)} / 물리 모델 → {code(r4)}")
        m = "데이터 소스가 없는 논리 모델은 구조 변경 진단을 예약할 수 없습니다."
        assert (code(r1), msg(r1), code(r2), msg(r2)) == (400, m, 400, m) and (code(r3), msg(r3)) == (400, "데이터 모델을 찾을 수 없습니다.") and code(r4) == 200
    step("S3-3. 구조 변경 진단 예약은 데이터 소스 있는 모델만 (6.10-11a)", _struct_need_ds)

    def _dict():
        r1 = mk(dataModelId=FX["mL"])
        r2 = mk(dataModelId=FX["mL"], dictId=FX["DA"], scheduleNm="ZZB3사전지정_" + TAG)
        r3 = mk(diagType="STRUCT", dataModelId=FX["mS"], dictId=FX["DA"], scheduleNm="ZZB3구조사전_" + TAG)
        s2_, s3_ = r2.get("contents"), r3.get("contents")
        lst = {x["scheduleId"]: x for x in get(sa, "/api/diag/schedule/list")}
        up = post(sa, "/api/diag/schedule/update", {**base, "scheduleId": s2_, "scheduleNm": "ZZB3사전지정_" + TAG, "dataModelId": FX["mL"], "dictId": FX["DB"]})
        up2 = post(sa, "/api/diag/schedule/update", {**base, "scheduleId": s3_, "scheduleNm": "ZZB3구조사전_" + TAG, "diagType": "STRUCT",
                                                     "dataModelId": FX["mS"], "dictId": FX["DB"]})
        db = psql(f"select schedule_id||':'||coalesce(dict_id,'-') from quality.tb_diag_schedule where schedule_id in ('{s2_}','{s3_}') order by schedule_nm")
        print(f"  사전 없는 모델·사전 비움 → {code(r1)} {msg(r1)[:40]}… / 사전 지정 → {code(r2)} 목록 사전={lst[s2_].get('dictId')} {lst[s2_].get('dictNm')}"
              f"\n  구조 진단에 사전 → {code(r3)} 목록 사전={lst[s3_].get('dictId')} / 수정 {code(up)}·{code(up2)} → DB {db.splitlines()}")
        assert code(r1) == 400 and msg(r1).startswith("이 데이터 모델에 표준사전이 지정되어 있지 않아 예약 진단을 등록할 수 없습니다.")
        assert code(r2) == 200 and lst[s2_]["dictId"] == FX["DA"] and lst[s2_]["dictNm"] == "ZZB3사전A_" + TAG
        assert code(r3) == 200 and lst[s3_].get("dictId") is None
        assert sorted(db.splitlines()) == sorted([f"{s3_}:-", f"{s2_}:{FX['DB']}"])
        # 리뷰 지적 — 배치(준수 현황 스냅샷)·범정부 메타 유형도 사전을 버림, 사용 중지된 사전·없는 모델은 거부
        r4 = mk(diagType="COMPLIANCE_SNAPSHOT", dataModelId=None, dictId=FX["DA"], scheduleNm="ZZB3배치사전_" + TAG)
        s4_ = r4.get("contents")
        db4 = psql(f"select coalesce(dict_id,'-') from quality.tb_diag_schedule where schedule_id='{s4_}'")
        dx = make_dict(sa, "ZZB3사전중지_" + TAG, "N")
        psql(f"update quality.tb_std_dict set use_yn='N' where dict_id='{dx}'")
        r5 = mk(dataModelId=FX["mR"], dictId=dx, scheduleNm="ZZB3중지사전_" + TAG)
        up5 = post(sa, "/api/diag/schedule/update", {**base, "scheduleId": s2_, "scheduleNm": "ZZB3사전지정_" + TAG, "dataModelId": FX["mL"], "dictId": dx})
        r6 = mk(dataModelId="ZZB3없는모델", dictId=FX["DA"], scheduleNm="ZZB3없는모델_" + TAG)
        db2 = psql(f"select coalesce(dict_id,'-') from quality.tb_diag_schedule where schedule_id='{s2_}'")
        print(f"  준수 현황 스냅샷에 사전 → {code(r4)} DB={db4} / 사용 중지된 사전 등록 → {code(r5)} {msg(r5)} · 수정 → {code(up5)} {msg(up5)} (DB {db2})"
              f"\n  사전 고르고 없는 모델 → {code(r6)} {msg(r6)}")
        assert code(r4) == 200 and db4 == "-"
        assert code(r5) == 400 and msg(r5) == "사용할 수 없는 표준사전입니다." and code(up5) == 400 and msg(up5) == "사용할 수 없는 표준사전입니다." and db2 == FX["DB"]
        assert code(r6) == 400 and msg(r6) == "데이터 모델을 찾을 수 없습니다."
    step("S3-4. 스케줄 표준사전 — 표준화 진단만 저장(배치 유형도 버림), 모델에 사전이 없어도 스케줄에 고르면 등록, 사용 중지된 사전·없는 모델은 거부 (6.10-2)", _dict)

    def _busy():
        sb = make_schedule(sa, "ZZB3진행중모델_" + TAG, "STANDARD", FX["mBusy"])
        r = post(s_sm, "/api/diag/schedule/runNow", {"scheduleId": sb})
        n = len(logs_of(sb))
        print(f"  화면 진단(READY Job) 진행 중인 모델 즉시 실행 → {code(r)} {msg(r)} / LOG {n}건")
        assert code(r) == 409 and msg(r) == "이미 진행 중인 진단이 있습니다." and n == 0
    step("S3-5. 즉시 실행 — 같은 모델에 진행 중인 진단이 있으면 409, LOG 를 만들지 않음 (6.10-12a)", _busy)

    def _runnow():
        sid = FX["SR"]
        psql(f"update quality.tb_diag_schedule set repeat_time='02:00' where schedule_id='{sid}'")
        r = post(s_sm, "/api/diag/schedule/runNow", {"scheduleId": sid})
        last = sched_last(sid)
        lid = r.get("contents")
        lg = wait_until(lambda: [l for l in logs_of(sid) if l[0] == lid and l[1] != "RUNNING"], 60, 3)
        lg = [l for l in logs_of(sid) if l[0] == lid][0]
        jr = job_row(lg[4]) if lg[4] else None
        sst = make_schedule(sa, "ZZB3구조즉시_" + TAG, "STRUCT", FX["mS"])
        r2 = post(s_sm, "/api/diag/schedule/runNow", {"scheduleId": sst})
        lg2 = wait_until(lambda: [l for l in logs_of(sst) if l[1] != "RUNNING"], 60, 3)
        hs = psql(f"select status||'|'||cret_user_id from quality.tb_struct_diag_history where diag_id='{lg2[0][5]}'") if lg2 else None
        print(f"  즉시 실행 → {code(r)} / 직후 스케줄(상태|로그|시각)={last}\n  LOG={lg[1]} {lg[2]} / Job(상태|수집|사전|실행자)={jr}"
              f"\n  구조 진단 즉시 실행 → {code(r2)} / LOG={lg2[0][1] if lg2 else None} / 구조 진단(상태|실행자)={hs}")
        assert code(r) == 200 and last.split("|")[1] == lid and last.split("|")[0] in ("RUNNING", "DONE")
        assert lg[1] == "DONE" and lg[2] == "MANUAL" and jr == f"DONE|-|{FX['DB']}|{FX['sm']}"
        assert code(r2) == 200 and lg2 and lg2[0][1] == "DONE" and hs == f"DONE|{FX['sm']}"
    step("S3-6. 즉시 실행 — 실행자는 누른 사람, 기동 때 마지막 실행을 RUNNING 으로, 스케줄 사전, 수집 이력 없는 모델도 진단 (6.10-13b·9·2·10)", _runnow)


# ---------------------------------------------------------------- S4 용어 추천·변환 이력
def s4(sa, su, sm):
    DA = FX["DA"]
    W1, W2 = FX["W1"], FX["W2"]

    def item(nm, eng, words, dom="ZZB3문자30", **kw):
        it = {"termsNm": nm, "termsEngAbrvNm": eng, "termsDesc": nm, "domainNm": dom, "words": words, "originalInput": nm}
        it.update(kw)
        return it

    def _batch():
        ok = item("ZZB3가나", "ZZBA_ZZBB", [{"wordId": W1, "wordOrd": 0}, {"wordId": W2, "wordOrd": 1}])
        old = item("ZZB3가신규", "ZZBA_ZZBN", [{"wordId": W1, "wordOrd": 0}, {"wordId": None, "wordOrd": 1, "newWordIndex": 0}],
                   newWords=[{"wordNm": "ZZB3신규", "wordEngAbrvNm": "ZZBN", "wordEngNm": "ZZBNEW", "wordDesc": "x"}])
        nodom = item("ZZB3가도메인없음", "ZZBA_ZZBX", [{"wordId": W1, "wordOrd": 0}], dom="")
        nodom2 = item("ZZB3가도메인키없음", "ZZBA_ZZBY", [{"wordId": W1, "wordOrd": 0}])
        nodom2.pop("domainNm")
        longd = item("ZZB3가설명김", "ZZBA_ZZBL", [{"wordId": W1, "wordOrd": 0}], termsDesc="가" * 1100)
        dup = item("ZZB3가중복", "C_KR", [{"wordId": W1, "wordOrd": 0}])
        r = post(sa, "/api/std/registerTermsBatch", {"dictId": DA, "items": [ok, old, nodom, nodom2, longd, dup]})
        det = {d["termsNm"]: (d["status"], d.get("message")) for d in r.get("details", [])}
        newword = psql(f"select count(*) from quality.tb_word where dict_id='{DA}' and word_nm='ZZB3신규'")
        hist = psql("select string_agg(target_nm||'='||coalesce(remark,''), ' / ' order by seq) from quality.tb_change_history_detail where change_id = "
                    "(select change_id from quality.tb_change_history where change_type='BULK_INSERT' and change_user_id='space' order by change_dt desc limit 1)")
        print(f"  응답 키={sorted(r.keys())} 등록={r.get('registeredTerms')} 실패={r.get('failed')}")
        for k, v in det.items():
            print(f"  {k}: {v}")
        print(f"  신규 단어 생김={newword} / 일괄 등록 이력 상세={hist}")
        assert "registeredWords" not in r and r.get("registeredTerms") == 1 and r.get("failed") == 5
        assert det["ZZB3가나"][0] == "SUCCESS"
        assert det["ZZB3가신규"] == ("FAIL", "단어 ID를 확인할 수 없습니다. 단어 등록 후 다시 시도해주세요. (순서: 1)") and newword == "0"
        assert det["ZZB3가도메인없음"] == ("FAIL", "도메인명은 필수입니다.") and det["ZZB3가도메인키없음"] == ("FAIL", "도메인명은 필수입니다.")
        assert det["ZZB3가설명김"] == ("FAIL", "용어 등록 중 오류가 발생했습니다. 입력값을 확인해주세요.")
        assert det["ZZB3가중복"][0] == "FAIL" and det["ZZB3가중복"][1].startswith("이미 등록된 용어 영문약어입니다")
        assert "PSQLException" not in (hist or "") and "ZZB3가설명김=용어 등록 중 오류가 발생했습니다. 입력값을 확인해주세요." in (hist or "")
    step("S4-1. 용어 일괄 등록 — 신규 단어 함께 등록 경로 없음, 도메인 필수, DB 오류 원문 대신 안내 문구(이력 비고도) (7.5-6·7·8)", _batch)

    def _greedy():
        ctl = open(os.path.join(ROOT, "q-center", "src", "main", "java", "qualitycenter", "controller", "DataStandardController.java"), encoding="utf-8").read()
        tst = open(os.path.join(ROOT, "q-center", "src", "test", "java", "qualitycenter", "controller", "DataStandardHelperTest.java"), encoding="utf-8").read()
        ex = os.path.exists(os.path.join(ROOT, "q-executor", "src", "main", "java", "qualityexecutor", "service", "std", "TermRecommendService.java"))
        code_only = "\n".join(ln for ln in tst.splitlines() if not ln.strip().startswith("//"))
        print(f"  center greedySplit={'greedySplit' in ctl} / 시험(주석 제외) greedySplit={'greedySplit' in code_only} "
              f"newWords={'newWords' in code_only} / executor TermRecommendService={ex}")
        assert "greedySplit" not in ctl and "greedySplit" not in code_only and "newWords" not in code_only and not ex
    step("S4-2. 쓰지 않던 분리 함수·executor 용어 분석 서비스·신규 단어 시험 사본 정리 (7.5-9c·3e·6)", _greedy)

    def _excel():
        bad = os.path.join(TMP, "bad.xlsx")
        open(bad, "wb").write(b"not an excel file" * 10)
        wb = openpyxl.Workbook()
        wb.active["B1"] = "고객명"
        empty = os.path.join(TMP, "emptyA.xlsx")
        wb.save(empty)
        wb2 = openpyxl.Workbook()
        for i, v in enumerate(["컬럼명", "ZZB3고객명", "  ", "ZZB3주소"], 1):
            wb2.active.cell(row=i, column=1, value=v)
        good = os.path.join(TMP, "good.xlsx")
        wb2.save(good)
        res = []
        for p in (bad, empty, good):
            with open(p, "rb") as f:
                r = sa.post(BASE + "/api/std/parseExcelColumnNames", files={"file": (os.path.basename(p), f)}, timeout=30)
            res.append((r.status_code, r.text[:120]))
        print(f"  엑셀 아님={res[0]} / A열 빈={res[1]} / 정상={res[2]}")
        assert res[0][0] == 400 and res[1] == (200, "[]") and json.loads(res[2][1]) == ["ZZB3고객명", "ZZB3주소"]
    step("S4-3. 한글컬럼 엑셀 읽기 — 읽지 못하면 400 (빈 목록으로 삼키지 않음) (7.5-11)", _excel)

    def _resolve():
        u = FX["su"]
        rows = [("ZZB3변환A1", " ZZB3변환A1  ", "space"), ("ZZB3변환A2", "", "space"), ("ZZB3변환A3", "ZZB3변환B3", "space"), ("ZZB3변환A4", None, "space")]
        rows += [(f"ZZB3변환P{i:02d}", f"ZZB3변환Q{i:02d}", u) for i in range(1, 24)]
        psql("insert into quality.tb_term_resolve_history (input_kr_nm, resolved_kr_nm, resolved_en_nm, change_user_id, change_dt) values " +
             ",".join(f"({q(a)}, {q(b)}, 'ZZB3_EN', '{c}', to_char(now(),'YYYYMMDDHH24MISS'))" for a, b, c in rows))
        cleanup.append(lambda: psql("delete from quality.tb_term_resolve_history where input_kr_nm like 'ZZB3변환%'"))
        cnt = {}
        for yn in ("", "Y", "N"):
            cnt[yn] = [x["inputKrNm"] for x in post(sa, "/api/termResolve/list", {"inputKrNm": "ZZB3변환A", "nmChangedYn": yn})]
        r = sa.get(BASE + "/api/termResolve/excel", params={"inputKrNm": "ZZB3변환A"}, timeout=30)
        p = os.path.join(TMP, "resolve.xlsx")
        open(p, "wb").write(r.content)
        ws = openpyxl.load_workbook(p).active
        flag = {}
        for row in ws.iter_rows(values_only=True):
            row = [c for c in row]
            nm = next((c for c in row if isinstance(c, str) and c.startswith("ZZB3변환A")), None)
            if nm:
                flag[nm] = next((c for c in row if c in ("변경", "유지")), None)
        mine = post(FX["s_su"], "/api/termResolve/list", {"inputKrNm": "ZZB3변환"})
        print(f"  전체={sorted(cnt[''])}\n  변경={sorted(cnt['Y'])} / 유지={sorted(cnt['N'])}\n  엑셀 한글명 변경={flag} / 일반 사용자 조회={len(mine)}건")
        assert sorted(cnt["Y"]) == ["ZZB3변환A3"] and sorted(cnt["N"]) == ["ZZB3변환A1", "ZZB3변환A2", "ZZB3변환A4"]
        assert sorted(cnt["Y"] + cnt["N"]) == sorted(cnt[""])
        assert flag == {"ZZB3변환A1": "유지", "ZZB3변환A2": "유지", "ZZB3변환A3": "변경", "ZZB3변환A4": "유지"}
        assert len(mine) == 23 and all(m["changeUserId"] == u for m in mine)
    step("S4-4. 한글 변환 이력 '한글명 변경' — 앞뒤 공백 무시·B 가 비면 유지, 목록 필터·엑셀 같은 기준, 일반 사용자는 본인 것만 (7.5-13c)", _resolve)


# ---------------------------------------------------------------- S5 114 §2
def s5(sa, su, sm):
    DA = FX["DA"]

    def dom(nm, tp, ln=None):
        b = {"domainNm": nm, "domainGrpNm": "ZZB3그룹", "domainClsfNm": "ZZB3분류A", "domainDesc": "110 3묶음 시험", "dataType": tp, "dictId": DA}
        if ln is not None:
            b["dataLen"] = ln
        return b

    def _binary():
        r = {
            "BINARY 빈칸": post(sa, "/api/std/createDomain", dom("ZZB3이진빈", "BINARY")),
            "VARBINARY 0": post(sa, "/api/std/createDomain", dom("ZZB3가변이진0", "VARBINARY", 0)),
            "BINARY 16": post(sa, "/api/std/createDomain", dom("ZZB3이진16", "BINARY", 16)),
            "NUMBER 빈칸": post(sa, "/api/std/createDomain", dom("ZZB3숫자빈", "NUMBER")),
        }
        did = psql(f"select domain_id from quality.tb_domain where dict_id='{DA}' and domain_nm='ZZB3이진16'")
        up = post(sa, "/api/std/updateDomain", {**dom("ZZB3이진16", "BINARY", 0), "id": did})
        saved = psql(f"select string_agg(domain_nm||':'||data_type||':'||coalesce(data_len::text,''), ',' order by domain_nm) from quality.tb_domain "
                     f"where dict_id='{DA}' and domain_nm in ('ZZB3이진빈','ZZB3가변이진0','ZZB3이진16','ZZB3숫자빈')")
        for k, v in r.items():
            print(f"  등록 {k:10s} → {code(v)} {msg(v)}")
        print(f"  BINARY 를 길이 0 으로 수정 → {code(up)} {msg(up)} / 저장된 도메인 {saved}")
        assert (code(r["BINARY 빈칸"]), msg(r["BINARY 빈칸"])) == (500, "데이터 타입 BINARY 은(는) 데이터 길이를 1 이상으로 입력해야 합니다.")
        assert (code(r["VARBINARY 0"]), msg(r["VARBINARY 0"])) == (500, "데이터 타입 VARBINARY 은(는) 데이터 길이를 1 이상으로 입력해야 합니다.")
        assert code(r["BINARY 16"]) == 200 and code(r["NUMBER 빈칸"]) == 200 and code(up) == 500
        assert saved == "ZZB3숫자빈:NUMBER:0,ZZB3이진16:BINARY:16"
    step("S5-1. 도메인 길이 필수 타입에 BINARY·VARBINARY (등록·수정) (114 §2 권고 — 도메인 길이 필수 타입)", _binary)

    def _pending():
        s_su = FX.get("s_su") or login_api(FX["su"], FX["suPw"])
        r = post(s_su, "/api/std/createWord", {"wordNm": "ZZB3내단어", "wordEngAbrvNm": "ZZBP", "wordEngNm": "ZZBPMINE", "wordDesc": "x",
                                               "wordClsfYn": "N", "domainClsfNm": "", "allophSynmLst": [], "forbdnWordLst": [],
                                               "commStndYn": "N", "dictId": DA})
        n = psql(f"select count(*) from quality.tb_word where dict_id='{DA}' and word_nm='ZZB3내단어'")
        print(f"  남의 승인 대기 단어와 같은 영문약어 등록 → {code(r)} {msg(r)!r} / 생김={n}")
        assert code(r) == 500 and msg(r).startswith("이미 등록된 단어 영문약어입니다") and n == "0"
    step("S5-2. 승인 권한 없는 사용자 — 남의 승인 대기 단어와 영문약어가 겹치면 등록 때 거부·안내 (114 §2 권고 — 지금대로)", _pending)

    def _doc():
        t = open(os.path.join(ROOT, "dataQ설계", "114_결정기록_2026-10-28.md"), encoding="utf-8").read()
        rows = [ln for ln in t.splitlines() if ln.startswith("| 10-31 |") and "`babfac7` 후속" in ln]
        print(f"  114 §1 의 10-31 권고 채택 행 {len(rows)}개 / §2 이동 표시={'후속 5건은 10-31 반영해 §1 로 옮겼다' in t}")
        assert len(rows) == 5 and all("권고 채택" in r for r in rows) and "후속 5건은 10-31 반영해 §1 로 옮겼다" in t
        assert any("4.12-27" in r and "그대로" in r for r in rows)
    step("S5-3. 114 결정 기록 — 후속 5건 §1 권고 채택, 4.12-27 표기는 그대로 (114 §2 권고 — 4.12-27 표기)", _doc)

    def _spec():
        # 리뷰 지적 — 지운 화면·API 가 102 에 남고 개수가 코드와 어긋났다, 114 §3·110 조치 이력에 3묶음이 없었다
        d = os.path.join(ROOT, "dataQ설계")
        md = open(os.path.join(d, "102_기능명세서_2026-10-12.md"), encoding="utf-8").read()
        html = open(os.path.join(d, "102_기능명세서_2026-10-12.html"), encoding="utf-8").read()
        t114 = open(os.path.join(d, "114_결정기록_2026-10-28.md"), encoding="utf-8").read()
        t110 = open(os.path.join(d, "110_기존기능_보완후보_2026-10-12판.md"), encoding="utf-8").read()
        lines = md.splitlines()
        gone_rows = [ln[:12] for ln in lines if re.match(r"\| (DG-039|DG-04[0-3]|DG-05[89]|DG-06[01]|DG-B24) \|", ln)]
        appx = [ln for ln in lines[next(i for i, ln in enumerate(lines) if ln.startswith("## 부록 A")):] if re.match(r"\| q-(center|executor) \|", ln)]
        gone_api = [ln for ln in appx if re.search(r"getDiagResultList|/api/diag/schedule/\{scheduleId\}|/api/diag/schedule/logs/\{logId\}|compareSchema|latestStatus|analyzeTermsBatch` \| analyze.*q-executor", ln)]
        n_c = sum(1 for ln in appx if ln.startswith("| q-center"))
        n_e = sum(1 for ln in appx if ln.startswith("| q-executor"))
        dg = sum(1 for ln in lines if ln.startswith("| DG-"))
        fn_total = sum(1 for ln in lines if re.match(r"\| [A-Z]+-[A-Z]?\d+ \|", ln))
        row6 = [ln for ln in lines if ln.startswith("| 6 | 진단")]
        r110 = [ln for ln in t110.splitlines() if ln.startswith("| 10-31 | 6.10-") or ln.startswith("| 10-31 | 7.5-")]
        print(f"  102 지운 기능 행={gone_rows} 지운 API 행={len(gone_api)} / 부록 A q-center {n_c}·q-executor {n_e} / DG 행 {dg}, 기능 행 {fn_total}"
              f"\n  6장 줄={row6[0] if row6 else None}\n  HTML 지운 ID={bool(re.search(r'DG-0(39|4[0-3]|5[89]|6[01])|DG-B24', html))} 322={'<b>322</b>' in html}"
              f"\n  114 §3 6.10-3b={'| 110 6.10-3b |' in t114} / 110 조치 이력 10-31 6.10·7.5 행 {len(r110)}개, 6.10-3b 결정 대기="
              f"{any(ln.startswith('| 10-31 | 6.10-3b') and '결정 대기' in ln for ln in r110)}")
        assert not gone_rows and not gone_api and (n_c, n_e) == (310, 12) and dg == 90
        assert "메서드·경로 조합 322개" in md and "q-center 310개, q-executor 12개" in md and "| | **합계** | **623** | |" in md
        assert row6 and row6[0].endswith("| 90 | 24 + 내부 4 |") and fn_total == 623
        assert "DSSchemaCompare.vue | 어느 파일에서도" not in md and "구조 진단 실행·스키마 비교" not in md
        assert not re.search(r"DG-0(39|4[0-3]|5[89]|6[01])|DG-B24", html) and "<b>322</b>" in html
        assert "| 110 6.10-3b |" in t114 and len(r110) >= 50
        assert any(ln.startswith("| 10-31 | 6.10-3b") and "결정 대기" in ln for ln in r110)
    step("S5-4. 문서 — 102 지운 화면·API 행 없음·API 322(310·12)·기능 623, HTML 다시 만듦, 114 §3 6.10-3b, 110 조치 이력 3묶음 (리뷰 지적)", _spec)


# ---------------------------------------------------------------- 화면(셀레니움)
VM = ("const vm = (tab, name) => { const e = [...document.querySelectorAll('#' + tab + ' *')].find(x => x.__vue__); "
      "const down = v => { if (v.$options.name === name) return v; for (const ch of v.$children) { const r = down(ch); if (r) return r; } return null; }; "
      "let c = e && e.__vue__; while (c) { const d = down(c); if (d) return d; c = c.$parent; } return null; };")
NAV = ("const nav = (() => { const e = document.querySelector('#nav_dashboard'); let v = e && e.__vue__; "
       "while (v && v.$options.name !== 'NdNav') v = v.$parent; return v; })();")
HOOK = ("if (!window.__errs) { window.__errs = []; const oe = console.error; console.error = function() { window.__errs.push([...arguments].map(String).join(' ')); "
        "oe.apply(console, arguments); }; window.addEventListener('error', e => window.__errs.push('onerror ' + e.message)); "
        "window.addEventListener('unhandledrejection', e => window.__errs.push('rejection ' + (e.reason && e.reason.message || e.reason))); }")


class UI:
    inst = None

    @classmethod
    def get(cls):
        if not cls.inst:
            cls.inst = UI()
        return cls.inst

    def __init__(self):
        from selenium import webdriver
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from edgeopt import headless
        opts = webdriver.EdgeOptions()
        opts.add_argument("--log-level=3")
        self.drv = webdriver.Edge(options=headless(opts))
        self.drv.set_window_size(1600, 1000)
        self.By, self.Wait = By, WebDriverWait
        self.errs = []
        self.user = None

    def js(self, s, *a):
        return self.drv.execute_script(s, *a)

    def login(self, user, pwd=None):
        if self.user == user:
            return
        if self.user:
            self.errs += self.js("return window.__errs || []") or []
        By = self.By
        self.drv.delete_all_cookies()
        self.drv.get(BASE)
        self.Wait(self.drv, 15).until(lambda d: d.find_elements(By.CSS_SELECTOR, "input[type='text']"))
        self.drv.find_element(By.CSS_SELECTOR, "input[type='text']").send_keys(user)
        self.drv.find_element(By.CSS_SELECTOR, "input[type='password']").send_keys(pwd or PW)
        self.drv.find_element(By.CSS_SELECTOR, "button[type='submit'], .v-btn").click()
        self.Wait(self.drv, 15).until(lambda d: "/app/main" in d.current_url)
        time.sleep(3)
        self.js(HOOK)
        self.user = user

    def open(self, key, title):
        self.js(NAV + "nav.addTabItem(arguments[0], arguments[1]);", title, key)
        time.sleep(2.5)

    def vm(self, tab, name, expr, *a):
        return self.js(VM + f"const c = vm('{tab}', '{name}'); " + expr, *a)

    def click_text(self, css, text, exact=True):
        els = [e for e in self.drv.find_elements(self.By.CSS_SELECTOR, css)
               if e.is_displayed() and ((e.text.strip() == text) if exact else (text in e.text))]
        assert els, f"'{text}' 없음 ({css})"
        self.js("arguments[0].click();", els[0])
        time.sleep(1)

    def snack(self, has="", t=10):
        def _t(d):
            for e in d.find_elements(self.By.CSS_SELECTOR, ".v-snack__content"):
                if e.is_displayed() and has in e.text:
                    return e.text.strip()
            return None
        return self.Wait(self.drv, t).until(_t)

    def swal(self, has="", t=10):
        """보이는 확인창의 제목\n본문 (아이콘 글자·버튼은 뺀다)"""
        def _t(d):
            for e in d.find_elements(self.By.CSS_SELECTOR, ".swal2-popup"):
                if e.is_displayed() and has in e.text:
                    parts = [x.text.strip() for x in e.find_elements(self.By.CSS_SELECTOR, ".swal2-title, .swal2-html-container")
                             if x.is_displayed() and x.text.strip()]
                    return "\n".join(parts)
            return None
        return self.Wait(self.drv, t).until(_t)

    def swal_close(self):
        self.js("document.querySelectorAll('.swal2-confirm').forEach(b => b.click());")
        time.sleep(0.8)

    def shot(self, nm):
        try:
            self.drv.save_screenshot(os.path.join(TMP, nm + ".png"))
        except Exception:
            pass

    def quit(self):
        try:
            self.errs += self.js("return window.__errs || []") or []
        except Exception:
            pass
        self.drv.quit()


def s6(sa, su, sm):
    ui = UI.get()
    By = ui.By
    ui.login("space")
    mu = FX["mU"]
    T = "tab_dataDiagResult"
    N = "DSDataDiagResult"

    def _result():
        r = post(sa, "/api/diag/startDiag", {"dataModelId": mu, "dictId": FX["DC"]})
        FX["J2"] = r.get("contents")
        wait_job(FX["J2"])
        add_attrs(mu, "zzown", "ZZ_T1", [("C_ADD", "나중추가", "VARCHAR", 10, None)])   # 진단 뒤 모델이 바뀜
        ui.open("dataDiagResult", "진단 결과")
        ui.vm(T, N, "c.loadJobList();")
        time.sleep(2)
        texts = ui.vm(T, N, "return c.jobList.filter(j => j.dataModelId === arguments[0]).map(j => [j.diagJobId, j.jobDisplayText, j.startDt, j.resultCnt]);", mu)
        ui.vm(T, N, "c.selectJob(arguments[0]);", FX["J1"])
        time.sleep(2.5)
        body = ui.drv.find_element(By.ID, T).text
        ov = ui.vm(T, N, "return c.overviewStats;")
        ui.shot("u1_overview")
        m = re.search(r"전체 \d+개 컬럼 중 \d+개 준수", body)
        print(f"  Job 표시={[(t[1]) for t in texts]}\n  요약={ov} / 화면 '{m.group(0) if m else None}'")
        assert len(texts) == 2 and texts[0][1] != texts[1][1]
        assert all(t[2] and t[1] == f"[{FX['mUNm']}] {t[2]} - 완료 (이슈 {t[3]}건)" for t in texts)
        assert ov["totalColCnt"] == 6 and ov["issueColCnt"] == 4 and ov["okColCnt"] == 2 and ov["complianceRate"] == 33.3
        assert "33.3%" in body and "전체 6개 컬럼 중 2개 준수" in body
    step("U1-1. 진단 결과 요약 — 그 Job 의 컬럼 수·이슈 컬럼 수·준수율(소수 1자리), Job 목록은 Job 마다 다른 시각 (6.10-14·20b)", _result)

    def row_of(col):
        return ui.js(f"return [...document.querySelectorAll('#{T} tbody tr')].find(tr => tr.offsetParent && "
                     "[...tr.cells].some(td => td.innerText.split(/\\s+/).includes(arguments[0])));", col)

    def scripts(dbms):
        psql(f"update ndata.tb_data_source set dbms_tp='{dbms}' where ds_id='{FX['dsUi']}'")
        ui.vm(T, N, "c.loadJobList();")
        time.sleep(1.5)
        ui.vm(T, N, "c.activeTab = 2;")
        time.sleep(1)
        out = {}
        for col, btn in (("C_KR", "코맨트 변경"), ("C_NUM", "컬럼 변경")):
            row = row_of(col)
            assert row, f"{col} 행 없음"
            b = [x for x in row.find_elements(By.CSS_SELECTOR, "button") if x.text.strip() == btn]
            assert b, f"{col} {btn} 버튼 없음 ({row.text!r})"
            ui.js("arguments[0].click();", b[0])
            time.sleep(1)
            pre = [p.text for p in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active pre") if p.is_displayed()]
            out[col] = pre[0] if pre else None
            ui.js("document.querySelectorAll('.v-dialog--active button').forEach(b => { if (b.innerText.trim() === '닫기') b.click(); });")
            ui.vm(T, N, "c.commentDialog = false; c.modifyDialog = false;")
            time.sleep(0.5)
        gen = ui.vm(T, N, "return c.detailList.filter(d => ['C_KR','C_NUM','C_DT'].includes(d.attrNm)).map(d => c.generateCommentScript(d));")
        return out, gen

    def _scripts():
        res = {}
        try:
            for dbms in ("PostgreSQL", "Cubrid", "MariaDB"):
                res[dbms] = scripts(dbms)
                print(f"  {dbms}: 코맨트={res[dbms][0]['C_KR']!r}\n           컬럼 변경={res[dbms][0]['C_NUM']!r}\n           일괄 코멘트={res[dbms][1]}")
        finally:
            psql(f"update ndata.tb_data_source set dbms_tp='PostgreSQL' where ds_id='{FX['dsUi']}'")
        pg, cub, mar = res["PostgreSQL"], res["Cubrid"], res["MariaDB"]
        assert pg[0]["C_KR"] == "COMMENT ON COLUMN zzown.ZZ_T1.C_KR IS 'ZZB3고객명';"
        assert pg[0]["C_NUM"].endswith("ALTER TABLE ZZ_T1 ALTER COLUMN C_NUM TYPE NUMBER(10,2);")
        assert cub[0]["C_KR"] == "-- CUBRID: COMMENT 는 ALTER 문에 직접 지원하지 않음. 카탈로그 별도 갱신 필요. (ZZ_T1.C_KR = 'ZZB3고객명')"
        assert cub[0]["C_NUM"].endswith("ALTER TABLE ZZ_T1 MODIFY C_NUM NUMBER(10,2);") and all(g.startswith("-- CUBRID") for g in cub[1])
        assert mar[0]["C_KR"] == "ALTER TABLE ZZ_T1 MODIFY COLUMN C_KR VARCHAR(30) COMMENT 'ZZB3고객명';"
        assert "ALTER TABLE ZZ_T1 MODIFY COLUMN C_NUM NUMBER(8,2) COMMENT 'ZZB3금액';" in mar[1]
        assert "ALTER TABLE ZZ_T1 MODIFY COLUMN C_DT DATETIME COMMENT 'ZZB3가까운일시';" in mar[1]
    step("U1-2. 코맨트·컬럼 변경 스크립트 — 스키마는 소유자, 소수 자릿수, CUBRID 문법·안내, MySQL 계열 실제 타입 (6.10-18·19b·19c·19d)", _scripts)

    def _termreg():
        ui.vm(T, N, "c.activeTab = 2;")
        time.sleep(1)
        row = row_of("ZQX_ZQY")
        btns = [x.text.strip() for x in row.find_elements(By.CSS_SELECTOR, "button")]
        ui.js("arguments[0].click();", [x for x in row.find_elements(By.CSS_SELECTOR, "button") if x.text.strip() == "용어 등록"][0])
        time.sleep(2.5)
        dlg = [d for d in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active") if "용어 빠른 등록" in d.text][0]
        cbs = dlg.find_elements(By.CSS_SELECTOR, "input[type='checkbox']")
        first = (len(cbs), cbs[0].is_selected() if cbs else None, "도메인 분류명" in dlg.text)
        ui.shot("u1_termreg")
        ui.js("arguments[0].click();", cbs[0])
        time.sleep(0.7)
        off = (cbs[0].is_selected(), "도메인 분류명" in dlg.text, ui.vm(T, N, "return c.termRegWords[1].domainClsfNm;"))
        ui.js("arguments[0].click();", cbs[0])
        time.sleep(0.7)
        ui.vm(T, N, "const w = c.termRegWords[1]; w.wordNm = 'ZZB3와이'; w.wordEngNm = 'ZQYNAME';")
        time.sleep(0.5)
        regs = [b for b in dlg.find_elements(By.CSS_SELECTOR, "button") if b.text.strip() == "단어 등록"]
        ui.js("arguments[0].click();", regs[-1])
        warn = ui.swal("형식 단어")
        ui.swal_close()
        items = ui.vm(T, N, "return c.termRegDomainClsfItems;")
        ui.vm(T, N, "c.termRegWords[1].domainClsfNm = 'ZZB3분류C';")
        time.sleep(0.5)
        ui.js("arguments[0].click();", regs[-1])
        time.sleep(2.5)
        try:
            ui.swal_close()
        except Exception:
            pass
        saved = psql(f"select word_clsf_yn||'|'||coalesce(domain_clsf_nm,'') from quality.tb_word where dict_id='{FX['DC']}' and word_eng_abrv_nm='ZQY'")
        ui.vm(T, N, "c.termRegDialog = false;")
        print(f"  행 버튼={btns} / 처음(체크박스 수, 체크, 분류명 칸)={first} / 끈 뒤(체크, 분류명 칸, 값)={off}\n  분류명 없이 등록 → {warn!r}"
              f"\n  분류명 목록={items} / 등록된 단어(형식|분류명)={saved}")
        assert btns == ["용어 등록"] and first == (1, True, True) and off == (False, False, "")
        assert "형식 단어는 도메인 분류명이 필수입니다." in warn and "ZZB3분류C" in items and saved == "Y|ZZB3분류C"
    step("U1-3. 용어 빠른 등록 — 남의 사전 용어는 매칭 안 됨(용어 등록 버튼), 마지막 미등록 단어 형식단어 체크(사전 강제 시 기본)·도메인 분류명 필수 (6.10-15·17)", _termreg)

    TD, ND = "tab_dataDiag", "DSDataDiag"

    def _diag_screen():
        jx = f"ZB3JX{TAG}"
        psql(f"insert into quality.tb_diag_job (diag_job_id, dm_id, status, total_cnt, process_cnt, cret_dt, start_dt, cret_user_id, dict_id) values "
             f"('{jx}', '{FX['mOv']}', 'RUNNING', 10, 3, to_char(now(),'YYYYMMDDHH24MISS'), to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}')")
        ui.open("dataDiag", "진단 실행")
        time.sleep(1)
        cur = ui.vm(TD, ND, "return c.currentJob && c.currentJob.diagJobId;")
        box = [e for e in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TD} input") if e.get_attribute("placeholder") == "데이터 모델 선택"][0]
        box.click()
        box.send_keys(FX["mRNm"])
        time.sleep(1.2)
        ui.click_text(".menuable__content__active .v-list-item", FX["mRNm"])
        time.sleep(1)
        dict_id = ui.vm(TD, ND, "return c.dictId;")
        dict_txt = [e.text for e in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TD} .v-select__selection") if e.is_displayed()]
        btn = [b for b in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TD} button") if b.text.strip() == "진단 시작"][0]
        en_free = btn.is_enabled()
        ui.vm(TD, ND, "c.selectedModel = arguments[0]; c.onModelChange(arguments[0]);", FX["mOv"])
        time.sleep(0.5)
        en_busy = btn.is_enabled()
        ui.vm(TD, ND, "c.selectedModel = arguments[0]; c.onModelChange(arguments[0]);", FX["mR"])
        time.sleep(0.5)
        ui.shot("u2_diag")
        before = job_ids(FX["mR"])
        # 리뷰 지적 — 목록을 다시 읽기 전에 버튼(starting)이 풀리면 두 번 눌러 409 가 난다.
        # 목록 읽기가 끝난 횟수를 세어 두고, starting 이 false 가 되는 순간의 횟수를 남긴다 (1 이상이어야 한다)
        ui.vm(TD, ND, "const orig = c.loadJobList; window.__zzLj = 0; window.__zzLjAtRelease = null; window.__zzOrigLj = orig;"
                      "c.loadJobList = function () { return orig.call(c).then(r => { window.__zzLj++; return r; }); };"
                      "const un = c.$watch('starting', v => { if (!v) { window.__zzLjAtRelease = window.__zzLj; un(); } });")
        ui.js("arguments[0].click();", btn)
        sn = ui.snack("진단 완료", 30)
        starting = ui.vm(TD, ND, "c.loadJobList = window.__zzOrigLj; return window.__zzLjAtRelease;")
        newj = one_new(before, job_ids(FX["mR"]))
        time.sleep(1)
        # 새 Job 이 끝나면(스낵바는 추적 중인 Job 에서만 뜬다) 목록을 다시 읽고 남은 진행 중 Job(jx)을 다시 추적한다
        after = ui.vm(TD, ND, "return [(c.jobList.find(j => j.diagJobId === arguments[0]) || {}).status, !!c.pollTimer && c.currentJob.diagJobId === arguments[1]];", newj, jx)
        print(f"  다른 모델 진행 중 추적 Job={cur == jx} / 모델 고른 뒤 사전={dict_id} 표시={dict_txt}\n  진단 시작 버튼: 다른 모델 진행 중={en_free}, 고른 모델 진행 중={en_busy}, 버튼이 풀릴 때 목록 다시 읽은 횟수={starting}"
              f"\n  스낵바={sn!r} / 끝난 뒤 (새 Job 목록 상태, 남은 진행 중 Job 추적)={after}")
        psql(f"delete from quality.tb_diag_job where diag_job_id='{jx}'")
        assert cur == jx and dict_id == FX["DA"] and ("ZZB3사전A_" + TAG) in dict_txt
        assert en_free and not en_busy and (starting or 0) >= 1 and sn.startswith("진단 완료 - 이슈") and after == ["DONE", True]
    step("U2-1. 진단 실행 — 모델 고르면 모델 사전, 다른 모델 진행 중이어도 시작 가능·고른 모델 진행 중이면 막음, 시작한 Job 추적, 끝나면 남은 진행 중 Job 추적 (6.10-21·22b)", _diag_screen)

    def _poll_ready():
        jy = f"ZB3JY{TAG}"
        psql(f"insert into quality.tb_diag_job (diag_job_id, dm_id, status, cret_dt, cret_user_id, dict_id) values "
             f"('{jy}', '{FX['mOrph']}', 'READY', to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}')")
        ui.vm(TD, ND, "c.pollJob(arguments[0]); c.startPoll(arguments[0]);", jy)
        time.sleep(4.5)
        on_ready = ui.vm(TD, ND, "return !!c.pollTimer;")
        psql(f"update quality.tb_diag_job set status='RUNNING', total_cnt=5, process_cnt=2 where diag_job_id='{jy}'")
        time.sleep(3.5)
        on_run = ui.vm(TD, ND, "return !!c.pollTimer && c.currentJob.status;")
        psql(f"update quality.tb_diag_job set status='ERROR' where diag_job_id='{jy}'")
        sn = ui.snack("진단 오류", 10)
        time.sleep(1.5)
        # 목록을 다시 읽으면 다른 RUNNING Job(시험의 가짜 Job 포함)이 있으면 그것을 추적하므로, 이 Job 추적이 끝났는지만 본다
        off = ui.vm(TD, ND, "return [!!c.pollTimer && c.currentJob.diagJobId === arguments[0], (c.jobList.find(j => j.diagJobId === arguments[0]) || {}).status];", jy)
        ui.shot("u2_error")
        print(f"  READY 첫 조회 뒤 폴링 계속={on_ready} / RUNNING={on_run} / ERROR → 스낵바={sn!r}, (이 Job 폴링 중, 목록 상태)={off}")
        assert on_ready and on_run == "RUNNING" and sn == "진단 오류 발생" and off == [False, "ERROR"]
    step("U2-2. 진단 실행 — 첫 조회가 READY 여도 폴링 계속, ERROR 로 끝나면 '진단 오류 발생'·목록 다시 읽기 (6.10-22a)", _poll_ready)

    def _untracked():
        # 리뷰 지적 — 추적은 Job 하나(B)뿐인데 다른 모델 Job(A RUNNING, R READY)이 끝나도 목록이 안 바뀌어
        # 그 모델의 '진단 시작' 이 탭을 다시 열 때까지 막혀 있었다
        ja, jb2, jr = f"ZB3UA{TAG}", f"ZB3UB{TAG}", f"ZB3UR{TAG}"
        psql(f"insert into quality.tb_diag_job (diag_job_id, dm_id, status, total_cnt, process_cnt, cret_dt, start_dt, cret_user_id, dict_id) values "
             f"('{ja}', '{FX['mOv']}', 'RUNNING', 10, 3, to_char(now(),'YYYYMMDDHH24MISS'), to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}'),"
             f"('{jb2}', '{FX['mAuto']}', 'RUNNING', 10, 3, to_char(now(),'YYYYMMDDHH24MISS'), to_char(now(),'YYYYMMDDHH24MISS'), 'space', '{FX['DA']}'),"
             f"('{jr}', '{FX['mDraft']}', 'READY', 0, 0, to_char(now(),'YYYYMMDDHH24MISS'), null, 'space', '{FX['DA']}')")
        try:
            ui.vm(TD, ND, "c.stopPoll(); c.loadJobList();")
            time.sleep(1)
            ui.vm(TD, ND, "c.startPoll(arguments[0]); c.pollJob(arguments[0]);", jb2)
            time.sleep(1)
            ui.vm(TD, ND, "c.selectedModel = arguments[0]; c.onModelChange(arguments[0]);", FX["mOv"])
            time.sleep(0.5)
            btn = [b for b in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TD} button") if b.text.strip() == "진단 시작"][0]
            en_before = btn.is_enabled()
            psql(f"update quality.tb_diag_job set status='DONE', process_cnt=10, end_dt=to_char(now(),'YYYYMMDDHH24MISS') where diag_job_id in ('{ja}','{jr}')")
            time.sleep(4.5)       # 추적 Job 폴링 한 번(3초) + 목록 다시 읽기
            st = ui.vm(TD, ND, "const s = id => (c.jobList.find(j => j.diagJobId === id) || {}).status; "
                               "return [s(arguments[0]), s(arguments[1]), c.currentJob.diagJobId === arguments[2], c.selectedModelBusy];", ja, jr, jb2)
            en_after = btn.is_enabled()
            row_btn = ui.js("const tr = [...document.querySelectorAll('#" + TD + " tbody tr')].filter(r => r.innerText.includes(arguments[0])); "
                            "return tr.map(r => [...r.querySelectorAll('button')].map(b => b.innerText.trim()).join(','));", FX["mOvNm"])
            ui.shot("u2_untracked")
            psql(f"update quality.tb_diag_job set status='DONE', process_cnt=10, end_dt=to_char(now(),'YYYYMMDDHH24MISS') where diag_job_id='{jb2}'")
            sn = ui.snack("진단 완료", 10)
            time.sleep(1.5)
            fin = ui.vm(TD, ND, "return [(c.jobList.find(j => j.diagJobId === arguments[0]) || {}).status, "
                                "c.jobList.filter(j => j.status === 'READY' || j.status === 'RUNNING').map(j => j.diagJobId), "
                                "c.pollTimer ? c.currentJob.diagJobId : null];", jb2)
        finally:
            psql(f"delete from quality.tb_diag_job where diag_job_id in ('{ja}','{jb2}','{jr}')")
        busy_left = fin[1]
        print(f"  B 추적 중 A 모델 '진단 시작' = {en_before} → A·R 이 끝난 뒤 {en_after} / (A, R 목록 상태, 추적 B 그대로, A 모델 막힘)={st}"
              f"\n  A 행 버튼={row_btn} / B 끝 → 스낵바={sn!r}, B 목록={fin[0]}, 남은 대기·진행 Job={busy_left}, 추적={fin[2]}")
        assert en_before is False and en_after is True and st == ["DONE", "DONE", True, False]
        assert any("결과보기" in b for b in row_btn)
        assert sn.startswith("진단 완료 - 이슈") and fin[0] == "DONE"
        # 남은 대기·진행 중 Job 이 있으면 그 하나를 추적하고, 없으면 추적하지 않는다
        assert (fin[2] in busy_left) if busy_left else fin[2] is None
    step("U2-3. 진단 실행 — 추적하지 않는 다른 모델 Job·대기 Job 이 끝나면 목록을 다시 읽어 그 모델 '진단 시작' 이 풀리고 결과보기가 생김, 추적 Job 이 끝나면 남은 Job 추적 (6.10-22b)", _untracked)

    def _struct409():
        hs = f"ZB3HR{TAG}"
        psql(f"insert into quality.tb_struct_diag_history (diag_id, data_model_id, status, cret_user_id) values ('{hs}', '{FX['mS']}', 'READY', 'space')")
        try:
            ui.open("structDiag", "진단 실행")
            ui.vm("tab_structDiag", "DSStructDiag", "c.selectedModel = arguments[0];", FX["mS"])
            time.sleep(0.8)
            ui.click_text("#tab_structDiag button", "진단 실행")
            sn = ui.snack("진단 요청 실패")
            ex = ui.vm("tab_structDiag", "DSStructDiag", "return c.executing;")
        finally:
            psql(f"delete from quality.tb_struct_diag_history where diag_id='{hs}'")
        print(f"  진행 중 구조 진단 → 스낵바={sn!r} / 진행 표시 꺼짐={not ex}")
        assert sn == "진단 요청 실패: 이미 진행 중인 구조 진단이 있습니다." and not ex
    step("U3-1. 구조 진단 — 진행 중이면 서버 사유(409)를 스낵바에 (6.10-28)", _struct409)

    TS, NS = "tab_scheduleManage", "DSScheduleManage"

    def _schedule_screen():
        long_nm = "ZZB3긴스케줄이름_" + TAG + "_" + "가나다라마바사아자차카타파하" * 4
        make_schedule(sa, long_nm, "STANDARD", FX["mR"], dict_id=FX["DB"])
        ui.open("scheduleManage", "스케줄 관리")
        time.sleep(1)
        total = ui.vm(TS, NS, "return c.filteredItems.length;")
        ui.vm(TS, NS, "c.itemsPerPage = 10;")
        time.sleep(1)
        pager = ui.drv.find_element(By.ID, "schPager")
        pages = [b.text for b in pager.find_elements(By.CSS_SELECTOR, ".v-pagination__item")]
        ui.js("arguments[0].click();", pager.find_elements(By.CSS_SELECTOR, ".v-pagination__item")[1])
        time.sleep(0.8)
        p2 = ui.vm(TS, NS, "return c.page;")
        ui.vm(TS, NS, "c.filterUseYn = 'N';")
        time.sleep(0.8)
        p_after = ui.vm(TS, NS, "return c.page;")
        ui.vm(TS, NS, "c.filterUseYn = ''; c.itemsPerPage = 30;")
        ui.shot("u4_schedule")
        # 등록 창 — 표준사전 칸
        ui.click_text(f"#{TS} button", "스케줄 추가")
        time.sleep(1)
        vis = {}
        for t in ("STANDARD", "STRUCT", "BOTH", "GOV_META"):
            ui.vm(TS, NS, "c.form.diagType = arguments[0];", t)
            time.sleep(0.5)
            el = ui.drv.find_elements(By.ID, "schDict")
            vis[t] = bool(el) and el[0].is_displayed()
        ui.vm(TS, NS, "c.form.diagType = 'STANDARD';")
        time.sleep(0.4)
        ph = ui.js("const e = document.getElementById('schDict'); return e ? (e.getAttribute('placeholder') || (e.closest('.v-input').querySelector('input') || {}).placeholder) : null;")
        lab = ui.js("const e = document.getElementById('schDict'); return e ? e.closest('.v-input').querySelector('label').innerText : null;")
        ui.shot("u4_schedule_dialog")
        warns = []
        base = "c.form.scheduleNm = 'ZZB3화면'; c.form.dataModelId = arguments[0]; c.form.scheduleType = 'SIMPLE';"
        for setv in ("c.form.repeatCycle = 'DAILY'; c.form.repeatTime = '24:00';",
                     "c.form.repeatCycle = 'WEEKLY'; c.form.repeatTime = '02:00'; c.form.repeatDayOfWeek = 8;",
                     "c.form.repeatCycle = 'MONTHLY'; c.form.repeatDayOfMonth = 29;",
                     "c.form.repeatCycle = 'MONTHLY'; c.form.repeatDayOfMonth = 2.5;"):
            ui.vm(TS, NS, base + setv + "c.saveSchedule();", FX["mR"])
            warns.append(ui.swal("").splitlines()[0])
            ui.swal_close()
        made = psql("select count(*) from quality.tb_diag_schedule where schedule_nm='ZZB3화면'")
        ui.vm(TS, NS, "c.editDialog = false;")
        time.sleep(0.5)
        # 수정 창 — 저장된 사전이 보인다
        item = ui.vm(TS, NS, "const it = c.items.find(x => x.scheduleNm === arguments[0]); c.openEditDialog(it); return [c.form.dictId, it.dictNm];", long_nm)
        time.sleep(1)
        shown = [e.text for e in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active .v-select__selection") if e.is_displayed()]
        ui.vm(TS, NS, "c.editDialog = false;")
        print(f"  전체 {total}건, 10건씩 페이지={pages} → 2쪽 클릭 page={p2}, 필터 바꾸면 page={p_after}\n  표준사전 칸 보임={vis} 라벨={lab!r} 안내={ph!r}"
              f"\n  저장 경고={warns} / 만들어진 스케줄={made}\n  수정 창 사전={item} 표시={shown}")
        assert total >= 11 and len(pages) >= 2 and p2 == 2 and p_after == 1
        assert vis == {"STANDARD": True, "STRUCT": False, "BOTH": True, "GOV_META": False} and lab == "표준사전" and ph == "모델 기본"
        assert warns == ["실행 시각 형식 오류 (HH:mm)", "요일은 1(월)~7(일) 이어야 합니다.", "일자는 1~28 이어야 합니다.", "일자는 1~28 이어야 합니다."] and made == "0"
        assert item == [FX["DB"], "ZZB3사전B_" + TAG] and ("ZZB3사전B_" + TAG) in shown
    step("U4-1. 스케줄 관리 — 페이지 이동, 표준화 진단만 표준사전 칸(모델 기본), 실행 시각·요일·일자 검사, 수정 창에 저장된 사전 (6.10-33·2·11b)", _schedule_screen)

    def _schedule_dict_off():
        # 리뷰 지적 — 저장된 사전이 사용 중지되면 칸은 '모델 기본' 으로 보이는데 저장하면 그 사전이 다시 저장됐다
        dnm = "ZZB3중지될사전_" + TAG + "_" + "가나다라마바사" * 3
        dy = make_dict(sa, dnm, "N")
        snm = "ZZB3사전중지스케줄_" + TAG
        sid = make_schedule(sa, snm, "STANDARD", FX["mR"], dict_id=dy)
        psql(f"update quality.tb_std_dict set use_yn='N' where dict_id='{dy}'")
        ui.vm(TS, NS, "c.loadDictList(); c.loadList();")
        time.sleep(1.5)
        got = ui.vm(TS, NS, "const it = c.items.find(x => x.scheduleNm === arguments[0]); c.openEditDialog(it); return [c.form.dictId, c.dictClearedNm];", snm)
        time.sleep(1)
        hint = [e.text.strip() for e in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active .sch-batch-hint") if e.is_displayed()]
        shown = [e.text for e in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active .v-select__selection") if e.is_displayed()]
        ui.drv.set_window_size(1366, 768)
        time.sleep(0.5)
        ui.shot("u4_schedule_dict_off_1366")
        ui.drv.set_window_size(1600, 1000)
        ui.click_text(".v-dialog--active button", "저장")
        sw = ui.swal("").splitlines()[0]
        ui.swal_close()
        db = psql(f"select coalesce(dict_id,'-') from quality.tb_diag_schedule where schedule_id='{sid}'")
        # 사전을 다시 고르면 안내가 사라진다
        ui.vm(TS, NS, "const it = c.items.find(x => x.scheduleId === arguments[0]); c.openEditDialog(it); c.form.dictId = arguments[1];", sid, FX["DA"])
        time.sleep(0.6)
        hint2 = [e.text for e in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active .sch-batch-hint") if e.is_displayed()]
        ui.vm(TS, NS, "c.editDialog = false;")
        print(f"  사용 중지된 사전 스케줄 수정 창 (form.dictId, 비운 사전)={got}\n  안내={hint} / 표시={shown}\n  저장 → {sw!r}, DB 사전={db} / 사전을 고르면 안내={hint2}")
        assert got == [None, dnm] and hint == [f"사용 중지된 표준사전({dnm})이라 비웠습니다. 저장하면 모델 기본 사전으로 진단합니다."]
        assert sw == "저장됨" and db == "-" and hint2 == []
    step("U4-2. 스케줄 관리 — 저장된 사전이 사용 중지되면 수정 창에서 비우고 안내, 저장하면 모델 기본 사전 (6.10-2)", _schedule_dict_off)

    TL, NL = "tab_scheduleLog", "DSScheduleLog"

    def _log_screen():
        sid = FX["SR"]
        ui.open("scheduleLog", "실행 이력")
        time.sleep(2)
        n_all = ui.vm(TL, NL, "return c.items.length;")
        note = [e.text for e in ui.drv.find_elements(By.ID, "schLogLimitNote") if e.is_displayed()]
        pages = [b.text for b in ui.drv.find_elements(By.CSS_SELECTOR, "#schLogPager .v-pagination__item")]
        ui.js("arguments[0].click();", ui.drv.find_elements(By.CSS_SELECTOR, "#schLogPager .v-pagination__item")[2])
        time.sleep(0.5)
        p3 = ui.vm(TL, NL, "return c.page;")
        ui.shot("u5_log_all")
        ui.vm(TL, NL, "c.filter.scheduleId = arguments[0]; c.loadList();", sid)
        time.sleep(2)
        n_f = ui.vm(TL, NL, "return [c.items.length, c.page];")
        note_f = [e for e in ui.drv.find_elements(By.ID, "schLogLimitNote") if e.is_displayed()]
        pager_f = ui.drv.find_elements(By.CSS_SELECTOR, "#schLogPager .v-pagination")
        print(f"  전체 조회 {n_all}건 안내={note} 페이지={pages} → 3쪽 page={p3}\n  스케줄 고른 조회 (건수, page)={n_f} 안내 보임={bool(note_f)} 페이지 이동 보임={bool(pager_f)}")
        assert n_all == 200 and note == ["최근 200건까지 보입니다 — 스케줄·기간으로 좁혀 조회하세요"] and pages == ["1", "2", "3", "4"] and p3 == 3
        assert n_f[0] < 50 and n_f[1] == 1 and not note_f and not pager_f
    step("U5-1. 실행 이력 — 200건 상한 안내, 50건씩 페이지 이동, 조건으로 좁히면 안내·페이지 사라지고 1쪽 (6.10-33)", _log_screen)

    TR, NR = "tab_termRecommend", "DSTermRecommend"

    def _recommend():
        ui.open("termRecommend", "한글컬럼 일괄 표준화")
        time.sleep(1.5)
        meth = ui.vm(TR, NR, "return [typeof c.$options.methods.resetAll, typeof c.$options.methods.onPaste, 'selectedItems' in c.$data];")
        # 엑셀 — 읽지 못함 / A열 비어 있음 / 정상
        fin = ui.drv.find_element(By.CSS_SELECTOR, f"#{TR} input[type='file']")
        got = []
        for f in ("bad.xlsx", "emptyA.xlsx"):
            p = os.path.join(TMP, f)
            if not os.path.exists(p):
                if f == "bad.xlsx":
                    open(p, "wb").write(b"not an excel file" * 10)
                else:
                    wb = openpyxl.Workbook()
                    wb.active["B1"] = "고객명"
                    wb.save(p)
            ui.js("arguments[0].value = '';", fin)
            fin.send_keys(p)
            got.append(ui.swal("엑셀 파싱 실패"))
            ui.swal_close()
        # 분석 결과 표에 선택 체크박스 없음
        ui.vm(TR, NR, "c.dictId = arguments[0]; c.inputText = 'ZZB3가ZZB3나';", FX["DA"])
        time.sleep(0.5)
        ui.vm(TR, NR, "c.startAnalysis();")
        ui.Wait(ui.drv, 30).until(lambda d: ui.vm(TR, NR, "return !c.isAnalyzing && c.analysisResults.length > 0;"))
        time.sleep(1)
        boxes = len([e for e in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TR} .v-data-table__checkbox, #{TR} th .v-simple-checkbox") if e.is_displayed()])
        ui.shot("u6_recommend")
        # 상태 재계산 — MATCHED + UNRECOGNIZED
        rc = ui.vm(TR, NR, "const it = {words: [{status: 'MATCHED', selected: {wordEngAbrvNm: 'ZZBA'}}, {status: 'UNRECOGNIZED', newWord: {wordEngAbrvNm: 'ZZBQ'}}]};"
                           "c.recalcItemStatus(it); const it2 = {words: [{status: 'MATCHED', selected: {wordEngAbrvNm: 'ZZBA'}}]}; c.recalcItemStatus(it2);"
                           "return [it.status, it.recommendedEngAbrvNm, it2.status];")
        # 형식단어 강제 사전 — 1단어 용어도 마지막 단어 확인
        ui.vm(TR, NR, "c.dictId = arguments[0];", FX["DB"])
        time.sleep(0.5)
        ui.vm(TR, NR, "const it = {inputNm: 'ZZB3다', termName: 'ZZB3다', status: 'AUTO', words: [{wordNm: 'ZZB3다', status: 'MATCHED', "
                      "selected: {wordId: arguments[0], wordNm: 'ZZB3다', wordEngAbrvNm: 'ZZBC'}}]}; c.editItem(it); c.confirmEdit();", FX["W3"])
        w = ui.swal("형식단어")
        ui.swal_close()
        dlg = ui.vm(TR, NR, "return c.editDialog;")
        ui.vm(TR, NR, "c.editDialog = false;")
        print(f"  지운 메서드·데이터(resetAll, onPaste, selectedItems)={meth}\n  엑셀 읽기 실패 창={[g.replace(chr(10), ' / ') for g in got]}"
              f"\n  분석 결과 표 체크박스={boxes} / 재계산(MATCHED+UNRECOGNIZED → 상태, 약어 / MATCHED 만)={rc}\n  1단어 용어 확인 → {w.splitlines()[0]!r} 창 그대로={dlg}")
        assert meth == ["undefined", "undefined", False]
        assert got[0].startswith("엑셀 파싱 실패") and "첫 번째 시트 A열에서 컬럼명을 찾지 못했습니다." in got[1]
        assert boxes == 0 and rc == ["PARTIAL", "ZZBA_ZZBQ", "AUTO"] and w.startswith("형식단어 필요") and dlg
    step("U6-1. 한글컬럼 일괄 표준화 — 쓰지 않던 선택·초기화·붙여넣기 정리, 엑셀 읽기 실패·A열 빈 안내, 미인식 단어 재계산, 1단어 용어 형식단어 확인 (7.5-9a·9b·11·10·5b)", _recommend)

    TH, NH = "tab_term_resolve_history", "DSTermResolveHistory"

    def _resolve_screen():
        ui.js("performance.setResourceTimingBufferSize(5000); performance.clearResourceTimings();")
        ui.open("termResolveHistory", "한글 변환 이력")
        time.sleep(2)
        calls = ui.js("return performance.getEntriesByType('resource').filter(e => e.name.indexOf('api/termResolve/list') >= 0).length;")
        ui.vm(TH, NH, "c.searchInput = 'ZZB3변환'; c.load();")
        time.sleep(2)
        n = ui.vm(TH, NH, "return [c.items.length, c.pageCount];")
        wrap = [e for e in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TH} .pagination_wrap") if e.is_displayed()]
        btns = wrap[0].find_elements(By.CSS_SELECTOR, ".v-pagination__item") if wrap else []
        ui.js("arguments[0].click();", btns[1])
        time.sleep(0.8)
        rows2 = len([tr for tr in ui.drv.find_elements(By.CSS_SELECTOR, "#term_resolve_table tbody tr") if tr.is_displayed()])
        ui.shot("u7_resolve")
        heads = [th.text.strip() for th in ui.drv.find_elements(By.CSS_SELECTOR, "#term_resolve_table thead th")]
        mono = len(ui.drv.find_elements(By.CSS_SELECTOR, "#term_resolve_table tbody span[style*='monospace']"))
        changed = re.search(r"한글명 변경 (\d+)건", ui.drv.find_element(By.ID, TH).text)
        ui.vm(TH, NH, "c.itemsPerPage = 50;")
        time.sleep(0.8)
        after = (ui.vm(TH, NH, "return c.page;"), bool([e for e in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TH} .pagination_wrap") if e.is_displayed()]))
        print(f"  탭 열 때 목록 조회 {calls}번 / 검색 (건수, 페이지 수)={n} 페이지 이동 보임={bool(wrap)} → 2쪽 행 수={rows2}\n  열={heads} 대상 칸 글꼴={mono} "
              f"한글명 변경={changed.group(1) if changed else None}건 / 50건씩 바꾼 뒤 (page, 페이지 이동 보임)={after}")
        assert calls == 1 and n == [27, 2] and wrap and rows2 == 7 and mono == 0 and changed and changed.group(1) == "24" and after == (1, False)
    step("U7-1. 한글 변환 이력 — 첫 진입 조회 1번, 20건 넘으면 페이지 이동, 대응 열 없는 대상 칸 없음, 변경 건수 (7.5-13a·13b·13c·13d)", _resolve_screen)

    def _resolve_user():
        ui.login(FX["su"], FX["suPw"])
        ui.open("termResolveHistory", "한글 변환 이력")
        time.sleep(2)
        n = ui.vm(TH, NH, "return [c.items.length, c.pageCount, c.items.every(i => i.changeUserId === arguments[0])];", FX["su"])
        ui.shot("u7_resolve_user")
        print(f"  일반 사용자 (건수, 페이지 수, 모두 본인)={n}")
        assert n[0] >= 23 and n[1] >= 2 and n[2]
    step("U7-2. 한글 변환 이력 — 일반 사용자는 본인 이력만, 페이지 이동 (7.5-13a)", _resolve_user)

    TM, NM = "tab_my_dm_changes", "MMMyDataModelChanges"

    def _my_changes():
        s_su = FX.get("s_su") or login_api(FX["su"], FX["suPw"])
        r = post(s_su, "/api/dm/saveAttrs", {"dataModelId": FX["mDraft"], "objNm": "ZZ_DRAFT", "objOwner": "",
                                             "attrs": [{"mode": "ADD", "attrNm": "C_DRAFT", "attrNmKr": "임시저장", "dataType": "VARCHAR", "dataLen": 10, "nullableYn": "Y"}]})
        assert code(r) == 200, r
        ui.open("myDmChanges", "내 변경 신청")
        time.sleep(2)
        tables = ui.js(f"return [...document.querySelectorAll('#{TM} .v-data-table')].map(t => [...t.querySelectorAll('thead th')].map(th => th.innerText.trim()));")
        drafts = ui.vm(TM, NM, "return c.drafts.length;")
        ui.shot("u8_my_changes")
        print(f"  DRAFT {drafts}건 / 표 머리글={tables}")
        assert drafts >= 1 and tables and [h for h in tables[0] if h] == ["유형", "대상"]
    step("U8-1. 내 변경 신청 — DRAFT 표는 유형·대상만 (신청 시각 열 없음) (114 §2 권고 — DRAFT 표)", _my_changes)

    ui.login("space")
    TC, NC = "tab_datamodelCollection", "DSDatamodelCollection"

    def _sys_change():
        ui.open("datamodelCollection", "데이터 모델 관리")
        ui.vm(TC, NC, "c.searchModel = arguments[0]; c.getDataModel();", FX["mSysNm"])
        time.sleep(2.5)
        link = [a for a in ui.drv.find_elements(By.CSS_SELECTOR, f"#{TC} tbody td span.ndColor--text") if a.text.strip() == FX["mSysNm"]]
        ui.js("arguments[0].click();", link[0])
        time.sleep(2)
        ui.click_text(f"#{TC} button", "수정")
        time.sleep(2)

        def state():
            return ui.vm(TC, NC, "return [c.update_dataModelSysCd, c.update_dataSource, c.update_dsNotice];")

        def hint():
            return [e.text for e in ui.drv.find_elements(By.CSS_SELECTOR, ".v-dialog--active .v-messages__message") if e.is_displayed() and e.text.strip()]
        s0 = state()
        box = ui.drv.find_element(By.CSS_SELECTOR, "#updateDataModelSysCd")
        ui.js("arguments[0].click();", box)
        box.send_keys(FX["sys2Nm"])
        time.sleep(1.2)
        ui.click_text(".menuable__content__active .v-list-item", FX["sys2Nm"], exact=False)
        time.sleep(2)
        s1_, h1 = state(), hint()
        ui.shot("u9_sys_change")
        ui.vm(TC, NC, "c.update_dataSource = arguments[0]; c.update_dsNotice = '';", FX["dsUi"])   # 데이터 소스를 다시 고름 (@change 와 같다)
        ui.vm(TC, NC, "c.update_dataModelSysCd = arguments[0];", FX["sys1"])
        time.sleep(2)
        s2_, h2 = state(), hint()
        ui.vm(TC, NC, "c.update_dataModelSysCd = null;")
        time.sleep(1.5)
        s3_ = state()
        ui.vm(TC, NC, "c.updateDataModelModal = false;")
        time.sleep(1)
        ui.click_text(f"#{TC} button", "수정")
        time.sleep(2)
        s4_ = state()
        ui.vm(TC, NC, "c.updateDataModelModal = false;")
        print(f"  창 열 때={s0}\n  속하지 않는 시스템으로 → {s1_} 안내={h1}\n  데이터 소스 다시 고르고 속한 시스템으로 → {s2_} 안내={h2}"
              f"\n  시스템 비움 → {s3_} / 다시 열 때={s4_}")
        m = "선택한 시스템에 속하지 않는 데이터 소스라 비웠습니다."
        assert s0 == [FX["sys1"], FX["dsUi"], ""] and s1_ == [FX["sys2"], None, m] and m in h1
        assert s2_ == [FX["sys1"], FX["dsUi"], ""] and m not in h2 and s3_ == [None, FX["dsUi"], ""] and s4_ == [FX["sys1"], FX["dsUi"], ""]
    step("U9-1. 데이터 모델 수정 창 — 시스템을 바꾸면 그 시스템에 없는 데이터 소스만 비우고 안내, 열 때·비울 때는 그대로 (114 §2 권고 — 시스템 변경)", _sys_change)

    def _domain_len():
        ui.open("domain", "도메인")
        time.sleep(1.5)
        res = []
        for tp, ln in (("BINARY", ""), ("VARBINARY", 0), ("BINARY", 16), ("DATE", ""), ("VARCHAR", "")):
            ok = ui.vm("tab_domain", "DSDomain", "return c.domainLenOk(arguments[0], arguments[1]);", tp, ln)
            t = None
            if not ok:
                t = ui.swal("데이터 타입").splitlines()[0]
                ui.swal_close()
            res.append((tp, ln, ok, t))
        print(f"  {res}")
        assert [r[2] for r in res] == [False, False, True, True, False]
        assert res[0][3] == "데이터 타입 BINARY 은(는) 데이터 길이를 1 이상으로 입력해야 합니다." and res[1][3].startswith("데이터 타입 VARBINARY")
    step("U10-1. 도메인 화면 길이 검사에 BINARY·VARBINARY (114 §2 권고 — 도메인 길이 필수 타입)", _domain_len)


# ---------------------------------------------------------------- 정리
def purge():
    """이름으로 남은 시험 데이터를 지운다 (--purge, 정리 실패 대비)"""
    dms = psql("select dm_id from quality.tb_data_model where dm_nm like 'ZZB3%'").split()
    for dm in dms:
        drop_model(dm)
    for d in psql("select dict_id from quality.tb_std_dict where dict_nm like 'ZZB3사전%' order by parent_dict_id nulls last").split():
        drop_dict(d)
    psql("""delete from quality.tb_change_history where change_id in (select change_id from quality.tb_change_history_detail where target_nm like 'ZZB3%')
              or target_nm like 'ZZB3%' or summary like '%ZZB3%';
            delete from quality.tb_change_history_detail where target_nm like 'ZZB3%';
            delete from quality.tb_aprv_stats where req_item_nm like 'ZZB3%';
            delete from quality.tb_diag_schedule_log where schedule_id in (select schedule_id from quality.tb_diag_schedule where schedule_nm like 'ZZB3%');
            delete from quality.tb_diag_schedule where schedule_nm like 'ZZB3%';
            delete from quality.tb_diag_job where diag_job_id like 'ZB3%';
            delete from quality.tb_struct_diag_history where diag_id like 'ZB3%';
            delete from quality.tb_gov_meta_send where send_id like 'ZB3%';
            delete from quality.tb_batch_run where run_id like 'ZB3%';
            delete from quality.tb_term_resolve_history where input_kr_nm like 'ZZB3%';
            delete from quality.tb_word where word_nm like 'ZZB3%';
            delete from quality.tb_sys_info where sys_nm like 'ZZB3%';
            delete from ndata.tb_data_source where dsn like 'ZZB3%';
            delete from quality.tb_user_role where user_id like 'zzb3%';
            delete from quality.tb_role_perm where role_id in (select role_id from quality.tb_role where role_nm like 'ZZB3%');
            delete from quality.tb_role where role_nm like 'ZZB3%';
            delete from quality.tb_user where user_id like 'zzb3%';""")
    for s in psql("select nspname from pg_namespace where nspname like 'zzb3%'").split():
        psql(f"drop schema if exists {s} cascade")
    for r in psql("select rolname from pg_roles where rolname like 'zzb3%'").split():
        psql(f"drop role if exists {r}")


def leftovers(start):
    return psql(f"select (select count(*) from quality.tb_data_model where dm_nm like 'ZZB3%')||'/'||"
                f"(select count(*) from quality.tb_std_dict where dict_nm like 'ZZB3%')||'/'||"
                f"(select count(*) from quality.tb_diag_schedule where schedule_nm like 'ZZB3%')||'/'||"
                f"(select count(*) from quality.tb_diag_job where diag_job_id like 'ZB3%' or dm_id not in (select dm_id from quality.tb_data_model))||'/'||"
                f"(select count(*) from quality.tb_term_resolve_history where input_kr_nm like 'ZZB3%')||'/'||"
                f"(select count(*) from ndata.tb_data_source where dsn like 'ZZB3%')||'/'||"
                f"(select count(*) from quality.tb_sys_info where sys_nm like 'ZZB3%')||'/'||"
                f"(select count(*) from quality.tb_user where user_id like 'zzb3%')||'/'||"
                f"(select count(*) from pg_namespace where nspname like 'zzb3%')||'/'||"
                f"(select count(*) from pg_roles where rolname like 'zzb3%')||'/'||"
                f"(select count(*) from quality.tb_diag_job where status in ('READY','RUNNING'))||'/'||"
                f"(select count(*) from quality.tb_diag_schedule_log where exec_status='RUNNING')")


def main(only):
    if "--purge" in sys.argv:
        purge()
        print("남은 시험 데이터 모델/사전/스케줄/Job/변환 이력/데이터 소스/시스템/사용자/스키마/역할/진행 중 Job/RUNNING LOG =", leftovers(""))
        return
    sa = login_api("space")
    start = time.strftime("%Y%m%d%H%M%S")
    assert psql("select count(*) from quality.tb_aprv_line where use_yn='Y'") == "0", "켜진 결재선이 있다 — 시험이 표준을 바로 등록하므로 멈춘다"
    assert psql("select count(*) from quality.tb_diag_job where status in ('READY','RUNNING')") == "0", "진행 중인 진단 Job 이 있다 — S2 재기동 정리가 건드리므로 멈춘다"
    groups = [("S2", s2), ("S1", s1), ("S3", s3), ("S4", s4), ("S5", s5), ("R", r_check), ("S6", s6)]
    try:
        step("S0. 시험 데이터 준비 (사전·도메인·용어·단어, 사용자, 시험 스키마·데이터 소스·시스템, 모델·수집)", lambda: setup(sa))
        su = sm = None
        for nm, fn in groups:
            if only and nm not in only:
                continue
            if nm == "R" and "T0" not in FX:
                continue
            fn(sa, su, sm)
    finally:
        if UI.inst:
            UI.inst.quit()
            errs = [e for e in UI.inst.errs if e]
            print(f"\n=== 화면 콘솔 오류 {len(errs)}건")
            for e in errs[:20]:
                print("  ", e[:200])
            results.append(("U11. 화면 콘솔 오류 없음", "PASS" if not [e for e in errs if re.search(r"TypeError|ReferenceError|Vue warn|rejection", e)] else "FAIL"))
        if not exec_up():
            print("  q-executor 가 내려가 있어 다시 올린다")
            start_executor()
        if KEEP:
            print("\n=== 정리 건너뜀 (--keep) — 끝나면 python verify_110_batch3.py --purge")
        else:
            print("\n=== 정리")
            for fn in reversed(cleanup):
                try:
                    fn()
                except Exception as e:
                    print("  정리 실패:", e)
            purge()
            psql(f"delete from quality.tb_notification where cret_dt >= '{start}'")
            print(f"  남은 시험 데이터 모델/사전/스케줄/Job/변환 이력/데이터 소스/시스템/사용자/스키마/역할/진행 중 Job/RUNNING LOG = {leftovers(start)}")
        print(f"  화면 캡처: {TMP}")


if __name__ == "__main__":
    t0 = time.time()
    main({a for a in sys.argv[1:] if not a.startswith("--")})
    elapsed = time.time() - t0
    p = sum(1 for _, st in results if st == "PASS")
    f = sum(1 for _, st in results if st == "FAIL")
    print(f"\n{'=' * 60}\n결과: {p} PASS / {f} FAIL  ({elapsed:.0f}초)\n{'=' * 60}")
    for n, st in results:
        print(f"  [{st}] {n}")
    sys.exit(0 if f == 0 else 1)
