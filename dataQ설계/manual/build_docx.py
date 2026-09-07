# -*- coding: utf-8 -*-
"""GS인증 제출용 사용자 매뉴얼(.docx) 렌더러.

본문은 manual_content.py 를 그대로 쓴다. 이 파일은 그걸 GS 제출 서식
(표지 · 판권 · 목차 · 서문 · 사용방법 · 부록 · 찾아보기)의 Word 문서로 만든다.

  build_manual.py  → index.html / artifact.html   (웹 · PDF)
  build_docx.py    → NavidM Meta 사용자 매뉴얼_v1.0.docx  (GS 제출)

실행: python build_docx.py
"""
import html
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from manual_content import (  # noqa: E402
    ERROR_DICT,
    KNOWN_LIMITS,
    PARTS,
    SCREEN_INDEX,
)

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(os.path.dirname(os.path.dirname(HERE)), "GS인증",
                   "NavidM Meta 사용자 매뉴얼_v1.0.docx")

PRODUCT = "NavidM Meta"
RELEASE = "Release 1.0"
COMPANY = "㈜나래데이터"
YEAR = "2026"

KO = "맑은 고딕"
EN = "Malgun Gothic"
INK = RGBColor(0x1A, 0x1D, 0x28)
MUTED = RGBColor(0x5A, 0x61, 0x78)
BRAND = RGBColor(0x24, 0x4A, 0x8A)

ROLE_LABEL = {"admin": "관리자 전용", "user": "일반 사용자", "both": "관리자 · 일반 사용자"}


# ───────────────────────────────────────────── 저수준 헬퍼
def _font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KO)
    return run


def _shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def _space(p, before=0, after=4, line=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if line:
        p.paragraph_format.line_spacing = line
    return p


def norm(text):
    """본문은 내부 표기 'Navid Meta' 를 쓴다. GS 제출본은 정식 제품명으로 통일."""
    return re.sub(r"Navid Meta(?!\w)", PRODUCT, text)


def rich(p, text, size=10.5, color=INK, base_bold=False):
    """manual_content 의 <b> 태그와 HTML 엔티티를 런으로 변환."""
    for chunk in re.split(r"(<b>.*?</b>)", norm(text)):
        if not chunk:
            continue
        bold = chunk.startswith("<b>")
        body = re.sub(r"</?b>", "", chunk)
        _font(p.add_run(html.unescape(body)), size, bold or base_bold, color)
    return p


def plain(text):
    return html.unescape(re.sub(r"<[^>]+>", "", norm(text)))


def field(p, instr):
    """Word 필드(TOC · PAGE 등) 삽입."""
    r = p.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    r._r.append(f1)
    r._r.append(it)
    r._r.append(f2)
    r._r.append(f3)
    return r


# 주의: w:updateFields 는 넣지 않는다. 넣으면 Word 가 문서를 열 때마다
# "다른 파일을 참조하는 필드가 있습니다. 업데이트할까요?" 를 묻는다.
# 대신 finalize_toc.py 가 Word 로 한 번 열어 목차를 채워 넣고 저장한다.


# ───────────────────────────────────────────── 문단 빌더
def heading(doc, text, level, page_break=False):
    style = {0: "Title", 1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(style=style)
    if page_break:
        p.paragraph_format.page_break_before = True
    size = {0: 24, 1: 18, 2: 14, 3: 12}[level]
    _font(p.add_run(text), size, True, BRAND if level <= 2 else INK)
    _space(p, before=14 if level > 1 else 6, after=6)
    return p


def body(doc, text, size=10.5, color=INK, indent=0, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    _space(p, after=after, line=1.4)
    rich(p, text, size, color)
    return p


def bullet(doc, text, indent=0.6, marker="▪"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    _space(p, after=3, line=1.35)
    _font(p.add_run(marker + " "), 10.5, False, BRAND)
    rich(p, text)
    return p


def numbered(doc, i, text, indent=0.75):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    _space(p, after=3, line=1.35)
    _font(p.add_run("%d) " % i), 10.5, True, BRAND)
    rich(p, text)
    return p


def label_block(doc, label, color="EEF2F9"):
    """'먼저 필요한 것' 같은 소제목 줄."""
    p = doc.add_paragraph()
    _space(p, before=6, after=2)
    _font(p.add_run(label), 10, True, BRAND)
    return p


def table(doc, head, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        _shade(c, "E7EDF7")
        _space(c.paragraphs[0], before=2, after=2)
        _font(c.paragraphs[0].add_run(plain(h)), 9.5, True, BRAND)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            _space(cells[i].paragraphs[0], before=2, after=2)
            rich(cells[i].paragraphs[0], str(v), 9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    _space(doc.add_paragraph(), after=6)
    return t


def picture(doc, name, caption=None, width=15.6):
    path = os.path.join(ASSETS, name + ".png")
    if not os.path.exists(path):
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(p, before=6, after=2)
    p.add_run().add_picture(path, width=Cm(width))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _space(c, after=8)
        _font(c.add_run(caption), 9, False, MUTED)
    return True


def header_footer(section, text):
    section.different_first_page_header_footer = True  # 표지에는 머리말·쪽번호 없음
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(hp.add_run(text), 9, False, MUTED)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(fp, " PAGE ")
    for r in fp.runs:
        _font(r, 9, False, MUTED)


# ───────────────────────────────────────────── 서문 데이터
DATASOURCES = [
    ["Oracle (SID / Service Name)", "지원", "지원"],
    ["PostgreSQL", "지원", "지원"],
    ["Cubrid", "지원", "미지원"],
    ["MariaDB", "지원", "미지원"],
    ["Tibero · SQLServer · Altibase · Goldilocks", "수집 불가", "수집 불가"],
]

ARCH_TERMS = [
    ["표준 사전",
     "단어 · 용어 · 도메인 · 코드로 이루어진 데이터 표준의 정의. 모든 진단의 기준이 된다."],
    ["단어",
     "용어를 구성하는 최소 단위. 표준단어명 · 영문명 · 영문약어명 · 형식단어 여부 · "
     "도메인 분류 · 이음동의어 · 금칙어를 관리한다."],
    ["용어",
     "단어를 조합해 만든 컬럼 수준의 표준 이름. 마지막 단어가 형식단어이며 도메인이 연결된다."],
    ["도메인",
     "데이터 타입과 길이의 표준 정의. 용어에 연결되어 컬럼의 타입 · 길이를 규정한다."],
    ["데이터 모델",
     "진단 대상이 되는 테이블 · 컬럼 · 인덱스 · 제약조건의 집합. "
     "운영 DB 수집 또는 설계 파일 임포트로 만든다."],
    ["표준 진단",
     "수집된 데이터 모델을 표준 사전과 대조해 컬럼마다 이슈 유형을 판정하는 작업."],
    ["구조 변경 진단",
     "직전 수집 시점의 스냅샷과 현재 운영 DB 를 대조해 추가 · 변경 · 삭제를 찾아내는 작업."],
    ["승인 워크플로",
     "일반 사용자가 신청하고 관리자가 승인해야 표준 사전과 데이터 모델에 반영되는 절차."],
]

PROCESSES = [
    ["q-center", "화면 · API · 승인",
     "사용자 브라우저의 요청을 받는 웹 서버. 표준 사전 CRUD, 승인 워크플로, "
     "진단 결과 조회를 담당한다. 시간이 걸리는 작업은 직접 처리하지 않고 "
     "q-executor 에 REST API 로 넘긴다."],
    ["q-executor", "수집 · 진단 · 스케줄",
     "백그라운드 워커. 대상 DB 접속이 필요한 모델 수집과 구조 변경 진단, "
     "시간이 걸리는 표준 진단, 예약 실행을 맡는다. 화면이 멈추지 않도록 분리되어 있다."],
    ["메타 데이터베이스", "PostgreSQL 13 (TimescaleDB 2.10)",
     "표준 사전, 수집된 모델 스냅샷, 진단 결과, 변경 이력을 저장한다. "
     "진단 대상이 되는 운영 DB 와는 별개다."],
    ["대상 운영 데이터베이스", "진단 대상",
     "모델 수집과 구조 변경 진단 시에만 JDBC 로 조회한다. 표준 진단은 접속하지 않는다. "
     "쓰기는 관리자가 DDL 반영을 실행했을 때만 일어난다."],
]

ENV_CLIENT = [
    ["OS", "MS Windows 10 (64 bit) 이상"],
    ["Web Browser", "Google Chrome 96.0 이상"],
    ["RAM", "8.00 GB 이상"],
    ["해상도", "1920 × 1080 이상 권장"],
]

ENV_SERVER = [
    ["OS", "RedHat(CentOS) 7.6 이상 (64 bit)"],
    ["CPU", "Intel Quad Core 2.4 GHz 이상"],
    ["RAM", "16.00 GB 이상"],
    ["HDD", "100 GB 이상"],
    ["SW", "docker 20.10.6 이상, OpenJDK(Temurin) 8"],
    ["DB", "PostgreSQL 13 (TimescaleDB 2.10)"],
]

# 찾아보기. 두 번째 값이 "@..." 면 고정 참조, 아니면 manual_content 의 작업 id.
# 작업 번호는 본문 렌더 순서에서 자동 산출하므로 작업이 추가·삭제돼도 어긋나지 않는다.
INDEX_TERMS = [
    ("영문", [
        ("DDL 내려받기", "t13b"), ("ERD", "t13c"), ("ERwin XML 임포트", "t12"),
        ("JDBC 조회", "@서문"), ("q-center", "@서문"), ("q-executor", "@서문"),
        ("XMI 2.1 임포트", "t12"),
    ]),
    ("ㄱ", [
        ("검색 (통합)", "t25"), ("구조 변경 진단", "t23"), ("권한", "t2"),
        ("금칙어", "t4"), ("기본값 보존", "t13"), ("공지사항", "t28"),
    ]),
    ("ㄷ", [
        ("단어 등록", "t4"), ("대시보드", "t1b"), ("데이터 모델 수집", "t11"),
        ("데이터 소스 등록", "t10"), ("도메인 등록", "t5"), ("도메인 분류", "t3"),
    ]),
    ("ㅁ", [("모델 변경 신청", "t21"), ("모델 구조 조회", "t13b")]),
    ("ㅂ", [("반려", "t9"), ("비밀번호 변경", "t26")]),
    ("ㅅ", [
        ("사용자 관리", "t27"), ("설계 파일 임포트", "t12"), ("스케줄 진단", "t24"),
        ("승인 (관리자)", "t9"), ("승인 대기 (사용자)", "t8"),
        ("시스템 요구사항", "@서문"),
    ]),
    ("ㅇ", [
        ("업무영역", "t27"), ("엑셀 일괄 등록", "t7b"), ("오류 메시지", "@부록 B"),
        ("용어 등록", "t7"), ("용어 미존재", "t16"), ("이음동의어", "t4"),
    ]),
    ("ㅈ", [
        ("제외 관리", "t20"), ("준수율", "t15"), ("지원 데이터소스", "@서문"),
        ("진단 결과", "t15"), ("진단 실행", "t14"),
    ]),
    ("ㅋ", [("코드 등록", "t6"), ("컬럼 변경 신청", "t21"),
            ("컬럼 변경 승인·DB 반영", "t22")]),
    ("ㅌ", [("타입·길이 불일치", "t18")]),
    ("ㅍ", [("표준 사전", "@서문"), ("표준 진단", "t14")]),
    ("ㅎ", [
        ("한글명 불일치", "t17"), ("한글 컬럼 일괄 표준화", "t19"),
        ("화면 목록", "@부록 A"), ("환경 제약", "@부록 C"),
    ]),
]


def task_numbers():
    """manual_content 등장 순서 → 작업 번호."""
    out, n = {}, 0
    for _, _, ts in PARTS:
        for t in ts:
            n += 1
            out[t["id"]] = n
    return out


# ───────────────────────────────────────────── 본문 조립
def build():
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = EN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), KO)

    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.4)
        s.top_margin = Cm(2.4)
        s.bottom_margin = Cm(2.2)

    tasks = [t for _, _, ts in PARTS for t in ts]

    # ── 표지
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(cover, before=170, after=0)
    _font(cover.add_run("NARAE DATA"), 15, True, MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(p, before=26, after=0)
    _font(p.add_run("%s User's Manual" % PRODUCT), 30, True, BRAND)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(p, before=10, after=0)
    _font(p.add_run("데이터 표준화 진단 및 메타데이터 관리 솔루션"), 12, False, MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(p, before=60, after=0)
    _font(p.add_run(RELEASE), 14, True, INK)

    # ── 판권
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    _space(p, before=0, after=14)
    _font(p.add_run("%s %s User's Manual" % (COMPANY.strip("㈜"), PRODUCT)), 12, True, INK)

    body(doc, RELEASE, 10.5, MUTED, after=14)
    body(doc, "Copyright ⓒ %s NARAE DATA Corporation." % YEAR, 10.5)
    body(doc, "All Rights Reserved.", 10.5, after=14)
    body(doc, "본 문서의 저작권은 %s에 있습니다. 당사의 동의 없이 무단으로 복제 또는 "
              "전용할 수 없습니다." % COMPANY, 10.5, after=18)
    body(doc, COMPANY, 10.5)
    body(doc, "07207", 10.5, MUTED, after=0)
    body(doc, "서울특별시 영등포구 선유서로25길 28, 703호 (양평동2가, 영등포디스테이트)",
         10.5, MUTED, after=0)
    body(doc, "전화: 070-8861-6806     팩스: 02-6455-9089", 10.5, MUTED, after=0)
    body(doc, "e-mail: support@naraedata.com     homepage: http://www.naraedata.com",
         10.5, MUTED)

    # ── 목차 (Heading 스타일을 쓰지 않는다 — 목차가 자기 자신을 항목으로 잡는다)
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    _space(p, before=6, after=6)
    _font(p.add_run("목  차"), 18, True, BRAND)
    _space(doc.add_paragraph(), after=6)
    tp = doc.add_paragraph()
    field(tp, ' TOC \\o "1-3" \\h \\z \\u ')

    # ── 서문
    heading(doc, "서문", 1, page_break=True)

    heading(doc, "소개", 2)
    body(doc, "“나비드엠 메타”(이하 “%s”)는 기업 및 공공기관이 보유한 데이터베이스의 "
              "테이블 · 컬럼 메타데이터를 수집하여 데이터 표준 사전과 대조하고, 표준을 "
              "벗어난 항목을 유형별로 판정하여 조치까지 연결하기 위한 데이터 표준화 및 "
              "메타데이터 관리 소프트웨어이다." % PRODUCT)
    body(doc, "표준화 지침과 단어 사전이 문서로 존재하더라도, 운영 데이터베이스에 실제로 "
              "만들어진 컬럼이 그 사전을 따랐는지는 컬럼을 하나씩 대조해야 알 수 있다. "
              "%s 는 이 대조를 자동화하여 격차를 수치로 측정하고, 이슈 유형마다 다른 "
              "조치 경로를 제공한다." % PRODUCT)

    heading(doc, "지원하는 데이터소스", 2)
    body(doc, "%s 가 지원하는 데이터소스는 아래와 같다. 연결이 되는 것과 모델 수집이 "
              "되는 것은 다르다." % PRODUCT)
    table(doc, ["DBMS", "테이블 · 컬럼 수집", "인덱스 · 제약조건 수집"],
          DATASOURCES, widths=[7.4, 4.3, 4.3])
    bullet(doc, "Tibero · SQLServer · Altibase · Goldilocks 는 데이터 소스 등록과 연결 "
                "테스트는 통과하지만 모델 수집은 실패한다. 이 경우 ERwin XML 또는 "
                "XMI 2.1 설계 파일 임포트로 모델을 등록한다.")
    bullet(doc, "재수집은 병합 방식이므로 운영 DB 에서 삭제된 대상을 자동으로 지우지 "
                "않는다. 구조 변경 진단으로 확인한 뒤 정리한다.")

    heading(doc, "데이터 아키텍처", 2)
    body(doc, "%s 는 아래와 같은 데이터 항목을 통해서 표준화 작업을 관리하고 수행한다."
         % PRODUCT)
    table(doc, ["데이터 항목", "설명"], ARCH_TERMS, widths=[3.6, 12.4])

    heading(doc, "프로세스 아키텍처 및 실행 환경", 2)
    body(doc, "%s 의 서버 프로세스는 컨테이너(Container) 기반으로 운영되며, 화면과 API 를 "
              "담당하는 서버와 무거운 작업을 처리하는 워커, 그리고 메타 정보를 담는 "
              "데이터베이스로 나뉜다. 각 프로세스별 기능은 아래와 같다." % PRODUCT)
    table(doc, ["프로세스", "역할", "기능 설명"], PROCESSES, widths=[3.2, 3.2, 9.6])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(p, before=8, after=2)
    diag = os.path.join(ASSETS, "arch_engine.png")
    if os.path.exists(diag):
        p.add_run().add_picture(diag, width=Cm(15.2))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _space(c, after=10)
        _font(c.add_run("[그림 1] 서버 엔진 아키텍처"), 9, False, MUTED)

    body(doc, "각 엔진 프로세스는 각각 컨테이너로 구성되고 개별적인 재시작과 버전 "
              "업그레이드가 가능하다.")

    heading(doc, "시스템 요구사항", 2)
    bullet(doc, "<b>사용자 PC 환경</b>")
    table(doc, ["항목", "사양"], ENV_CLIENT, widths=[4.0, 12.0])
    bullet(doc, "<b>서버 운영 환경</b>")
    table(doc, ["항목", "사양"], ENV_SERVER, widths=[4.0, 12.0])

    heading(doc, "이 매뉴얼을 읽는 법", 2)
    body(doc, "본 매뉴얼은 화면 단위가 아니라 <b>작업 단위</b>로 구성되어 있다. "
              "하려는 일을 목차에서 찾으면 되고, 처음부터 순서대로 읽지 않아도 된다. "
              "각 작업은 아래 순서로 서술된다.")
    for k, v in [
        ("먼저 필요한 것", "그 작업을 시작하기 전에 끝나 있어야 하는 선행 작업"),
        ("하는 법", "화면에 실제로 있는 라벨 기준의 조작 순서"),
        ("이렇게 되면 성공", "정상 처리되었을 때 화면에 나타나는 상태"),
        ("자주 겪는 문제", "증상 · 원인 · 조치"),
    ]:
        bullet(doc, "<b>%s</b> — %s" % (k, v))
    body(doc, "작업 제목 옆에는 그 작업을 수행할 수 있는 권한이 표시된다.")

    # ── 1. 사용방법
    heading(doc, "1. %s 사용방법" % PRODUCT, 1, page_break=True)
    body(doc, "이 장은 %s 가 제공하는 기능 및 사용방법에 대해 설명한다. "
              "총 %d개 작업을 %d개 부로 나누어 서술한다."
         % (PRODUCT, len(tasks), len(PARTS)))

    n = 0
    missing = []
    for pi, (title, sub, ts) in enumerate(PARTS, 1):
        heading(doc, "1.%d %s" % (pi, title), 2, page_break=(pi > 1))
        body(doc, sub, 10.5, MUTED, after=8)

        for t in ts:
            n += 1
            heading(doc, "작업 %d. %s" % (n, plain(t["title"])), 3)

            r = doc.add_paragraph()
            _space(r, after=6)
            _font(r.add_run("대상 권한 : "), 9.5, True, MUTED)
            _font(r.add_run(ROLE_LABEL[t["role"]]), 9.5, False, MUTED)

            body(doc, t["goal"])

            if t.get("pre"):
                label_block(doc, "먼저 필요한 것")
                for x in t["pre"]:
                    bullet(doc, x)

            if t.get("steps"):
                label_block(doc, "하는 법")
                for i, s in enumerate(t["steps"], 1):
                    numbered(doc, i, s)

            if t.get("note"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.4)
                _space(p, before=5, after=5)
                _font(p.add_run("※ "), 10, True, BRAND)
                rich(p, t["note"], 10, MUTED)

            if t.get("table"):
                head, rows = t["table"]
                table(doc, head, rows)

            if t.get("shot"):
                if not picture(doc, t["shot"], "[화면] %s" % plain(t["title"])):
                    missing.append(t["shot"])

            if t.get("verify"):
                label_block(doc, "이렇게 되면 성공")
                body(doc, t["verify"], 10.5, indent=0.4)

            if t.get("traps"):
                label_block(doc, "자주 겪는 문제")
                rows = []
                for sym, why, fix in t["traps"]:
                    rows.append([sym, why or "—", fix or "—"])
                table(doc, ["증상", "원인", "조치"], rows, widths=[4.4, 5.6, 6.0])

    # ── 부록
    heading(doc, "부록 A. 화면 찾아보기", 1, page_break=True)
    body(doc, "화면 이름을 알 때 어느 메뉴로 가는지 찾는 표이다.")
    table(doc, ["화면", "메뉴 경로"],
          [[nm, path] for nm, path, _ in SCREEN_INDEX], widths=[5.4, 10.6])

    heading(doc, "부록 B. 오류 메시지 사전", 1, page_break=True)
    body(doc, "화면에 뜨는 메시지를 그대로 찾아보면 된다.")
    table(doc, ["메시지", "뜻", "할 일"],
          [[m, mean, act] for m, mean, act in ERROR_DICT], widths=[5.2, 5.2, 5.6])

    heading(doc, "부록 C. 알려진 제약", 1, page_break=True)
    body(doc, "고칠 수 있는 결함이 아니라, 현재 제품이 의도적으로 하지 않거나 아직 "
              "지원하지 않는 것들이다.")
    table(doc, ["항목", "내용", "영향 범위"],
          [[t_, d, sc] for t_, d, sc in KNOWN_LIMITS], widths=[3.6, 8.4, 4.0])

    # ── 찾아보기
    heading(doc, "찾아보기", 1, page_break=True)
    body(doc, "숫자는 작업 번호이다. 해당 작업은 “1. %s 사용방법” 장에서 찾을 수 있다."
         % PRODUCT, 9.5, MUTED, after=10)
    nums = task_numbers()
    unresolved = []
    for group, items in INDEX_TERMS:
        g = doc.add_paragraph()
        _space(g, before=8, after=3)
        _font(g.add_run(group), 11, True, BRAND)
        for term, ref in items:
            if ref.startswith("@"):
                label = ref[1:]
            elif ref in nums:
                label = "작업 %d" % nums[ref]
            else:
                unresolved.append((term, ref))
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            _space(p, after=1, line=1.2)
            _font(p.add_run(term), 10)
            _font(p.add_run("  " + label), 10, False, MUTED)

    if unresolved:
        print("찾아보기 미해결:", unresolved)

    header_footer(doc.sections[0], "%s %s 사용자 매뉴얼" % (PRODUCT, RELEASE))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    return len(tasks), missing


if __name__ == "__main__":
    cnt, missing = build()
    print("작업 %d개 렌더링" % cnt)
    print("스크린샷 누락:", missing or "없음")
    print("저장:", OUT)
